"""BookForge exporters: imposition, cover, EPUB, DOCX, HTML, TXT."""
import dataclasses
import html
import io
import math
import re
import uuid
import zipfile
from datetime import datetime, timezone

from bf_core import PAPERS, page_size_inches, plain, esc
from bf_pdf import font_roles, apply_case, auto_label, LABELISH, title_layout, _num
import bf_matter


def edited(book, s, key):
    """The edited content of a front/back page (bf_matter.section_blocks), or None while it is automatic."""
    if bf_matter.is_edited(book, key) and bf_matter.override(book, key) is not None:
        return bf_matter.section_blocks(book, s, key)
    return None


PICTURE_SECTIONS = [k for k, v in bf_matter.SECTIONS.items() if v[4]]


def title_look(s):
    """The Layout tab's title settings as the EPUB, Word document and web page use them - read and checked by
    bf_pdf.title_layout, the same function the PDF uses, so every format works from the same values:
    alignments "left" / "center" / "right", positions in % of the text area, sizes in pt (0 = automatic) and
    the chapter drop in inches (none in Continuous, as in the PDF)."""
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    tl = title_layout(s)
    name = lambda a: "left" if a == TA_LEFT else "right" if a == TA_RIGHT else "center"   # noqa: E731
    return {"tp_align": name(tl["tp_align"]), "tp_pos": tl["tp_pos"], "tp_title_size": tl["tp_title_size"],
            "ch_align": name(tl["ch_align"]), "part_align": name(tl["part_align"]), "part_pos": tl["part_pos"],
            "part_size": tl["part_size"],
            "sink": 0.0 if s.get("chapter_start") == "Continuous" else tl["sink"]}


def cover_look(s):
    """The front cover's title settings (Output tab, Cover): the alignment of the title, subtitle and author
    lines ("Same as title page" = the Layout tab's title page alignment) and how far down the front panel the
    title line sits (% of the cover's height; 32 = where it always was)."""
    a = s.get("cover_align", "Same as title page")
    al = {"Left": "left", "Right": "right", "Centered": "center"}.get(a) or title_look(s)["tp_align"]
    return {"align": al, "pos": _num(s.get("cover_pos"), 32.0, 5.0, 85.0)}


def matter_pictures(book, s):
    """The pictures in edited front/back pages (they go into EPUB / web pages with the chapters' ones)."""
    out = []
    for key in PICTURE_SECTIONS:
        for b in edited(book, s, key) or []:
            if b.kind == "image" and b.image:
                out.append(b)
    return out


def front_text_blocks(book, s):
    """The text before chapter 1 as the PDF prints it: the edited page (when it has something in it), or the
    paragraphs read from the source ("Keep text before the first chapter"). [] = none."""
    ed = edited(book, s, "front_text")
    if ed is not None:
        return ed if bf_matter.has_text(ed) else []
    return [b for b in book.front_blocks if b.kind == "para"]


def blocks_text(blocks, s):
    """Plain-text lines for the blocks of an edited front/back page (text export)."""
    out = []
    for b in blocks or []:
        k = b.kind
        if k in ("para", "quote", "head", "subtitle", "flush", "ded", "endline", "tocentry") + MATTER_LINES:
            words = plain((b.text or "").replace("<br/>", "\n"))
            if words.strip():
                out += [words, ""]
        elif k == "meta":
            out.append(plain(b.text or ""))
        elif k == "anote" and not an_mode(s).startswith("Leave"):
            out += [("[" + b.label + "] " if b.label else "") + plain(b.text), ""]
        elif k == "item":
            out.append("  " * b.depth + (b.bullet + " " if b.bullet else "  ") + plain(b.text))
        elif k == "verse":
            out += [plain(x) for x in b.lines] + [""]
        elif k == "scene":
            out += [scene_text(b, s), ""]
    if out and out[-1] != "":
        out.append("")
    return out


# --------------------------------------------------------------------------
# Finished page size on the printed sheets (for the cover)
# --------------------------------------------------------------------------
def trim_size(s):
    pw, ph = page_size_inches(s)
    if s.get("_trim"):
        return s["_trim"]
    if s.get("fmt_impose"):
        # Estimate only (the build uses the size the imposition really gives): a folio on the sheet turned
        # landscape. Any paper preset works here, not just the four oldest ones.
        from bf_impose import sheet_points
        W0, H0 = sheet_points(s)
        SW, SH = max(W0, H0) / 72, min(W0, H0) / 72
        k = min((SW / 2) / pw, SH / ph)
        return pw * k, ph * k
    return pw, ph


# --------------------------------------------------------------------------
# Wraparound cover template
# --------------------------------------------------------------------------
def build_cover(book, s, out_path, n_pages):
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor, black
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.platypus import Paragraph, Frame
    from reportlab.pdfbase.pdfmetrics import stringWidth

    (R, B, I, BI), (HR, HB, _hi, _hbi), _d = font_roles(s, lambda m: None)
    if s.get("heading_bold", True):
        B = HB
    else:
        B = HR
    from reportlab.lib.utils import ImageReader
    from bf_pictures import cover_art, fit_box, COVER_FITS
    tw, th = trim_size(s)
    spine = estimated_spine(s, n_pages)
    if float(s.get("spine_width") or 0) > 0:
        spine = float(s["spine_width"])                   # measured by the user
    W, H = (2 * tw + spine) * 72, th * 72
    c = canvas.Canvas(out_path, pagesize=(W, H))
    c.setTitle(f"{book.title} - cover")
    fill = s.get("cover_image_fit", COVER_FITS[0]) == COVER_FITS[0]
    arts = {}
    for which, x0 in (("back", 0), ("front", (tw + spine) * 72)):
        try:
            data = cover_art(book, s, which)
        except ValueError:
            data = None
        if not data:
            continue
        try:
            img = ImageReader(io.BytesIO(data))
            iw, ih = img.getSize()
            bx, by, bw, bh = fit_box(iw, ih, x0, 0, tw * 72, H, fill)
            c.saveState()
            path = c.beginPath()
            path.rect(x0, 0, tw * 72, H)
            c.clipPath(path, stroke=0, fill=0)
            c.drawImage(img, bx, by, bw, bh, mask="auto")
            c.restoreState()
            arts[which] = True
        except Exception:
            pass
    c.setStrokeColor(HexColor("#bbbbbb"))
    c.setDash(3, 3)
    c.setLineWidth(0.5)
    for x in (tw * 72, (tw + spine) * 72):
        if arts:                                          # short marks only, so the pictures stay clean
            c.line(x, 0, x, 0.25 * 72)
            c.line(x, H - 0.25 * 72, x, H)
        else:
            c.line(x, 0, x, H)
    c.setDash()

    def fit(text, font, maxsize, maxw):
        sz = maxsize
        while sz > 8 and stringWidth(text, font, sz) > maxw:
            sz -= 1
        return sz

    fx0 = (tw + spine) * 72
    cx = fx0 + tw * 72 / 2
    pad = 0.6 * 72
    # Front cover title lines (Output tab, Cover): their alignment ("Same as title page" = the Layout tab's title
    # page alignment) and how far down the front the title line sits (32 % = where it always was). The title
    # never goes so high that it leaves the cover or so low that the subtitle runs into the author line.
    look = cover_look(s)
    t = book.title
    sz = fit(t, B, 40, tw * 72 - 2 * pad)
    s2 = fit(book.subtitle, I, 16, tw * 72 - 2 * pad) if book.subtitle else 0
    s3 = fit(book.author, R, 18, tw * 72 - 2 * pad) if book.author else 0
    low = (H * 0.18 + s3 * 1.6 if book.author else pad) + (sz * 1.3 if book.subtitle else 0)
    y = max(low, min(H * (1 - look["pos"] / 100), H - pad - sz * 0.8))

    def line(text, yy):
        if look["align"] == "left":
            c.drawString(fx0 + pad, yy, text)
        elif look["align"] == "right":
            c.drawRightString(fx0 + tw * 72 - pad, yy, text)
        else:
            c.drawCentredString(cx, yy, text)
    c.setFillColor(black)
    if not arts.get("front") or s.get("cover_text_over_art"):
        c.setFont(B, sz)
        line(t, y)
        if book.subtitle:
            c.setFont(I, s2)
            line(book.subtitle, y - sz * 1.3)
        if book.author:
            c.setFont(R, s3)
            line(book.author, H * 0.18)
    if spine * 72 >= 14:
        label = book.title + (f"    {book.author}" if book.author else "")
        ssz = min(spine * 72 * 0.55, 14)
        ssz = fit(label, B, ssz, H - 72)
        c.saveState()
        c.translate((tw + spine / 2) * 72, H / 2)
        c.rotate(-90)
        c.setFont(B, ssz)
        c.drawCentredString(0, -ssz * 0.35, label)
        c.restoreState()
    blurb = edited(book, s, "blurb")
    if blurb is not None or (s.get("blurb") or "").strip():
        sty = ParagraphStyle("blurb", fontName=R, fontSize=11, leading=15, alignment=TA_JUSTIFY,
                             spaceAfter=8)
        if blurb is not None:                      # edited: its bold, italics, colours and alignment
            from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
            al = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}
            paras = []
            for b in blurb:
                if b.kind == "image" or not plain(b.text or "").strip():
                    continue
                bs = ParagraphStyle(f"blurb{len(paras)}", parent=sty, alignment=al.get(b.align, TA_JUSTIFY))
                markup = re.sub(r'<font face="[^"]*"', "<font", b.text)   # cover fonts: the book's own
                try:
                    paras.append(Paragraph(markup, bs))
                except Exception:
                    paras.append(Paragraph(esc(plain(b.text)), bs))
        else:
            paras = [Paragraph(esc(" ".join(p.split())), sty)
                     for p in re.split(r"\n\s*\n", s["blurb"].strip()) if p.strip()]
        Frame(pad, pad, tw * 72 - 2 * pad, H * 0.75 - pad, showBoundary=0).addFromList(paras, c)
    c.showPage()
    c.save()
    return {"trim": (tw, th), "spine": spine, "width": W / 72, "height": H / 72,
            "measured": float(s.get("spine_width") or 0) > 0, "art": sorted(arts)}


def estimated_spine(s, n_pages):
    """Spine width in inches worked out from the page count and the paper's thickness."""
    return math.ceil(n_pages / 2) * PAPERS.get(s["paper"], 0.004)


# --------------------------------------------------------------------------
# Shared helpers for text exports
# --------------------------------------------------------------------------
def chapter_heads(book, s):
    """Yield (chapter, label, title) with the same labelling rules as the PDF."""
    n = 0
    for ch in book.active():
        title = apply_case(ch.title, s["title_case"])
        if ch.level == 0:
            yield ch, ch.label, title
            continue
        n += 1
        lbl = ch.label
        if not lbl and s["chapter_label"] not in ("None", "") and not LABELISH.match(ch.title + " "):
            lbl = auto_label(s["chapter_label"], n)
        yield ch, lbl, title


def contents_list(book, s, heads, sep=" \u2014 "):
    """The contents list for EPUB, HTML and plain text: [(chapter number in heads or None, text)].
    Follows "Edit contents" (renamed, hidden, moved and added entries). Page numbers are left out
    because these formats reflow. With no edits this is the plain list of chapter headings."""
    items = [(i, f"{lbl}{sep}{title}" if lbl and title else (title or lbl))
             for i, (_ch, lbl, title) in enumerate(heads, 1)]
    if not getattr(book, "toc", None):
        return items
    from bf_pdf import contents_entries, apply_contents_edits
    by_uid = {ch.uid: n for n, (ch, _l, _t) in enumerate(heads, 1)}
    out = []
    for e in apply_contents_edits(contents_entries(book, s), book.toc):
        if e.get("custom"):
            out.append((None, e["text"]))
        elif e["uid"] in by_uid:
            n = by_uid[e["uid"]]
            out.append((n, e["text"] if e.get("edited") else items[n - 1][1]))
    return out


MATTER_LINES = ("btitle", "bsub", "bauth", "bcred", "plate")


def an_mode(s):
    return s.get("an_mode", "Set apart (smaller, indented)")


def scene_text(b, s):
    sb = s.get("scene_break", "* * *")
    if sb == "Blank space":
        return ""
    if sb == "Keep the source's own":
        return plain(b.text) or "* * *"
    return sb


def _font_span(m):
    a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(0)))
    css = []
    if a.get("face"):
        css.append(f"font-family:'{a['face']}'")
    if a.get("size"):
        css.append(f"font-size:{a['size']}" + ("" if a["size"].endswith("%") else "pt"))
    if a.get("color"):
        css.append(f"color:{a['color']}")
    return f'<span style="{";".join(css)}">' if css else "<span>"


def para_style(b):
    """CSS for formatting set in the editor (alignment, indent, space above)."""
    css = []
    if getattr(b, "align", ""):
        css.append(f"text-align:{b.align}")
    if getattr(b, "indent", 0):
        css.append(f"margin-left:{0.25 * int(b.indent):g}in")
    if getattr(b, "space", 0):
        css.append(f"margin-top:{int(b.space)}pt")
    return f' style="{";".join(css)}"' if css else ""


def front_art(book, s):
    """The front cover picture's bytes (see bf_pictures.cover_art), or None if there isn't a usable one."""
    from bf_pictures import cover_art
    try:
        return cover_art(book, s, "front")
    except ValueError:
        return None


def picture_html(b, src):
    """A picture for EPUB / web pages, with the size, alignment and placement chosen in the editor."""
    pcss, icss = [], []
    if b.align in ("left", "center", "right"):
        pcss.append(f"text-align:{b.align}")
    if b.img_place == "full":
        icss.append("width:100%")
    elif b.img_width:
        icss.append(f"width:{min(100, max(5, int(b.img_width)))}%")
    if b.img_place in ("page", "full"):
        pcss += ["page-break-before:always", "page-break-after:always"]
        icss.append("max-height:95vh")
    ps = f' style="{";".join(pcss)}"' if pcss else ""
    ims = f' style="{";".join(icss)}"' if icss else ""
    return f'<p class="img"{ps}><img src="{src}" alt=""{ims} /></p>'


def to_xhtml(markup):
    m = markup.replace("<super>", "<sup>").replace("</super>", "</sup>").replace("<br/>", "<br />")
    m = m.replace("<strike>", "<s>").replace("</strike>", "</s>")
    m = re.sub(r"<font\b[^>]*>", _font_span, m).replace("</font>", "</span>")
    return m.replace("&nbsp;", "&#160;")


def img_type(data):
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png", "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "jpg", "image/jpeg"
    if data[:4] == b"GIF8":
        return "gif", "image/gif"
    if b"<svg" in data[:300]:
        return "svg", "image/svg+xml"
    return None, None


CSS = """body{font-family:Georgia,'Times New Roman',serif;line-height:1.55;margin:0 6%;}
p{text-indent:1.5em;margin:0 0 .45em;text-align:justify;}
.label{text-align:center;margin:3em 0 .3em;font-size:1.1em;letter-spacing:.05em;}
h1,h2{text-align:center;margin:.3em 0 1.2em;}
h3{text-align:center;margin:1.5em 0 .8em;}
.part{margin-top:30%;}
.verse{text-align:center;font-style:italic;margin:1em 0;text-indent:0;}
.quote{margin:1em 2em;text-indent:0;}
.subtitle{text-align:center;font-style:italic;margin:0 2em 1.5em;text-indent:0;}
.scene{text-align:center;margin:1em 0;text-indent:0;}
.item{text-indent:-1.2em;text-align:left;}
.anote{font-size:.88em;margin:0 2em .3em;text-indent:0;text-align:left;}
.anlabel{font-size:.82em;font-weight:bold;margin:1em 2em .2em;text-indent:0;}
.meta{text-indent:-1em;margin:0 0 .3em 1em;text-align:left;}
.titlepage{text-align:center;margin-top:25%;}
.titlepage .t{font-size:2.2em;font-weight:bold;text-indent:0;text-align:center;}
.titlepage p{text-align:center;text-indent:0;}
.ded{text-align:center;font-style:italic;margin-top:25%;text-indent:0;}
.end{text-align:center;margin-top:2em;text-indent:0;}
img{max-width:100%;}
.img{text-align:center;text-indent:0;}
nav ol{list-style:none;} a{color:inherit;}
"""


def title_css(s):
    """Stylesheet lines for the titles BookForge makes, from the Layout tab's title settings (title_look), added
    after CSS so they win: the title page block, chapter titles (number line, title, subtitles), part titles and
    the headings BookForge adds (Contents, Work info, Summary, Notes, the extra front page, end pages: h2, as in
    the PDF they follow the chapter title alignment). Headings typed in the text (h3) keep their own alignment.
    A screen has no fixed page, so positions become space above the title: the title page's top margin grows
    from the standard 25 % (at 20 %) in step with its %, a part's top margin is its % (30 % = as before, % of the
    page width, as CSS works), the chapter drop is padding in ems of the text (the drop / the text size), and
    sizes set on the Layout tab are ems of the text size (0 = automatic keeps the standard sizes)."""
    tl = title_look(s)
    body = _num(s.get("body_size"), 13.0, 4.0, 100.0)

    def em(pt):
        return f"{pt / body:.4g}em"
    tp, ch, pa = tl["tp_align"], tl["ch_align"], tl["part_align"]
    side = {"left": ";margin-left:0", "right": ";margin-right:0"}.get(ch, "")
    out = [f".titlepage{{text-align:{tp};margin-top:{tl['tp_pos'] * 1.25:.4g}%}}",
           f".titlepage p,.titlepage .t{{text-align:{tp}}}",
           f".label,h2{{text-align:{ch}}}",
           ".label{text-indent:0}",                # (the number line lines up with its title, as in the PDF)
           f".subtitle{{text-align:{ch}{side}}}",
           f".part{{margin-top:{tl['part_pos']:.4g}%}}",
           f".part,.part .label,.part h1,.part h2{{text-align:{pa}}}"]
    if tl["sink"]:
        out.append(f".chead{{padding-top:{em(tl['sink'] * 72)}}}")
    if tl["tp_title_size"]:
        out.append(f".titlepage .t{{font-size:{em(tl['tp_title_size'])}}}")
    if tl["part_size"]:
        out.append(f".part h1,.part h2{{font-size:{em(tl['part_size'])}}}")
    return "\n".join(out) + "\n"


def blocks_html(blocks, imgmap, s=None):
    s = s or {}
    out = []
    prev_label = None
    for b in blocks:
        if b.kind != "anote":
            prev_label = None
        if b.kind == "anote":
            if an_mode(s).startswith("Leave"):
                continue
            keep = an_mode(s).startswith("Keep")
            if b.label and b.label != prev_label:
                out.append(f'<p class="{"quote" if keep else "anlabel"}">{esc(b.label)}</p>')
            out.append(f'<p class="{"quote" if keep else "anote"}"{para_style(b)}>{to_xhtml(b.text)}</p>')
            prev_label = b.label
        elif b.kind == "item":
            ind = 1.5 * (b.depth + 1)
            bullet = esc(b.bullet) + "&#160;" if b.bullet else ""
            out.append(f'<p class="item" style="margin-left:{ind}em">{bullet}{to_xhtml(b.text)}</p>')
        elif b.kind == "para":
            out.append(f"<p{para_style(b)}>{to_xhtml(b.text)}</p>")
        elif b.kind == "quote":
            out.append(f'<p class="quote"{para_style(b)}>{to_xhtml(b.text)}</p>')
        elif b.kind == "subtitle":
            out.append(f'<p class="subtitle"{para_style(b)}>{to_xhtml(b.text)}</p>')
        elif b.kind == "verse":
            out.append('<p class="verse">' + "<br />".join(to_xhtml(l) if l else "&#160;"
                                                          for l in b.lines) + "</p>")
        elif b.kind == "head":
            out.append(f"<h3{para_style(b)}>{to_xhtml(b.text)}</h3>")
        elif b.kind == "scene":
            out.append(f'<p class="scene">{esc(scene_text(b, s)) or "&#160;"}</p>')
        elif b.kind == "image" and id(b) in imgmap:
            out.append(picture_html(b, imgmap[id(b)]))
        elif b.kind in MATTER_CSS:                 # title page and other front/back page styles
            cls, css = MATTER_CSS[b.kind]
            if b.kind in ("btitle", "bsub", "bauth", "bcred"):       # the title page alignment (Layout tab)
                tl = title_look(s)
                css = css.replace("text-align:center", f"text-align:{tl['tp_align']}")
                if b.kind == "btitle" and tl["tp_title_size"]:
                    body = _num(s.get("body_size"), 13.0, 4.0, 100.0)
                    css = css.replace("font-size:2.2em", f"font-size:{tl['tp_title_size'] / body:.4g}em")
            mine = para_style(b)[len(' style="'):-1] if para_style(b) else ""
            css = ";".join(x for x in (css, mine) if x)
            css = f' style="{css}"' if css else ""
            out.append(f'<p{cls}{css}>{to_xhtml(b.text or "") or "&#160;"}</p>')
    return "\n".join(out)


# The front/back page styles in EPUB / web pages: (class, CSS); the paragraph's own alignment, indent and
# space (Specific Text Edit) come after, so they win.
_C = "text-indent:0;text-align:center"
MATTER_CSS = {"btitle": ("", _C + ";font-size:2.2em;font-weight:bold;margin:0 0 .4em"),
              "bsub": ("", _C + ";font-style:italic"), "bauth": ("", _C + ";margin-top:1.5em"),
              "bcred": ("", _C + ";font-size:.9em"), "plate": ("", _C + ";font-style:italic"),
              "ded": ("", _C + ";font-style:italic"), "endline": (' class="end"', ""),
              "meta": (' class="meta"', ""), "flush": ("", "text-indent:0"), "tocentry": ("", "text-indent:0")}





def fic_sections(book, s):
    """[(heading, [(css_class, markup)], edited blocks or None)] for the fanfic info / summary / notes pages
    (an edited page is written from its blocks, with their alignment, indents and styles)."""
    out = []
    for key, on, orig, cls in (("work_info", s.get("fic_info", "Leave out") != "Leave out", book.fic_meta, "meta"),
                               ("summary", s.get("fic_summary", "Add a summary page") != "Leave out",
                                book.fic_summary, "flush"),
                               ("work_notes", not an_mode(s).startswith("Leave"), book.fic_notes, "flush")):
        if not on:
            continue
        ed = bf_matter.override(book, key) if bf_matter.is_edited(book, key) else None
        if ed is not None:                         # edited: every line, with its own look
            items = [("meta" if b.kind == "meta" else "flush", b.text) for b in ed
                     if b.kind not in ("image", "verse", "scene") and plain(b.text or "").strip()]
        else:
            items = [(cls, m if key == "work_info" else m.text) for m in orig]
        if items:
            out.append((bf_matter.section_heading(book, s, key), items,
                        [b for b in ed if b.kind not in ("image",)] if ed is not None else None))
    return out


def fic_html(fic, s):
    """The fanfic info / summary / notes pages for EPUB and web pages."""
    out = []
    for h, items, blocks in fic:
        out.append(f"<h2>{esc(h)}</h2>")
        if blocks is not None:            # edited: text without an indent is set apart like the original
            out.append(blocks_html([dataclasses.replace(b, kind="quote") if b.kind == "flush" else b
                                    for b in blocks], {}, s))
        else:
            out += [f'<p class="{"meta" if c == "meta" else "quote"}">{to_xhtml(m)}</p>' for c, m in items]
    return "\n".join(out)


# --------------------------------------------------------------------------
# EPUB 3
# --------------------------------------------------------------------------
def build_epub(book, s, out_path):
    uid = "urn:uuid:" + str(uuid.uuid4())
    lang = book.language or "en"
    heads = list(chapter_heads(book, s))

    def page(title, body):
        return ('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
                f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
                f'xml:lang="{lang}" lang="{lang}"><head><meta charset="utf-8" /><title>{esc(title)}</title>'
                '<link rel="stylesheet" type="text/css" href="style.css" /></head>'
                f"<body>\n{body}\n</body></html>")

    files, manifest, spine, navitems = {}, [], [], []
    imgmap, n_img = {}, 0
    for ch in book.active():
        for b in ch.blocks:
            if b.kind == "image" and b.image:
                ext, mt = img_type(b.image)
                if ext:
                    n_img += 1
                    name = f"images/img{n_img:04d}.{ext}"
                    files["OEBPS/" + name] = b.image
                    manifest.append(f'<item id="img{n_img}" href="{name}" media-type="{mt}"/>')
                    imgmap[id(b)] = name
    for b in matter_pictures(book, s):
        ext, mt = img_type(b.image)
        if ext:
            n_img += 1
            name = f"images/img{n_img:04d}.{ext}"
            files["OEBPS/" + name] = b.image
            manifest.append(f'<item id="img{n_img}" href="{name}" media-type="{mt}"/>')
            imgmap[id(b)] = name
    cover_meta = ""
    art = front_art(book, s)
    if art and img_type(art)[0]:
        ext, mt = img_type(art)          # e-book readers show this picture as the book's cover
        files[f"OEBPS/images/cover.{ext}"] = art
        manifest.append(f'<item id="cover-image" href="images/cover.{ext}" media-type="{mt}" '
                        'properties="cover-image"/>')
        files["OEBPS/cover.xhtml"] = page(book.title, '<div style="text-align:center;margin:0;padding:0">'
                                          f'<img src="images/cover.{ext}" alt="{esc(book.title)}" '
                                          'style="max-width:100%;max-height:100vh" /></div>')
        manifest.append('<item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/>')
        spine.append("cover")
        cover_meta = '<meta name="cover" content="cover-image"/>'
    title_ed = edited(book, s, "title")
    if title_ed is not None:
        tp = ['<div class="titlepage">', blocks_html(title_ed, imgmap, s), "</div>"]
    else:
        tp = ['<div class="titlepage">', f'<p class="t">{esc(book.title)}</p>']
        if book.subtitle:
            tp.append(f"<p><i>{esc(book.subtitle)}</i></p>")
        if book.author:
            tp.append(f"<p>{esc(book.author)}</p>")
        if book.credit:
            tp.append(f"<p>{esc(book.credit)}</p>")
        tp.append("</div>")
    ded_ed = edited(book, s, "dedication")
    if ded_ed is not None:
        if bf_matter.has_text(ded_ed):
            tp.append('<div class="ded">' + blocks_html(ded_ed, imgmap, s) + "</div>")
    elif (s.get("dedication") or "").strip():
        tp.append('<p class="ded">' + esc(s["dedication"].strip()).replace("\n", "<br />") + "</p>")
    files["OEBPS/title.xhtml"] = page(book.title, "\n".join(tp))
    manifest.append('<item id="title" href="title.xhtml" media-type="application/xhtml+xml"/>')
    spine.append("title")
    fic = fic_sections(book, s)
    if fic:
        if any(bl is not None for _h, _i, bl in fic):
            body = fic_html(fic, s)
        else:                                       # (as before)
            body = "".join(f"<h2>{esc(h)}</h2>" + "".join(
                f'<p class="{"meta" if c == "meta" else "quote"}">{to_xhtml(m)}</p>' for c, m in items)
                for h, items, _b in fic)
        files["OEBPS/info.xhtml"] = page("About this work", body)
        manifest.append('<item id="info" href="info.xhtml" media-type="application/xhtml+xml"/>')
        spine.append("info")
    note_ed = edited(book, s, "note")

    def add_note():
        if note_ed is not None and (bf_matter.has_text(note_ed) or s.get("note_title")):
            body = (f"<h2>{esc(s.get('note_title') or '')}</h2>" if s.get("note_title") else "") + \
                blocks_html(note_ed, imgmap, s)
        elif note_ed is None and (s.get("note_text") or "").strip():
            body = (f"<h2>{esc(s.get('note_title') or '')}</h2>" if s.get("note_title") else "") + "".join(
                f"<p>{esc(' '.join(p.split()))}</p>" for p in re.split(r"\n\s*\n", s["note_text"].strip()))
        else:
            return
        files["OEBPS/note.xhtml"] = page(s.get("note_title") or "Note", body)
        manifest.append('<item id="note" href="note.xhtml" media-type="application/xhtml+xml"/>')
        spine.append("note")

    # the PDF's order: work info..., the extra front page before or after the contents (as set), the text
    # before chapter 1, the chapters
    if s.get("note_place") == "Before contents":
        add_note()
    nav_at = len(spine)                              # (the contents go here)
    if s.get("note_place") != "Before contents":
        add_note()
    front = front_text_blocks(book, s)
    if front:
        files["OEBPS/front.xhtml"] = page(book.title, blocks_html(front, imgmap, s))
        manifest.append('<item id="front" href="front.xhtml" media-type="application/xhtml+xml"/>')
        spine.append("front")
    for i, (ch, lbl, title) in enumerate(heads, 1):
        fn = f"ch{i:03d}.xhtml"
        parts = []
        wrap = ch.level == 0 or lbl or title         # part title block / chapter title block (the drop)
        if wrap:
            parts.append('<div class="part">' if ch.level == 0 else '<div class="chead">')
        if lbl:
            parts.append(f'<p class="label">{esc(lbl)}</p>')
        if title:
            parts.append(f'<h{1 if ch.level == 0 else 2}>{esc(title)}</h{1 if ch.level == 0 else 2}>')
        if wrap:
            parts.append("</div>")
        parts.append(blocks_html(ch.blocks, imgmap, s))
        if i == len(heads) and s.get("the_end"):
            end_ed = edited(book, s, "the_end")
            parts.append(blocks_html(end_ed, imgmap, s) if end_ed is not None else '<p class="end">THE END</p>')
        files["OEBPS/" + fn] = page(title or lbl, "\n".join(parts))
        manifest.append(f'<item id="c{i}" href="{fn}" media-type="application/xhtml+xml"/>')
        spine.append(f"c{i}")
        disp = f"{lbl} \u2014 {title}" if lbl and title else (title or lbl)
        navitems.append((fn, disp))
    xm = edited(book, s, "end_matter")
    xhead = bf_matter.section_heading(book, s, "end_matter")
    if xm is not None and (bf_matter.has_text(xm) or xhead):
        files["OEBPS/endpages.xhtml"] = page(xhead or "End", (f"<h2>{esc(xhead)}</h2>" if xhead else "") +
                                             blocks_html(xm, imgmap, s))
        manifest.append('<item id="endpages" href="endpages.xhtml" media-type="application/xhtml+xml"/>')
        spine.append("endpages")
    nav = ("<nav epub:type=\"toc\" id=\"toc\"><h2>" + esc(s.get("toc_title") or "Contents") + "</h2><ol>" +
           "".join(f'<li><a href="{navitems[n - 1][0]}">{esc(t)}</a></li>' if n else f"<li><span>{esc(t)}</span></li>"
                   for n, t in contents_list(book, s, heads)) + "</ol></nav>")
    files["OEBPS/nav.xhtml"] = page("Contents", nav)
    manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>')
    if s.get("toc"):
        spine.insert(nav_at, "nav")
    ncx = ('<?xml version="1.0" encoding="utf-8"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" '
           f'version="2005-1"><head><meta name="dtb:uid" content="{uid}"/></head>'
           f"<docTitle><text>{esc(book.title)}</text></docTitle><navMap>" +
           "".join(f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{esc(t)}</text></navLabel>'
                   f'<content src="{fn}"/></navPoint>' for i, (fn, t) in enumerate(navitems, 1)) +
           "</navMap></ncx>")
    files["OEBPS/toc.ncx"] = ncx
    manifest.append('<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>')
    files["OEBPS/style.css"] = CSS + title_css(s)
    manifest.append('<item id="css" href="style.css" media-type="text/css"/>')
    mod = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    opf = ('<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" '
           'version="3.0" unique-identifier="bookid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           f'<dc:identifier id="bookid">{uid}</dc:identifier><dc:title>{esc(book.title)}</dc:title>'
           f'<dc:language>{esc(lang)}</dc:language>' +
           (f"<dc:creator>{esc(book.author)}</dc:creator>" if book.author else "") +
           f'{cover_meta}<meta property="dcterms:modified">{mod}</meta></metadata><manifest>' + "".join(manifest) +
           '</manifest><spine toc="ncx">' + "".join(f'<itemref idref="{x}"/>' for x in spine) +
           "</spine></package>")
    files["OEBPS/content.opf"] = opf
    with zipfile.ZipFile(out_path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   "</rootfiles></container>", compress_type=zipfile.ZIP_DEFLATED)
        for name, data in files.items():
            z.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


# --------------------------------------------------------------------------
# Word DOCX
# --------------------------------------------------------------------------
WORD_FONT = {"Times (built-in)": "Times New Roman", "Helvetica (built-in)": "Arial",
             "Courier (built-in)": "Courier New"}


def add_docx_picture(d, b, s, pw, ph):
    """A picture in the Word file, with the size, alignment and placement chosen in the editor.
    Word can't print to the paper's edge, so a full-page picture fills the text area of its own page."""
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from bf_pictures import size
    tw = pw - float(s["margin_inner"]) - float(s["margin_outer"])
    th = ph - float(s["margin_top"]) - float(s["margin_bottom"])
    own = b.img_place in ("page", "full")
    if b.img_place == "full":
        w = tw
    elif b.img_width:
        w = tw * min(100, max(5, int(b.img_width))) / 100
    else:
        w = min(tw, 5)
    px = size(b.image)
    if px and px[0] and w * px[1] / px[0] > th * (0.95 if own else 0.9):
        w = th * (0.95 if own else 0.9) * px[0] / px[1]          # keep it on one page, same shape
    if own and d.paragraphs and d.paragraphs[-1].text.strip():
        d.add_page_break()
    d.add_picture(io.BytesIO(b.image), width=Inches(w))
    d.paragraphs[-1].alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(
        b.align, WD_ALIGN_PARAGRAPH.CENTER)
    if own:
        d.add_page_break()


def build_docx(book, s, out_path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches, RGBColor

    d = docx.Document()
    pw, ph = page_size_inches(s)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(pw), Inches(ph)
    sec.top_margin, sec.bottom_margin = Inches(float(s["margin_top"])), Inches(float(s["margin_bottom"]))
    sec.left_margin, sec.right_margin = Inches(float(s["margin_inner"])), Inches(float(s["margin_outer"]))
    st_el = d.settings.element
    st_el.append(OxmlElement("w:mirrorMargins"))
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    st_el.append(uf)

    font = WORD_FONT.get(s["font_family"], s["font_family"])
    hfam = s.get("heading_font") or "Same as body text"
    hfont = font if hfam == "Same as body text" else WORD_FONT.get(hfam, hfam)
    size = float(s["body_size"])
    normal = d.styles["Normal"]
    normal.font.name = font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.font.size = Pt(size)
    pf = normal.paragraph_format
    pf.first_line_indent = Inches(float(s["indent"]))
    pf.space_after = Pt(float(s["para_space"]))
    pf.line_spacing = float(s["line_spacing"])
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if s["align"] == "Justified" else WD_ALIGN_PARAGRAPH.LEFT
    pf.widow_control = True
    # Titles (Layout tab, "Titles: alignment & position"; the same checked values as the PDF): Word has real
    # pages of the book's size, so positions are the PDF's - space above the title, as a % of the text area.
    tl = title_look(s)
    AL = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    fh_pt = (ph - float(s["margin_top"]) - float(s["margin_bottom"])) * 72
    ts = float(s["title_size"])
    has_parts = any(c.level == 0 for c in book.active())
    for hn in ("Heading 1", "Heading 2"):
        # Heading 1 = part titles when the book has parts (Heading 2 = chapters), otherwise chapter titles
        part = has_parts and hn == "Heading 1"
        h = d.styles[hn]
        h.font.name, h.font.bold = hfont, bool(s.get("heading_bold", True))
        h.font.size = Pt(tl["part_size"] or ts * 1.2) if part else Pt(ts)
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.alignment = AL[tl["part_align"] if part else tl["ch_align"]]
        h.paragraph_format.space_after = Pt(16)
        h.paragraph_format.first_line_indent = Inches(0)

    def top_space(pos_pt, block_pt):
        """Space above a title block that starts pos_pt down the text area, never so much that the block
        (block_pt high, estimated) is pushed off its page (as in the PDF)."""
        room = fh_pt - block_pt - 2
        return pos_pt if pos_pt <= room else max(0.0, room)

    def line_pt(pt):                                   # a line's height in Word (font height x line spacing)
        return pt * 1.15 * float(s["line_spacing"])

    def center(text, sz=None, bold=False, italic=False, before=0, after=6, brk=False, align="center"):
        p = d.add_paragraph()
        p.alignment = AL.get(align, WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.page_break_before = brk
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if sz:
            r.font.size = Pt(sz)
        return p

    def fmt_para(p, b):
        """Alignment / indent / space above set in the editor."""
        al = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
              "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
        if getattr(b, "align", "") in al:
            p.alignment = al[b.align]
        if getattr(b, "indent", 0):
            p.paragraph_format.left_indent = Inches(0.25 * int(b.indent))
        if getattr(b, "space", 0):
            p.paragraph_format.space_before = Pt(int(b.space))

    def add_markup(p, markup, base_italic=False):
        fl = {"i": 0, "b": 0, "super": 0, "sub": 0, "strike": 0, "u": 0}
        fonts = []                                    # editor fonts: (face, size, colour) stack
        for tok in re.split(r"(<[^>]+>)", markup):
            if not tok:
                continue
            if tok.startswith("<"):
                name = re.sub(r"[</>]", " ", tok).split()
                name = name[0].lower() if name else ""
                if name == "br":
                    p.add_run().add_break()
                elif name in fl:
                    fl[name] += -1 if tok.startswith("</") else 1
                elif name == "font":
                    if tok.startswith("</"):
                        if fonts:
                            fonts.pop()
                    else:
                        fonts.append(dict(re.findall(r'(\w+)="([^"]*)"', tok)))
                continue
            r = p.add_run(html.unescape(tok))
            for f in fonts:
                if f.get("face"):
                    r.font.name = html.unescape(f["face"]).replace(" (built-in)", "")
                if f.get("size"):
                    try:
                        v = f["size"]
                        r.font.size = Pt(float(s["body_size"]) * float(v[:-1]) / 100 if v.endswith("%")
                                         else float(v))
                    except ValueError:
                        pass
                if f.get("color", "").startswith("#") and len(f["color"]) == 7:
                    r.font.color.rgb = RGBColor.from_string(f["color"][1:].upper())
            r.italic = bool(fl["i"]) != base_italic
            r.bold = bool(fl["b"])
            if fl["strike"]:
                r.font.strike = True
            if fl["u"]:
                r.underline = True
            if fl["super"]:
                r.font.superscript = True
            if fl["sub"]:
                r.font.subscript = True

    bt = tl["tp_title_size"] or max(30, ts * 1.5)          # book title size, as in the PDF
    TITLE_LINES = ("btitle", "bsub", "bauth", "bcred")       # (the title page's alignment)
    MATTER_DOCX = {"btitle": (bt, True, False, 0, 12), "bsub": (float(s["sub_size"]) + 1, False, True, 0, 6),
                   "bauth": (float(s["sub_size"]), False, False, 30, 6), "bcred": (size * 0.9, False, False, 0, 6),
                   "plate": (float(s["sub_size"]), False, True, 0, 6), "ded": (None, False, True, 0, 8),
                   "endline": (None, False, False, 24, 6)}

    def add_blocks(blocks):
        """Paragraphs for blocks: the chapter styles and the front/back page ones (title page lines...)."""
        prev_label = None
        for b in blocks:
            if b.kind != "anote":
                prev_label = None
            if b.kind == "anote":
                if an_mode(s).startswith("Leave"):
                    continue
                keep = an_mode(s).startswith("Keep")
                if b.label and b.label != prev_label:
                    p = d.add_paragraph()
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.left_indent = Inches(0 if keep else 0.35)
                    add_markup(p, esc(b.label) if keep else f"<b>{esc(b.label)}</b>")
                p = d.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.35 if not keep else 0.4)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_markup(p, b.text)
                fmt_para(p, b)
                if not keep:
                    for r in p.runs:
                        r.font.size = Pt(size * 0.88)
                prev_label = b.label
            elif b.kind == "item":
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3 * (b.depth + 1))
                p.paragraph_format.first_line_indent = Inches(-0.25 if b.bullet else 0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_markup(p, (esc(b.bullet) + "\t" if b.bullet else "") + b.text)
                fmt_para(p, b)
            elif b.kind == "para":
                p = d.add_paragraph()
                add_markup(p, b.text)
                fmt_para(p, b)
            elif b.kind == "subtitle":                   # (chapter subtitles: the chapter title alignment)
                p = center("", after=14, align=tl["ch_align"])
                add_markup(p, b.text, base_italic=True)
                fmt_para(p, b)
            elif b.kind == "quote":
                p = d.add_paragraph()
                p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.4)
                p.paragraph_format.first_line_indent = Inches(0)
                add_markup(p, b.text)
                fmt_para(p, b)
            elif b.kind == "verse":
                p = d.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(0)
                if s["verse_style"] == "Centered italic":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_markup(p, "<br/>".join(b.lines), base_italic=True)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.left_indent = Inches(0.5)
                    add_markup(p, "<br/>".join(b.lines))
            elif b.kind == "head":
                p = center("", bold=True, before=12, after=8)
                add_markup(p, f"<b>{b.text}</b>")
                fmt_para(p, b)
            elif b.kind == "scene":
                center(scene_text(b, s))
            elif b.kind == "image" and b.image:
                try:
                    add_docx_picture(d, b, s, pw, ph)
                except Exception:
                    pass
            elif b.kind in MATTER_DOCX:
                sz, bold, italic, before, after = MATTER_DOCX[b.kind]
                p = center("", sz, bold=bold, italic=italic, before=before, after=after,
                           align=tl["tp_align"] if b.kind in TITLE_LINES else "center")
                add_markup(p, f"<b>{b.text}</b>" if bold else (b.text or ""), base_italic=italic)
                if sz:
                    for r in p.runs:
                        if r.font.size is None:
                            r.font.size = Pt(sz)
                fmt_para(p, b)
            elif b.kind in ("flush", "meta", "tocentry"):
                p = d.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(-0.25 if b.kind == "meta" else 0)
                p.paragraph_format.left_indent = Inches(0.25 if b.kind == "meta" else 0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_markup(p, b.text or "")
                fmt_para(p, b)

    def page_of(blocks, before=0, brk=True):
        """blocks on a new page (the first paragraph starts the page, before points down)."""
        start = len(d.paragraphs)
        add_blocks(blocks)
        if len(d.paragraphs) > start:
            pf = d.paragraphs[start].paragraph_format
            pf.page_break_before = brk or pf.page_break_before
            if before and not pf.space_before:
                pf.space_before = Pt(before)

    art = front_art(book, s) if s.get("cover_art_in_book") else None
    if art:
        from bf_core import Block
        try:
            add_docx_picture(d, Block("image", image=art, img_place="full"), s, pw, ph)
        except Exception:
            pass
    plate_ed = edited(book, s, "bookplate")

    def plate(before, brk):
        if plate_ed is not None:
            page_of(plate_ed, before=before, brk=brk)
        else:
            center(s["bookplate_text"], float(s["sub_size"]), italic=True, before=before, brk=brk)

    if s["title_page"]:
        # The title block starts tp_pos % down the text area, with the title page alignment (the PDF's place;
        # the bookplate stays centred, as in the PDF)
        title_ed = edited(book, s, "title")
        ss = float(s["sub_size"])
        est = sum(line_pt(MATTER_DOCX.get(b.kind, (size,))[0] or size) + 8 for b in (title_ed or [])
                  if b.kind != "image") if title_ed is not None else (
            line_pt(bt) + 12 + (line_pt(ss + 1) + 6 if book.subtitle else 0) +
            (30 + line_pt(ss) + 6 if book.author else 0) + (line_pt(size * 0.9) + 6 if book.credit else 0))
        if s["bookplate"] == "On title page":
            est += 60 + line_pt(ss) + 6 + line_pt(size)
        top = top_space(fh_pt * tl["tp_pos"] / 100, est)
        tpa = tl["tp_align"]
        if title_ed is not None:                       # edited in Specific Text Edit
            page_of(title_ed, before=top, brk=False)
        else:
            center(book.title, bt, bold=True, before=top, after=12, align=tpa)
            if book.subtitle:
                center(book.subtitle, ss + 1, italic=True, align=tpa)
            if book.author:
                center(book.author, ss, before=30, align=tpa)
            if book.credit:
                center(book.credit, size * 0.9, align=tpa)
        if s["bookplate"] == "On title page":
            plate(60, False)
            center("_" * 40, after=0)
    if s["bookplate"] == "Separate page":
        plate(220, True)
        center("_" * 40)
    ded_ed = edited(book, s, "dedication")
    if ded_ed is not None:
        if bf_matter.has_text(ded_ed):
            page_of(ded_ed, before=180)
    elif (s.get("dedication") or "").strip():
        center(s["dedication"].strip(), italic=True, before=180, brk=True)
    def field(p, instr, result=""):
        """A Word field (TOC, PAGEREF...): Word works out its result when the fields are updated."""
        r = p.add_run()
        for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", result), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
                r._r.append(el)
            elif kind == "text":
                if txt:
                    p.add_run(txt)
                r = p.add_run()
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                r._r.append(el)

    def mark(uid):
        return "bf_" + re.sub(r"\W", "", str(uid))[:30]

    def contents_page():
        center(s["toc_title"], ts, bold=True, after=18, brk=True, align=tl["ch_align"])
        if not getattr(book, "toc", None):           # automatic: Word builds it from the chapter headings
            p = d.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0)
            field(p, 'TOC \\o "1-2" \\h \\z \\u', "Press F9 (or right-click > Update Field) to build the contents.")
            return
        # Changed in Edit contents: those entries, in that order, with their words and formatting; a typed page
        # number as typed, otherwise the chapter's real page in this document (a PAGEREF field Word fills in).
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        from bf_pdf import contents_entries, apply_contents_edits
        width = Inches(pw - float(s["margin_inner"]) - float(s["margin_outer"]))
        for e in apply_contents_edits(contents_entries(book, s), book.toc):
            p = d.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Inches(0)
            pf.left_indent = Inches(0.3 * int(e.get("level") or 0))
            pf.space_after = Pt(4)
            if e.get("space"):
                pf.space_before = Pt(int(e["space"]))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT,
                                      WD_TAB_LEADER.DOTS if s.get("toc_dots") else WD_TAB_LEADER.SPACES)
            add_markup(p, e.get("markup") or esc(e["text"]))
            if (e.get("page") or "").strip():
                p.add_run("\t" + e["page"].strip())
            elif e.get("uid") is not None:
                p.add_run("\t")
                field(p, f" PAGEREF {mark(e['uid'])} \\h ")

    fic = fic_sections(book, s)                       # (the PDF's order: work info..., note, contents, text
    for heading, items, blocks in fic:                #  before chapter 1, chapters)
        center(heading, float(s["sub_size"]), bold=True, brk=True, after=12, align=tl["ch_align"])
        if blocks is not None:
            add_blocks(blocks)
            continue
        for cls, m in items:
            p = d.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(-0.25 if cls == "meta" else 0)
            p.paragraph_format.left_indent = Inches(0.25 if cls == "meta" else 0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_markup(p, m)
    note_ed = edited(book, s, "note")

    def note_page():
        if note_ed is not None:
            if bf_matter.has_text(note_ed) or (s.get("note_title") or "").strip():
                if (s.get("note_title") or "").strip():
                    center(s["note_title"], float(s["sub_size"]), bold=True, brk=True, after=12,
                           align=tl["ch_align"])
                page_of(note_ed, brk=not (s.get("note_title") or "").strip())
        elif (s.get("note_text") or "").strip():
            first = True
            if s.get("note_title"):
                center(s["note_title"], float(s["sub_size"]), bold=True, brk=True, after=12, align=tl["ch_align"])
                first = False
            for para in re.split(r"\n\s*\n", s["note_text"].strip()):
                p = d.add_paragraph(" ".join(para.split()))
                if first:
                    p.paragraph_format.page_break_before = True
                    first = False

    if s.get("note_place") == "Before contents":
        note_page()
    if s["toc"]:
        contents_page()
    if s.get("note_place") != "Before contents":
        note_page()
    front = front_text_blocks(book, s)
    if front:
        page_of(front)

    heads = list(chapter_heads(book, s))
    lead_pt = line_pt(size)
    for i, (ch, lbl, title) in enumerate(heads):
        hstyle = "Heading 1" if (ch.level == 0 or not has_parts) else "Heading 2"
        h = d.add_paragraph(style=hstyle)
        h.paragraph_format.page_break_before = True
        block = (line_pt(float(s["sub_size"])) if lbl else 0) + (line_pt(ts * (1.2 if ch.level == 0 else 1)) + 16
                                                                 if title else 0)
        if ch.level == 0:                              # a part: its title starts part_pos % down the page
            h.paragraph_format.space_before = Pt(top_space(fh_pt * tl["part_pos"] / 100, block))
        elif tl["sink"] and (lbl or title):            # "Drop chapter start down" (leaving room for 2 lines):
            own = d.styles[hstyle].paragraph_format.space_before      # the title's own space above, plus the drop
            h.paragraph_format.space_before = Pt(top_space(tl["sink"] * 72 + (own.pt if own is not None else 0),
                                                           block + 2 * lead_pt))
        linked = bool(s["toc"] and getattr(book, "toc", None))     # an edited contents' page numbers point here
        if linked:
            bm_start = OxmlElement("w:bookmarkStart")
            bm_start.set(qn("w:id"), str(i + 1))
            bm_start.set(qn("w:name"), mark(ch.uid))
            h._p.append(bm_start)
        if lbl:
            rl = h.add_run(lbl)
            rl.bold = False
            rl.font.size = Pt(float(s["sub_size"]))
            if title:
                rl.add_break()
        if title:
            h.add_run(title)
        if linked:
            bm_end = OxmlElement("w:bookmarkEnd")
            bm_end.set(qn("w:id"), str(i + 1))
            h._p.append(bm_end)
        add_blocks(ch.blocks)
        if i == len(heads) - 1 and s.get("the_end"):
            end_ed = edited(book, s, "the_end")
            if end_ed is not None:
                add_blocks(end_ed)
            else:
                center("THE END", before=24)
    xm = edited(book, s, "end_matter")                 # end pages (afterword, notes...)
    xhead = bf_matter.section_heading(book, s, "end_matter")
    if xm is not None and (bf_matter.has_text(xm) or xhead):
        if xhead:
            center(xhead, float(s["sub_size"]), bold=True, brk=True, after=12, align=tl["ch_align"])
        page_of(xm, brk=not xhead)

    if s["pn_position"] != "None":
        fp = sec.footer.paragraphs[0] if not s["pn_position"].startswith("Top") else sec.header.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        for kind in ("begin", "instr", "end"):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = "PAGE"
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
    d.core_properties.title = book.title
    d.core_properties.author = book.author
    d.save(out_path)


# --------------------------------------------------------------------------
# HTML and TXT
# --------------------------------------------------------------------------
def build_html(book, s, out_path):
    import base64
    heads = list(chapter_heads(book, s))
    imgmap = {}
    for ch in book.active():
        for b in ch.blocks:
            if b.kind == "image" and b.image:
                ext, mt = img_type(b.image)
                if ext:
                    imgmap[id(b)] = f"data:{mt};base64," + base64.b64encode(b.image).decode()
    for b in matter_pictures(book, s):
        ext, mt = img_type(b.image)
        if ext:
            imgmap[id(b)] = f"data:{mt};base64," + base64.b64encode(b.image).decode()
    out = [f'<!DOCTYPE html><html lang="{book.language or "en"}"><head><meta charset="utf-8">'
           f"<title>{esc(book.title)}</title><style>{CSS}{title_css(s)}"
           " body{max-width:40em;margin:0 auto;padding:2em}"
           "</style></head><body>"]
    art = front_art(book, s) if s.get("cover_art_in_book") else None
    if art and img_type(art)[0]:
        out.append(f'<p class="img"><img src="data:{img_type(art)[1]};base64,{base64.b64encode(art).decode()}" '
                   f'alt="{esc(book.title)}" style="max-width:100%" /></p>')
    title_ed = edited(book, s, "title")
    if title_ed is not None:
        out += ['<div class="titlepage">', blocks_html(title_ed, imgmap, s), "</div>"]
    else:
        out += ['<div class="titlepage">', f'<p class="t">{esc(book.title)}</p>']
        for v, tag in ((book.subtitle, "i"), (book.author, "span"), (book.credit, "span")):
            if v:
                out.append(f"<p><{tag}>{esc(v)}</{tag}></p>")
        out.append("</div>")
    ded = bf_matter.section_blocks(book, s, "dedication")      # edited, or made from the settings
    if bf_matter.has_text(ded):
        out.append('<div class="ded">' + blocks_html(ded, imgmap, s) + "</div>")
    note_ed = bf_matter.section_blocks(book, s, "note")

    def note():
        head = (s.get("note_title") or "").strip()
        if head or bf_matter.has_text(note_ed):
            out.append((f"<h2>{esc(head)}</h2>" if head else "") + blocks_html(note_ed, imgmap, s))

    fic = fic_sections(book, s)                 # (the PDF's order: work info..., the note, contents)
    if fic:
        out.append(fic_html(fic, s))
    if s.get("note_place") == "Before contents":
        note()
    if s["toc"]:
        out.append(f"<h2>{esc(s['toc_title'])}</h2><ol>")
        for i, disp in contents_list(book, s, heads):
            out.append(f'<li><a href="#c{i}">{esc(disp)}</a></li>' if i else f"<li>{esc(disp)}</li>")
        out.append("</ol>")
    if s.get("note_place") != "Before contents":
        note()
    front = front_text_blocks(book, s)
    if front:
        out.append('<div class="front">' + blocks_html(front, imgmap, s) + "</div>")
    for i, (ch, lbl, title) in enumerate(heads, 1):
        out.append(f'<hr id="c{i}">')
        if lbl or title:                               # part title block / chapter title block, as in the EPUB
            out.append('<div class="part">' if ch.level == 0 else '<div class="chead">')
            if lbl:
                out.append(f'<p class="label">{esc(lbl)}</p>')
            if title:
                out.append(f"<h2>{esc(title)}</h2>")
            out.append("</div>")
        out.append(blocks_html(ch.blocks, imgmap, s))
    if s.get("the_end"):
        end_ed = edited(book, s, "the_end")
        out.append(blocks_html(end_ed, imgmap, s) if end_ed is not None else '<p class="end">THE END</p>')
    xm = edited(book, s, "end_matter")
    xhead = bf_matter.section_heading(book, s, "end_matter")
    if xm is not None and (bf_matter.has_text(xm) or xhead):
        out.append("<hr>" + (f"<h2>{esc(xhead)}</h2>" if xhead else "") + blocks_html(xm, imgmap, s))
    out.append("</body></html>")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))


def build_txt(book, s, out_path):
    title_ed = edited(book, s, "title")
    if title_ed is not None:
        out = [plain(b.text or "") for b in title_ed if b.kind != "image" and plain(b.text or "").strip()]
    else:
        out = [book.title]
        for v in (book.subtitle, book.author, book.credit):
            if v:
                out.append(v)
    out.append("")
    ded = bf_matter.section_blocks(book, s, "dedication")      # edited, or made from the settings
    if bf_matter.has_text(ded):
        out += [""] + blocks_text(ded, s)
    note_ed = bf_matter.section_blocks(book, s, "note")

    def note():
        head = (s.get("note_title") or "").strip()
        if head or bf_matter.has_text(note_ed):
            out.extend(([head, ""] if head else []) + blocks_text(note_ed, s))

    heads = list(chapter_heads(book, s))
    for heading, items, blocks in fic_sections(book, s):       # (the PDF's order: work info..., note, contents)
        out += [heading] + (blocks_text(blocks, s) if blocks is not None else [plain(m) for _c, m in items] + [""])
    if s.get("note_place") == "Before contents":
        note()
    if s["toc"]:
        out.append(s["toc_title"])
        for _i, disp in contents_list(book, s, heads, sep=" - "):
            out.append("  " + disp)
        out.append("")
    if s.get("note_place") != "Before contents":
        note()
    front = front_text_blocks(book, s)
    if front:
        out += [""] + blocks_text(front, s)
    for ch, lbl, title in heads:
        out += ["", ""]
        if lbl:
            out.append(lbl)
        if title:
            out.append(title)
        out.append("")
        for b in ch.blocks:
            if b.kind in ("para", "quote", "head", "subtitle"):
                out.append(plain(b.text))
                out.append("")
            elif b.kind == "anote" and not an_mode(s).startswith("Leave"):
                out += [("[" + b.label + "] " if b.label else "") + plain(b.text), ""]
            elif b.kind == "item":
                out.append("  " * b.depth + (b.bullet + " " if b.bullet else "  ") + plain(b.text))
            elif b.kind == "verse":
                out += [plain(l) for l in b.lines] + [""]
            elif b.kind == "scene":
                out += [scene_text(b, s), ""]
            elif b.kind in MATTER_CSS:
                out += blocks_text([b], s)
    if s.get("the_end"):
        end_ed = edited(book, s, "the_end")
        out += blocks_text(end_ed, s) if end_ed is not None else ["THE END"]
    xm = edited(book, s, "end_matter")
    xhead = bf_matter.section_heading(book, s, "end_matter")
    if xm is not None and (bf_matter.has_text(xm) or xhead):
        out += ["", ""] + ([xhead, ""] if xhead else []) + blocks_text(xm, s)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
