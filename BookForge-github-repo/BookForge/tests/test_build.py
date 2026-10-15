"""Full-build regression tests (plain Python - run:  python tests/test_build.py).

Uses the books in tests/inputs/ (see test_chapters.py); missing ones are skipped.
"""
import os
import re
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
INPUTS = os.environ.get("BOOKFORGE_TEST_INPUTS", os.path.join(HERE, "inputs"))

import pdfplumber                                        # noqa: E402
from pypdf import PdfReader                              # noqa: E402
import bf_impose                                         # noqa: E402,F401
from bf_core import DEFAULTS, load_book, apply_info, plain, parse_text      # noqa: E402
from bf_build import run_build                           # noqa: E402

FAILS = []
OUT = tempfile.mkdtemp(prefix="bf_build_")


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


def build(name, **kw):
    s = dict(DEFAULTS)
    s.update(source_path=os.path.join(INPUTS, name), out_dir=os.path.join(OUT, kw.pop("tag", "x")),
             fmt_epub=True, fmt_docx=True, fmt_html=True, fmt_txt=True, fmt_cover=True)
    s.update(kw)
    log = []
    book = apply_info(load_book(s), s)
    folder, made, notes = run_build(book, s, log.append)
    return book, folder, made, log, s


def pdf_text_pages(path):
    with pdfplumber.open(path) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def toc_and_outline_agree(pdf_path, book):
    """Every chapter's bookmark points at the page that really starts with that chapter's title."""
    r = PdfReader(pdf_path)
    pages = pdf_text_pages(pdf_path)
    bad = []
    outline = [o for o in r.outline if not isinstance(o, list)]
    titles = {o.title: r.get_destination_page_number(o) for o in outline}
    def norm(x):
        return re.sub(r"[^a-z0-9]", "", plain(x).lower())
    for ch in book.active():
        want = norm(ch.title)[:24]
        key = next((t for t in titles if want and want in norm(t)), None)
        if key is None:
            bad.append(("no bookmark", ch.title[:30]))
            continue
        if want[:12] not in norm(pages[titles[key]]):
            bad.append(("bookmark page lacks title", ch.title[:30]))
    return bad


# ---------------------------------------------------------------- classic: Grimms, every format
if os.path.exists(os.path.join(INPUTS, "pg2591-images-3.epub")):
    book, folder, made, log, s = build("pg2591-images-3.epub", tag="grimms")
    exts = sorted(os.path.splitext(m)[1] for m in made)
    check(exts.count(".pdf") == 3 and all(e in exts for e in (".epub", ".docx", ".html", ".txt")),
          f"Grimms: all formats made ({len(made)} files)")
    pdf = [m for m in made if m.endswith("Print.pdf")][0]
    bad = toc_and_outline_agree(pdf, book)
    check(not bad, f"Grimms: every chapter bookmark lands on its chapter's page ({len(book.active())} chapters) {bad[:2]}")
    pages = pdf_text_pages(pdf)
    toc = "\n".join(pages[1:5])
    nums = re.findall(r"THE GOLDEN BIRD\s*\.*\s*(\d+)|The Golden Bird\s*\.*\s*(\d+)", toc)
    gb = next((int(a or b) for a, b in nums), None)
    real = next((i + 1 for i, t in enumerate(pages) if t.strip().startswith("THE GOLDEN BIRD") or
                 "THE GOLDEN BIRD\n" in t[:60]), None)
    check(gb is not None and real is not None and gb == real,
          f"Grimms: contents says The Golden Bird is on page {gb}, it is on page {real}")
    import docx
    d = docx.Document([m for m in made if m.endswith(".docx")][0])
    check(len(d.paragraphs) > 500, f"Grimms: Word file opens ({len(d.paragraphs)} paragraphs)")
    ep = [m for m in made if m.endswith(".epub")][0]
    s2 = dict(DEFAULTS)
    s2["source_path"] = ep
    check(len(load_book(s2).chapters) == 62, "Grimms: our EPUB re-reads as 62 chapters")
    check(zipfile.ZipFile(ep).namelist()[0] == "mimetype", "Grimms: EPUB starts with mimetype entry")
    cov = PdfReader([m for m in made if m.endswith("Cover.pdf")][0]).pages[0]
    trim_w = float(cov.mediabox.width) / 72
    n = len(PdfReader(pdf).pages)
    spine = (n + 1) // 2 * 0.004
    check(abs(trim_w - (2 * 5.5 + spine)) < 0.05,
          f"Cover width {trim_w:.3f} in = 2 x 5.50 trim + spine {spine:.3f} in")

# ---------------------------------------------------------------- AO3: notes / work info / summary settings
EOTL = "End_of_the_Line.epub"
if os.path.exists(os.path.join(INPUTS, EOTL)):
    probe_note = "cookie clicker"           # text of chapter 1's end note
    for mode, expect in (("Set apart (smaller, indented)", True), ("Keep as written", True), ("Leave out", False)):
        book, folder, made, log, s = build(EOTL, tag="eotl_" + mode.split()[0], an_mode=mode,
                                           fic_info="Add a work info page", fmt_impose=False,
                                           fmt_epub=False, fmt_docx=False, fmt_html=False, fmt_txt=True,
                                           fmt_cover=False)
        pdf = [m for m in made if m.endswith("Print.pdf")][0]
        text = "\n".join(pdf_text_pages(pdf))
        check((probe_note in text) == expect, f"EotL, author's notes = {mode}: note text present = {expect}")
        txt = open([m for m in made if m.endswith(".txt")][0], encoding="utf-8").read()
        check((probe_note in txt) == expect, f"EotL TXT export follows author's notes = {mode}")
        if mode.startswith("Set"):
            check("Work Info" in text and "Rating:" in text, "EotL: work info page present when chosen")
            check("Derek lost." in text, "EotL: summary page present")
            check("Please drop by the Archive" not in text, "EotL: AO3 'Please drop by' line removed")
            bad = toc_and_outline_agree(pdf, book)
            check(not bad, f"EotL: bookmarks match chapter pages ({len(book.active())} chapters) {bad[:2]}")
    book, folder, made, log, s = build(EOTL, tag="eotl_noinfo", fic_info="Leave out", fic_summary="Leave out",
                                       fmt_impose=False, fmt_epub=False, fmt_docx=False, fmt_html=False,
                                       fmt_txt=False, fmt_cover=False)
    text = "\n".join(pdf_text_pages([m for m in made if m.endswith("Print.pdf")][0]))
    check("Work Info" not in text and "Derek lost." not in text, "EotL: work info and summary left out on request")

# ---------------------------------------------------------------- included fonts in a real book
ULMO = "Of_the_Tales_of_Ulmos.html"
if os.path.exists(os.path.join(INPUTS, ULMO)):
    book, folder, made, log, s = build(ULMO, tag="ulmo_fonts", font_family="EB Garamond",
                                       heading_font="Cinzel", dropcap_font="Playfair Display", drop_cap=True,
                                       fmt_epub=False, fmt_docx=False, fmt_html=False, fmt_txt=False)
    pdf = [m for m in made if m.endswith("Print.pdf")][0]
    names = set()
    for p in PdfReader(pdf).pages:
        for v in (p.get("/Resources", {}).get("/Font", {}) or {}).values():
            names.add(str(v.get_object()["/BaseFont"]).split("+")[-1])
    check({"EBGaramond-Regular", "Cinzel-Bold", "PlayfairDisplay-Regular"} <= names,
          f"Ulmo: chosen included fonts are embedded {sorted(names)}")
    check(not any("WARNING" in l for l in log), f"Ulmo: no font warnings ({[l for l in log if 'WARN' in l]})")
    check(len(book.active()) == 22 and any(m.endswith("Imposed for binding.pdf") for m in made),
          "Ulmo: 22 chapters, imposed PDF made")

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
