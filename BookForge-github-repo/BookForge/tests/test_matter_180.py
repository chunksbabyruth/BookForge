"""Checks for editable front & back matter and contents in 1.8.0 (plain Python - run: python tests/test_matter_180.py).

The pages around the chapters (title page, bookplate, dedication, extra front page, work info, summary, author's
notes, contents, text before chapter 1, THE END, end pages, Notes pages, back-cover blurb) can be edited in
Specific Text Edit like chapters. These checks cover:

1. Every page: edited words and formatting (fonts, sizes, bold, italic, underline, colour, alignment, space)
   reach the Print PDF, the cover and the EPUB / Word / web page / text exports; a page saved unchanged prints
   exactly as before.
2. Contents edited as text: reorder, hide, add, rename and format entries; chapters, their order and their ids
   are never changed.
3. Typed contents page numbers: printed beside the entry, the chapter stays on its real page; clearing goes
   back to the real number.
4. Fanfic work info, summary and work-level notes survive project save/reopen.
5. Edits survive light/dark, tab switching, previews and settings changes (window checks).
6. Pictures and rich formatting on the pages that take pictures.
7. Preview = final Print PDF (and the imposed sheets) with edited pages.
8. Older projects (no edited pages stored) open and print as before.
9. Blank pages, right-hand chapter starts, page numbering and binding are unaffected.
Window checks run only where a display is available. Nothing needs the internet.
"""
import copy
import dataclasses
import io
import json
import os
import re
import sys
import tempfile
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="bf_matter_")
os.environ["APPDATA"] = os.path.join(TMP, "appdata")
os.environ["HOME"] = os.path.join(TMP, "home")
os.environ["BOOKFORGE_NO_UPDATE_CHECK"] = "1"
os.makedirs(os.environ["APPDATA"], exist_ok=True)
os.makedirs(os.path.join(os.environ["HOME"], "Documents"), exist_ok=True)

import pdfplumber                                                     # noqa: E402
from PIL import Image                                                 # noqa: E402

import bf_edit                                                        # noqa: E402
import bf_export                                                      # noqa: E402
import bf_matter                                                      # noqa: E402
from bf_build import build_preview, run_build                         # noqa: E402
from bf_core import DEFAULTS, Book, Chapter, Block, apply_info, page_size_inches  # noqa: E402
from bf_pdf import build_pdf, contents_entries, contents_rows, rows_to_edits  # noqa: E402

FAILS = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg, flush=True)
    if not cond:
        FAILS.append(msg)


def settings(**kw):
    s = copy.deepcopy(DEFAULTS)
    s.update(open_when_done=False, out_dir=tempfile.mkdtemp(dir=TMP), source_mode="paste")
    s.update(kw)
    return s


def png(w, h, color):
    b = io.BytesIO()
    Image.new("RGB", (w, h), color).save(b, "PNG")
    return b.getvalue()


RED_PNG, BLUE_PNG = png(120, 80, (220, 30, 30)), png(80, 120, (30, 30, 220))


def fic_book():
    """A book with every kind of front matter the source can give: fanfic work info, summary, notes, and text
    before the first chapter."""
    return Book(title="Matter Title", subtitle="Matter Subtitle", author="Matter Author", credit="Matter Credit",
                fic_meta=["<b>Rating:</b> General", "<b>Fandom:</b> Original Work"],
                fic_summary=[Block("para", text="The original <i>summary</i>.", label="Summary")],
                fic_notes=[Block("para", text="The original work notes.", label="Notes")],
                front_blocks=[Block("para", text="Original text before the first chapter.")],
                chapters=[Chapter("First Chapter", blocks=[Block("para", text="First chapter words. " * 60)]),
                          Chapter("Second Chapter", blocks=[Block("para", text="Second chapter words. " * 60)]),
                          Chapter("Third Chapter", blocks=[Block("para", text="Third chapter words. " * 60)])])


ALL_ON = dict(dedication="For the original reader", note_title="A Note", note_text="The original note.",
              the_end=True, fic_info="Add a work info page", bookplate="Separate page", notes_pages=2,
              blurb="The original blurb.")


def words(path):
    """[(page number, word dict)] with font, size and colour."""
    out = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            for w in pg.extract_words(extra_attrs=["fontname", "size", "non_stroking_color"]):
                out.append((pg.page_number, w))
    return out


def find(path, text):
    """The words of the first line containing text: [(page, word)] (empty when it isn't there)."""
    ws = words(path)
    first = text.split()[0]
    for i, (p, w) in enumerate(ws):
        if w["text"].startswith(first) or first in w["text"]:
            line = sorted([(pp, x) for pp, x in ws if pp == p and abs(x["top"] - w["top"]) < 1.5],
                          key=lambda e: e[1]["x0"])
            if text in " ".join(x["text"] for _p, x in line):
                return [(pp, x) for pp, x in line]
    return []


def page_texts(path):
    with pdfplumber.open(path) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def chars(path):
    with pdfplumber.open(path) as pdf:
        return [[(c["text"], round(c["x0"], 2), round(c["top"], 2), c["fontname"].split("+")[-1], round(c["size"], 2))
                 for c in p.chars] for p in pdf.pages]


def images_on(path):
    with pdfplumber.open(path) as pdf:
        return [len(p.images) for p in pdf.pages]


def font_has(ws, *bits):
    return bool(ws) and all(any(b in w["fontname"] for b in bits) for _p, w in ws)


def area(page_no, s):
    """(left, right) of the text area on a page, in points (odd pages are right-hand pages)."""
    pw = page_size_inches(s)[0] * 72
    inner, outer = float(s["margin_inner"]) * 72, float(s["margin_outer"]) * 72
    return (inner, pw - outer) if page_no % 2 else (outer, pw - inner)


def cpages(res):
    """The page each chapter starts on, from a preview result."""
    return [pg for _id, pg in res["chapter_pages"]]


def make(book, s, name):
    out = os.path.join(TMP, name + ".pdf")
    build_pdf(book, s, out, log=lambda m: None)
    return out


# ============================================================================ 1. every page, unchanged and edited
s1 = settings(**ALL_ON)
base_pdf = make(apply_info(fic_book(), s1), s1, "base")
base_chars = chars(base_pdf)
same = []
for key in bf_matter.SECTIONS:
    b = apply_info(fic_book(), s1)
    bf_matter.set_override(b, key, bf_matter.auto_blocks(b, s1, key))
    if chars(make(b, s1, "same_" + key)) == base_chars:
        same.append(key)
check(same == list(bf_matter.SECTIONS), f"each page saved unchanged prints exactly as before ({len(same)} of "
                                        f"{len(bf_matter.SECTIONS)}: {sorted(set(bf_matter.SECTIONS) - set(same))})")

book = apply_info(fic_book(), s1)
book.title = "Edited Book Title"            # the title page's words are the Book info fields
bf_matter.set_override(book, "title", [Block("btitle", text="<i>Matter Title</i>", align="left"),
                                       Block("bauth", text='<font color="#cc0000">Matter Author</font>'),
                                       Block("para", text="An extra <b>title page</b> line", align="right")])
bf_matter.set_override(book, "bookplate", [Block("plate", text="<b>This Book Belongs To</b>")])
bf_matter.set_override(book, "dedication", [Block("ded", text="For <b>Ann</b>, who read it first"),
                                            Block("para", text="and for Ben", align="right", space=24)])
bf_matter.set_override(book, "note", [Block("head", text="Edited Note Heading"),
                                      Block("para", text='Some <u>underlined</u> and <font color="#cc0000">red</font> '
                                                         'words in the note.'),
                                      Block("image", image=RED_PNG, image_name="red.png", img_width=40)])
bf_matter.set_override(book, "work_info", [Block("meta", text="<b>Fandom:</b> Edited Fandom"),
                                           Block("flush", text="An added info line")])
bf_matter.set_heading(book, "work_info", "About This Story")
bf_matter.set_override(book, "summary", [Block("flush", text="The edited <b>summary</b>.")])
bf_matter.set_override(book, "work_notes", [Block("flush", text="The edited work notes.")])
bf_matter.set_heading(book, "work_notes", "Notes From The Author")
bf_matter.set_override(book, "front_text", [Block("para", text='Edited text <font size="16">before</font> chapter one.')])
bf_matter.set_override(book, "the_end", [Block("endline", text="<b>FIN</b>")])
bf_matter.set_override(book, "end_matter", [Block("para", text="Thanks to everyone who helped with this book."),
                                            Block("image", image=BLUE_PNG, image_name="blue.png", img_width=30)])
bf_matter.set_heading(book, "end_matter", "Afterword")
bf_matter.set_override(book, "notes_page", [Block("flush", text="Write your thoughts here.")])
bf_matter.set_heading(book, "notes_page", "Jottings")
bf_matter.set_override(book, "blurb", [Block("para", text="An <b>edited</b> blurb for the back.", align="center")])
ed_pdf = make(book, s1, "edited")
txt = "\n".join(page_texts(ed_pdf))

tl = find(ed_pdf, "Edited Book Title")
check(tl and font_has(tl, "Italic", "Oblique") and abs(tl[0][1]["x0"] - area(tl[0][0], s1)[0]) < 2 and
      "Matter Title" not in txt,
      "title page: the words follow Book info (new title), the line keeps its italics and its own left alignment")
au = find(ed_pdf, "Matter Author")
check(au and all(tuple(round(c, 2) for c in w["non_stroking_color"]) == (0.8, 0.0, 0.0) for _p, w in au),
      "title page: a colour set on the author line prints")
ex = find(ed_pdf, "An extra title page line")
check(ex and ex[0][0] == tl[0][0] and abs(ex[-1][1]["x1"] - area(ex[0][0], s1)[1]) < 2,
      "title page: an added line, right-aligned, on the title page")
s1r = dict(s1, tp_align="Right", tp_pos=50.0)
b1r = apply_info(fic_book(), s1r)
bf_matter.set_override(b1r, "title", [Block("btitle", text="<i>Matter Title</i>"), Block("bauth", text="Matter Author")])
r_pdf = make(b1r, s1r, "title_right")
b1c = apply_info(fic_book(), s1)
bf_matter.set_override(b1c, "title", [Block("btitle", text="<i>Matter Title</i>"), Block("bauth", text="Matter Author")])
c_pdf = make(b1c, s1, "title_centre")
rt, ct = find(r_pdf, "Matter Title"), find(c_pdf, "Matter Title")
check(rt and ct and abs(rt[-1][1]["x1"] - area(rt[0][0], s1)[1]) < 2 and rt[0][1]["top"] > ct[0][1]["top"] + 100,
      "an edited title page still follows Layout's title page alignment (Right) and position (50 %)")
bp = find(ed_pdf, "This Book Belongs To")
check(bp and font_has(bp, "Bold"), "bookplate: its formatting (bold) prints")
dd = find(ed_pdf, "For Ann")
dd2 = find(ed_pdf, "and for Ben")
check(dd and font_has([w for w in dd if w[1]["text"].startswith("Ann")], "Bold") and dd2 and
      abs(dd2[-1][1]["x1"] - area(dd2[0][0], s1)[1]) < 2,
      "dedication: bold words and a right-aligned second line")
nt = find(ed_pdf, "Edited Note Heading")
nr = find(ed_pdf, "underlined and red words")
red = [w for _p, w in nr if w["text"].startswith("red")]
with pdfplumber.open(ed_pdf) as pdf:
    np_ = nr[0][0] if nr else 1
    pg = pdf.pages[np_ - 1]
    und = [w for _p, w in nr if w["text"].startswith("underlined")]
    has_line = bool(und) and any(abs(ln["top"] - und[0]["bottom"]) < 4 and ln["x0"] < und[0]["x1"]
                                 for ln in pg.lines + pg.rects)
check(nt and nr and red and tuple(round(c, 2) for c in red[0]["non_stroking_color"]) == (0.8, 0.0, 0.0) and has_line,
      "extra front page: heading, underline and colour print")
check(nr and images_on(ed_pdf)[nr[0][0] - 1] >= 1, "extra front page: the picture is on the page")
check("About This Story" in txt and "Edited Fandom" in txt and "An added info line" in txt and "Original Work" not in txt,
      "work info: edited lines and heading replace what the source had")
check("The edited summary." in txt and "The original summary" not in txt and "Notes From The Author" in txt and
      "The edited work notes." in txt and "original work notes" not in txt, "summary and author's notes: edited")
ft = find(ed_pdf, "Edited text")
check(ft and any(w["text"] == "before" and w["size"] > 15 for p, w in words(ed_pdf) if p == ft[0][0]) and
      "Original text before" not in txt,
      "text before chapter 1: edited, with a bigger size on one word")
fin = find(ed_pdf, "FIN")
check(fin and font_has(fin, "Bold") and "THE END" not in txt, "THE END: replaced by the edited closing line (bold)")
af = find(ed_pdf, "Afterword")
th_ = find(ed_pdf, "Thanks to everyone")
jt = find(ed_pdf, "Jottings")
check(af and th_ and fin and af[0][0] > fin[0][0] and jt and jt[0][0] > af[0][0],
      "end pages come after THE END and before the Notes pages")
check(th_ and images_on(ed_pdf)[th_[0][0] - 1] >= 1, "end pages: the picture is on the page")
check(jt and "Write your thoughts here." in txt and "Notes" not in [w["text"] for _p, w in words(ed_pdf)
                                                                     if _p == jt[0][0]],
      "Notes pages: edited heading and text above the lines")
cov = os.path.join(TMP, "cover.pdf")
bf_export.build_cover(book, s1, cov, 40)
cw = find(cov, "blurb for the back")
check(cw and any(w["text"] == "edited" and "Bold" in w["fontname"] for _p, w in find(cov, "An edited blurb")) and
      "original blurb" not in page_texts(cov)[0], "back-cover blurb: the edited text and its bold on the cover")

# ---- exports
d = tempfile.mkdtemp(dir=TMP)
bf_export.build_epub(book, s1, os.path.join(d, "b.epub"))
z = zipfile.ZipFile(os.path.join(d, "b.epub"))
epub = {n: z.read(n) for n in z.namelist()}
etxt = "\n".join(v.decode("utf-8", "replace") for n, v in epub.items() if n.endswith(".xhtml"))
check("Edited Book Title" in etxt and "<i>Edited Book Title</i>" in etxt and "For <b>Ann</b>" in etxt and
      "<u>underlined</u>" in etxt and "<b>FIN</b>" in etxt and "Afterword" in etxt and "Edited Fandom" in etxt and
      "The edited <b>summary</b>." in etxt, "EPUB: title page, dedication, note, work info, summary, THE END, end pages")
check(sum(1 for n in epub if n.startswith("OEBPS/images/img")) == 2 and RED_PNG in epub.values() and
      BLUE_PNG in epub.values(), "EPUB: the pictures of the note and the end pages are in the book")
import docx                                                            # noqa: E402
bf_export.build_docx(book, s1, os.path.join(d, "b.docx"))
dx = docx.Document(os.path.join(d, "b.docx"))
dtx = [p.text for p in dx.paragraphs]
ded_p = next((p for p in dx.paragraphs if p.text.startswith("For Ann")), None)
check("Edited Book Title" in dtx and ded_p is not None and any(r.bold and r.text == "Ann" for r in ded_p.runs) and
      "FIN" in dtx and "Afterword" in dtx and "Edited Note Heading" in dtx and len(dx.inline_shapes) == 2,
      "Word: title page, dedication (bold kept), note, THE END, end pages and both pictures")
bf_export.build_html(book, s1, os.path.join(d, "b.html"))
h = open(os.path.join(d, "b.html"), encoding="utf-8").read()
check("Edited Book Title" in h and "For <b>Ann</b>" in h and "Edited Note Heading" in h and "<b>FIN</b>" in h and
      "Thanks to everyone" in h and h.count("data:image/png") == 2, "web page: the edited pages and their pictures")
bf_export.build_txt(book, s1, os.path.join(d, "b.txt"))
tt = open(os.path.join(d, "b.txt"), encoding="utf-8").read()
check(all(x in tt for x in ("Edited Book Title", "For Ann, who read it first", "Edited Note Heading", "FIN",
                            "Afterword", "Edited Fandom", "The edited summary.")) and "THE END" not in tt,
      "plain text: the edited pages' words")

# ============================================================================ 2./3. contents edited as text
s2 = settings(toc=True)
b2 = apply_info(fic_book(), s2)
uids = [c.uid for c in b2.chapters]
titles = [c.title for c in b2.chapters]
auto = contents_entries(b2, s2)
rows = contents_rows(auto, [])
rows = [rows[2], rows[0], rows[1]]                 # third chapter first
rows[1] = dict(rows[1], markup="<b>First Chapter</b>", space=12)
rows[2] = dict(rows[2], show=False)
rows.insert(1, {"uid": None, "custom": True, "auto": "", "title": "A Map", "page": "ix", "show": True, "level": 1,
                "auto_level": 1, "markup": '<i>A Map</i>', "space": 0})
b2.toc = rows_to_edits(rows, auto)
p2 = make(b2, s2, "contents")
pages = page_texts(p2)
ci = next(i for i, p in enumerate(pages) if "CONTENTS" in p or "Contents" in p)
lines = [ln for ln in pages[ci].splitlines() if ln.strip() and "ONTENTS" not in ln.upper() and
         not re.fullmatch(r"[0-9ivxlc]+", ln.strip())]
check([c.uid for c in b2.chapters] == uids and [c.title for c in b2.chapters] == titles,
      "contents edits never change the chapters, their order or their ids")
check(len(lines) == 3 and lines[0].startswith("Third Chapter") and lines[1].startswith("A Map") and
      lines[1].rstrip().endswith("ix") and lines[2].startswith("First Chapter") and "Second Chapter" not in pages[ci],
      f"contents: reordered, a typed entry with its page, a hidden entry left out {lines}")
fc = find(p2, "First Chapter")
check(fc and fc[0][0] == ci + 1 and font_has(fc[:2], "Bold") and font_has(find(p2, "A Map")[:2], "Italic", "Oblique"),
      "contents: entry formatting (bold, italic) prints")
body = [i for i, p in enumerate(pages) if p.strip().startswith(("First Chapter", "Second Chapter", "Third Chapter"))]
order = [pages[i].strip().split("\n")[0] for i in body]
check(order == ["First Chapter", "Second Chapter", "Third Chapter"], f"the chapters still print in book order {order}")
# a typed page number is only what the contents prints
rows2 = contents_rows(auto, [])
rows2[1]["page"] = "42"
b2.toc = rows_to_edits(rows2, auto)
res = build_preview(b2, s2, log=lambda m: None)
cp = page_texts(res["pdf"])
second = next(ln for ln in cp[ci].splitlines() if ln.startswith("Second Chapter"))
real = cpages(res)
check(second.rstrip().endswith("42") and res["pages"] == len(pages), "a typed page number prints beside its entry")
b2.toc = rows_to_edits(contents_rows(auto, []), auto)
res0 = build_preview(b2, s2, log=lambda m: None)
second0 = next(ln for ln in page_texts(res0["pdf"])[ci].splitlines() if ln.startswith("Second Chapter"))
check(cpages(res0) == real and not second0.rstrip().endswith("42") and b2.toc == [],
      "the chapter stays on its real page; clearing the number goes back to the real one (automatic contents)")
# markup words follow a renamed entry
rows3 = contents_rows(auto, [])
rows3[0]["markup"] = "<b>First</b> Chapter"
b2.toc = rows_to_edits(rows3, auto)
b2.chapters[0].title = "Renamed Chapter"
p3 = make(b2, s2, "contents_renamed")
rc = find(p3, "Renamed Chapter")
check(rc and rc[0][0] == ci + 1 and font_has(rc[:1], "Bold"),
      "a formatted entry follows its chapter when the chapter is renamed (keeping its look)")

# ============================================================================ 4./8. projects
proj = os.path.join(TMP, "matter.bookforge")
bf_edit.save_project(proj, book, s1, "1.8.0")
b4, s4 = bf_edit.load_project(proj)
check(b4.matter.keys() == book.matter.keys() and all(
    bf_matter.override(b4, k) == bf_matter.override(book, k) and bf_matter.heading_override(b4, k) ==
    bf_matter.heading_override(book, k) for k in book.matter), "every edited page is saved in the project and comes back")
check(bf_matter.override(b4, "work_info")[0].text == "<b>Fandom:</b> Edited Fandom" and
      bf_matter.heading_override(b4, "work_notes") == "Notes From The Author" and
      bf_matter.override(b4, "note")[2].image == RED_PNG, "fanfic work info, notes and pictures survive save/reopen")
check(chars(make(b4, s4, "reopened")) == chars(ed_pdf), "the reopened project prints the same")
data = json.load(open(proj, encoding="utf-8"))
del data["book"]["matter"]                          # a project from before 1.8.0
json.dump(data, open(proj, "w", encoding="utf-8"))
b5, s5 = bf_edit.load_project(proj)
fresh = apply_info(fic_book(), s1)
fresh.title = "Edited Book Title"
check(b5.matter == {} and chars(make(b5, s5, "old")) == chars(make(fresh, s1, "fresh")),
      "an older project (no edited pages) opens with every page automatic and prints as before")

# ============================================================================ 7./9. preview = print, binding intact
s7 = settings(**ALL_ON, blank_pages_start="2", blank_pages_end="1", chapter_start="Right-hand page", fmt_impose=True,
              fmt_pdf=True, fmt_cover=False, fmt_epub=False, fmt_docx=False, fmt_html=False, fmt_txt=False,
              pn_mode="Roman front matter, 1 on first chapter")


def look_only(b):
    """Formatting-only edits (same lines, same sizes): nothing may move to another page."""
    for key, kind in (("title", "btitle"), ("the_end", "endline")):
        bl = bf_matter.auto_blocks(b, s7, key)
        bf_matter.set_override(b, key, [dataclasses.replace(x, text=f"<b>{x.text}</b>") if x.kind == kind else x
                                        for x in bl])
    bl = bf_matter.auto_blocks(b, s7, "dedication")
    bf_matter.set_override(b, "dedication", [dataclasses.replace(x, text=f"<i>{x.text}</i>") for x in bl])
    return b


plain_res = build_preview(apply_info(fic_book(), s7), dict(s7), log=lambda m: None)
b7 = look_only(apply_info(fic_book(), s7))
res7 = build_preview(b7, dict(s7), log=lambda m: None)
check(res7["pages"] == plain_res["pages"] and cpages(res7) == cpages(plain_res) and
      list(res7["toc_labels"].values()) == list(plain_res["toc_labels"].values()),
      "formatting the title page, dedication and THE END moves no page (count, chapter pages, contents numbers)")
pt = page_texts(res7["pdf"])
check(not pt[0].strip() and not pt[1].strip() and not pt[-1].strip(), "blank pages at the start and end stay blank")
starts = cpages(res7)
check(starts and all(p % 2 == 1 for p in starts), f"chapters still start on right-hand pages {starts}")
check(len(res7["sheets"]) == len(plain_res["sheets"]) and
      [len(pdfplumber.open(f).pages) for f in res7["sheets"]] == [len(pdfplumber.open(f).pages) for f in plain_res["sheets"]],
      "the imposed sheets have the same pages")
b7b = look_only(apply_info(fic_book(), s7))
bf_matter.set_override(b7b, "end_matter", [Block("para", text="An afterword. " * 30)])
res7b = build_preview(b7b, dict(s7), log=lambda m: None)
pt7b = page_texts(res7b["pdf"])
check(not pt7b[-1].strip() and cpages(res7b) == cpages(plain_res) and
      any("An afterword." in p for p in pt7b[-6:]), "end pages go before the blank end pages; chapters don't move")
out7 = run_build(look_only(apply_info(fic_book(), s7)), dict(s7), log=lambda m: None)
printed = [os.path.join(r, f) for r, _d, fs in os.walk(s7["out_dir"]) for f in fs if f.endswith("Print.pdf")]
check(printed and chars(printed[0]) == chars(res7["pdf"]), "preview = Make the book: same Print PDF with edited pages")

# ======================================================================================== window checks
try:
    import tkinter as tk
    _probe = tk.Tk()
    _probe.destroy()
    have_display = True
except Exception as e:                     # noqa: BLE001
    have_display = False
    print(f"SKIP  window checks (no display: {e})")

if have_display:
    from tkinter import messagebox, filedialog
    SAID = []
    ANSWER = {}
    for k in ("askyesno", "askyesnocancel", "showinfo", "showerror", "showwarning", "askokcancel"):
        setattr(messagebox, k, (lambda kind: (lambda title="", message="", **kw: SAID.append((kind, title, message))
                                              or ANSWER.get(kind, True)))(k))
    PICK = []
    filedialog.askopenfilename = lambda **kw: PICK.pop(0) if PICK else ""
    filedialog.asksaveasfilename = lambda **kw: PICK.pop(0) if PICK else ""
    import BookForge as BF
    import bf_theme as th
    app = BF.App()
    app._no_save = True
    errors = []
    app.report_callback_exception = lambda *a: errors.append("".join(__import__("traceback").format_exception(*a)))
    ASK = []
    app.ask_text = lambda parent, title, prompt, initial="": ASK.pop(0) if ASK else None
    ed, t = app.editor, app.editor.txt

    def pump(sec=0.2):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            time.sleep(0.01)

    def wait(cond, sec=180):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            if cond():
                return True
            time.sleep(0.02)
        return False

    def idle(sec=240):
        return wait(lambda: not app.busy, sec)

    def tab(key):
        app.shelf.select(th.tab_index(key))
        pump(0.15)

    def preview():
        idle()
        n = len(SAID)
        app.preview_book()
        idle()
        pump(0.2)
        pw = getattr(app, "preview_win", None)
        return (pw.result if pw is not None and pw.alive() else None), SAID[n:]

    def shown():
        return t.get("1.0", "end-1c")

    BOOK = "".join(f"Chapter {i}\n\n" + "\n\n".join(f"Paragraph {i}.{j}: " + "words of text here " * 10
                                                    for j in range(1, 4)) + "\n\n" for i in range(1, 4))
    app.var("source_mode").set("paste")
    app._mode_changed()
    app.texts["paste_text"].insert("1.0", BOOK)
    app.var("title").set("Window Book")
    app.var("author").set("Window Author")
    app.texts["dedication"].insert("1.0", "For the first reader")
    app.var("the_end").set(True)
    idle()
    app.scan(force=True)
    idle(60)
    tab("edit")
    names = [ed._item_text(r).strip() for r in ed.rows]
    check(names[0] == "FRONT PAGES" and "CHAPTERS" in names and "BACK PAGES" in names and
          len(ed.chapter_rows()) == 3 and any(n.startswith("Title page") for n in names) and
          any(n.startswith("THE END") for n in names), "Specific Text Edit lists the front pages, chapters and back pages")
    res, said = preview()
    before = chars(res["pdf"]) if res else None
    for key in ("title", "bookplate", "dedication", "contents", "the_end"):
        ed.select_section(key)
        pump(0.05)
        ed.dirty = True                            # saved without a change
        ed.commit()
    res, said = preview()
    check(res and chars(res["pdf"]) == before, "opening and saving pages without changing them changes nothing printed")
    for key in list(app.book.matter):
        app.clear_section(key)

    # ---- title page: words <-> Book info, formatting kept
    ed.select_section("title")
    pump(0.05)
    check(shown().split("\n")[:2] == ["Window Book", "Window Author"], "the title page shows the Book info words")
    t.tag_add("f_i", "1.0", "1.end")
    t.delete("1.0", "1.end")
    t.insert("1.0", "Edited In Editor", ("f_i",))
    ed.dirty = True
    tab("info")
    check(app.var("title").get() == "Edited In Editor", "changing the title line changes the Title on Book info")
    app.var("author").set("Changed Author")
    tab("edit")
    check(ed.sec == "title" and shown().split("\n")[:2] == ["Edited In Editor", "Changed Author"] and
          "f_i" in t.tag_names("1.2"), "an author changed on Book info shows on the edited title page; italics stay")
    check(app.unsaved and app.section_edited("title"), "an edited page counts as an unsaved edit")

    # ---- dedication: formatted here, the plain box can't be typed in
    ed.select_section("dedication")
    pump(0.05)
    t.tag_add("f_b", "1.4", "1.9")
    ed.dirty = True
    ed.select_section("the_end")
    box = app.texts["dedication"]
    check(str(box.cget("state")) == "disabled" and box.get("1.0", "end-1c") == "For the first reader" and
          "<b>" in bf_matter.override(app.book, "dedication")[0].text,
          "a formatted dedication: kept with the book; the Book info box shows its words and can't be typed in")
    t.delete("1.0", "end")
    t.insert("1.0", "The Very End")
    t.tag_add("f_b", "1.0", "1.end")
    ed.dirty = True

    # ---- contents as text
    ed.select_section("contents")
    pump(0.05)
    uids = [c.uid for c in app.book.chapters]
    check(shown().split("\n")[:3] == ["Chapter 1", "Chapter 2", "Chapter 3"] and len(ed.marks) == 3,
          "the contents show one line per entry, each with its page number")
    t.focus_force()
    t.mark_set("insert", "3.end")                   # after Chapter 3's page number
    for ks in ("BackSpace", "x", "BackSpace"):
        t.event_generate("<KeyPress>", keysym=ks)
        t.event_generate("<KeyRelease>", keysym=ks)
        pump(0.05)
    t.mark_set("insert", "3.end-1c")
    t.event_generate("<KeyPress>", keysym="Delete")
    pump(0.05)
    check(shown().split("\n")[2] == "Chapter " and ed._mark_on_line(3) == f"u:{uids[2]}",
          "Backspace and typing at the end of an entry change its words, never its page number; Delete keeps it too")
    t.insert("3.end-1c", "3")
    t.tag_add("sel", "2.0", "3.0")                  # Chapter 2's whole line selected (as a triple-click does)...
    t.mark_set("insert", "3.0")
    for ks in ("Z", "BackSpace"):                   # ...and typed over, then the typing taken back
        t.event_generate("<KeyPress>", keysym=ks)
        t.event_generate("<KeyRelease>", keysym=ks)
        pump(0.05)
    t.insert("2.0", "Chapter 2")
    check(shown().split("\n")[:3] == ["Chapter 1", "Chapter 2", "Chapter 3"] and
          [ed._mark_on_line(n) for n in (1, 2, 3)] == [f"u:{u}" for u in uids],
          "typing over an entry's selected line keeps its page number and its link to the chapter")
    t.insert("1.9", " (the start)")
    t.tag_add("f_b", "1.0", "1.7")
    t.insert("2.0", "Map of the Island\n")
    t.delete("4.0", "5.0")                          # Chapter 3's line
    t.tag_add("in|1", "3.0", "3.end+1c")
    marks = [w for w in ed.mark_widgets.values()]
    ASK.append("99")
    ed.contents_page(marks[0])
    ed.dirty = True
    ed.commit()
    toc = app.book.toc
    check([c.uid for c in app.book.chapters] == uids and [c.title for c in app.book.chapters] ==
          ["Chapter 1", "Chapter 2", "Chapter 3"], "editing the contents never changes the chapters or their ids")
    r = contents_rows(contents_entries(app.book, app.collect()), toc)
    check([x["uid"] for x in r if not x["custom"]] == uids and
          [(x["title"] or x["auto"], x["show"], x["level"]) for x in r] ==
          [("Chapter 1 (the start)", True, 0), ("Map of the Island", True, 0), ("Chapter 2", True, 1),
           ("Chapter 3", False, 0)] and r[0]["page"] == "99" and "<b>" in (r[0]["markup"] or ""),
          "renamed + bold, added, indented, hidden (line removed) and a typed page number")
    check(ed.sec == "contents" and shown().split("\n")[:3] == ["Chapter 1 (the start)", "Map of the Island",
                                                                "Chapter 2"], "the text shows the contents as saved")
    idle()
    app.edit_contents()
    pump(0.1)
    cw = app.contents_win
    check(cw is not None and [x["title"] or x["auto"] for x in cw.rows][:2] == ["Chapter 1 (the start)",
                                                                                 "Map of the Island"] and
          str(t.cget("state")) == "disabled", "Edit contents shows the same entries; the text waits while it is open")
    cw.tree.selection_set("3")
    cw.toggle()                                     # Chapter 3 back on the contents page
    cw.ok()
    pump(0.1)
    check(str(t.cget("state")) == "normal" and shown().split("\n")[3] == "Chapter 3",
          "after Edit contents closes, the text shows its changes")
    res, said = preview()
    pages = page_texts(res["pdf"]) if res else []
    cp = next((p for p in pages if "Map of the Island" in p), "")
    check(re.search(r"Chapter 1 \(the start\)[ .]*99", cp) and cpages(res)[0] != 99,
          "the preview prints the typed page number; the chapter stays on its real page")
    marks = list(ed.mark_widgets.values())
    ASK.append("")
    ed.contents_page(marks[0])
    ed.commit()
    check(contents_rows(contents_entries(app.book, app.collect()), app.book.toc)[0]["page"] == "",
          "clearing the typed number goes back to the real page")
    ANSWER["askyesno"] = True
    ed.back_to_automatic()
    check(app.book.toc == [] and shown().split("\n")[:3] == ["Chapter 1", "Chapter 2", "Chapter 3"],
          "Back to automatic: the automatic contents again")
    t.mark_set("insert", "3.2")
    ed.move_chapter(-1)                             # the Move up button: the entry, not a chapter
    ed.move_chapter(-1)
    ed.commit()
    r = contents_rows(contents_entries(app.book, app.collect()), app.book.toc)
    check([x["uid"] for x in r] == [uids[2], uids[0], uids[1]] and [c.uid for c in app.book.chapters] == uids and
          shown().split("\n")[:3] == ["Chapter 3", "Chapter 1", "Chapter 2"],
          "Move up moves a contents entry (still linked to its chapter); the chapters stay in their order")
    ed.back_to_automatic()

    # ---- pictures: allowed on the extra front page, not on the title page
    ed.select_section("title")
    PICK.append(os.path.join(TMP, "never.png"))
    ed.insert_picture()
    check(PICK and "Pictures can go" in str(ed.status.cget("text")), "no pictures on the title page (it says where)")
    PICK.clear()
    pic = os.path.join(TMP, "red.png")
    Image.new("RGB", (200, 120), (220, 30, 30)).save(pic)
    app.var("note_title").set("Window Note")
    ed.select_section("note")
    t.insert("1.0", "Note text with a picture below.")
    t.tag_add("f_u", "1.0", "1.4")
    t.mark_set("insert", "1.end")
    PICK.append(pic)
    ed.insert_picture()
    ed.commit()
    nb = bf_matter.override(app.book, "note")
    check(nb and [b.kind for b in nb] == ["para", "image"] and nb[1].image == open(pic, "rb").read() and
          "<u>Note</u>" in nb[0].text, "extra front page: text with underline and a picture")

    # ---- light / dark and tabs
    ed.select_section("the_end")
    pump(0.05)
    app.toggle_dark()
    pump(0.5)
    ed, t = app.editor, app.editor.txt
    check(ed.sec == "the_end" and shown().strip() == "The Very End" and "f_b" in t.tag_names("1.2"),
          "light/dark: the page shown stays, with its text and bold")
    ed.select(1)
    first = shown()
    app.toggle_dark()
    pump(0.5)
    ed, t = app.editor, app.editor.txt
    check(ed.cur == 1 and shown() == first and first.startswith("Paragraph 2.1"),
          "light/dark: a chapter shown in Specific Text Edit is shown again (its text box is no longer empty)")
    for key in ("src", "info", "front", "chap", "lay", "out", "edit"):
        tab(key)
    check(str(app.texts["dedication"].cget("state")) == "disabled" and
          "<b>" in bf_matter.override(app.book, "dedication")[0].text and
          bf_matter.override(app.book, "the_end")[0].text == "<b>The Very End</b>",
          "the edited pages are kept through tab switching and both theme changes")

    # ---- preview = Make the book; project save / reopen
    res, said = preview()
    outdir = tempfile.mkdtemp(dir=TMP)
    app.var("out_dir").set(outdir)
    app.var("open_when_done").set(False)
    for k, v in dict(fmt_pdf=True, fmt_impose=True, fmt_cover=False, fmt_epub=False, fmt_docx=False,
                     fmt_html=False, fmt_txt=False).items():
        app.var(k).set(v)
    res, said = preview()
    idle()
    app.build()
    idle()
    printed = [os.path.join(r, f) for r, _d, fs in os.walk(outdir) for f in fs if f.endswith("Print.pdf")]
    ptxt = "\n".join(page_texts(printed[0])) if printed else ""
    check(printed and res and chars(printed[0]) == chars(res["pdf"]) and "Edited In Editor" in ptxt and
          "The Very End" in ptxt and "Window Note" in ptxt, "Make the book: the Print PDF is the preview, edits included")
    proj = os.path.join(TMP, "window.bookforge")
    idle()
    PICK.append(proj)
    app.save_project_dialog()
    check(not app.unsaved, "the project is saved")
    PICK.append(proj)
    idle()
    app.open_project_dialog()
    pump(0.3)
    check(app.book.matter.keys() >= {"title", "dedication", "the_end", "note"} and
          str(app.texts["dedication"].cget("state")) == "disabled" and
          bf_matter.override(app.book, "note")[1].image == open(pic, "rb").read(),
          "the edited pages come back with the project (the dedication box still follows the edited page)")

    # ---- Back to automatic for the dedication: the box can be typed in again
    ed.select_section("dedication")
    ed.back_to_automatic()
    check(str(app.texts["dedication"].cget("state")) == "normal" and not app.section_edited("dedication") and
          app.texts["dedication"].get("1.0", "end-1c") == "For the first reader",
          "Back to automatic: the dedication is plain text in the Book info box again")

    # ---- fanfic pages from a real (AO3-style) download: edited, kept in projects, replaced (after asking) when
    # the source is read again; a chapter's own notes stay with the chapter
    def ao3_chapter(n):
        return (f'<div class="meta group"><h2 class="heading">Chapter {n}</h2><p>Chapter Notes</p>'
                f'<blockquote class="userstuff"><p>A note for chapter {n} only.</p></blockquote></div>'
                f'<div class="userstuff"><p>' + f"Chapter {n} words of the story. " * 30 + "</p></div>")
    ao3 = os.path.join(TMP, "story.html")
    with open(ao3, "w", encoding="utf-8") as f:
        f.write('<!DOCTYPE html><html><head><meta charset="UTF-8"><title>The Keeper - Test Author - Original Work'
                '</title></head><body><div id="preface"><div class="meta"><dl class="tags">'
                "<dt>Rating:</dt><dd>General Audiences</dd><dt>Archive Warning:</dt><dd>No Archive Warnings Apply</dd>"
                "<dt>Fandom:</dt><dd>Original Work</dd><dt>Language:</dt><dd>English</dd>"
                "<dt>Stats:</dt><dd>Published: 2026-01-01 Chapters: 2/2 Words: 1000</dd></dl>"
                '<h1>The Keeper</h1><div class="byline">by <a rel="author" href="x">Test Author</a></div>'
                '<p>Summary</p><blockquote class="userstuff"><p>A keeper looks after a light on a rock.</p>'
                '</blockquote><p>Notes</p><blockquote class="userstuff"><p>Thanks to my beta reader.</p></blockquote>'
                '</div></div><div id="chapters" class="userstuff">' + ao3_chapter(1) + ao3_chapter(2) +
                "</div></body></html>")
    app.var("source_mode").set("file")
    app._mode_changed()
    app.var("source_path").set(ao3)
    app.var("fic_info").set("Add a work info page")
    idle()                                          # (a preview being made after the last change finishes first)
    app.scan(force=True)
    wait(lambda: not app.busy and app.book is not None and app.book.fic_meta, 120)
    pump(0.2)
    check(len(app.book.fic_meta) == 5 and app.book.fic_summary and app.book.fic_notes,
          "an AO3-style download: work info, summary and work notes are read")
    tab("edit")
    names = [ed._item_text(r).strip() for r in ed.rows]
    check(all(any(n.startswith(x) for n in names) for x in ("Work info", "Summary", "Author's notes")),
          "its work info, summary and author's notes are in the list")
    ed.select_section("work_info")
    t.insert("1.end", " And Up")
    t.tag_add("f_i", "3.0", "3.end")
    ed.select_section("summary")
    t.insert("1.end", " An edited summary.")
    t.tag_add("al|center", "1.0", "1.end+1c")
    ed.select_section("work_notes")
    app.editor.title_var.set("Notes From The Author")
    t.insert("1.0", "Edited: ")
    ed.select(0)
    check(app.edited and bf_matter.override(app.book, "work_info")[0].text.endswith("And Up") and
          bf_matter.override(app.book, "summary")[0].align == "center" and
          bf_matter.heading_override(app.book, "work_notes") == "Notes From The Author",
          "work info, summary and notes edited (counted as changes to what was read from the source)")
    res, said = preview()
    txt = "\n".join(page_texts(res["pdf"])) if res else ""
    fd = find(res["pdf"], "Fandom: Original Work") if res else []
    check("General Audiences And Up" in txt and "An edited summary." in txt and "Notes From The Author" in txt and
          "Edited: Thanks to my beta reader." in txt and font_has(fd, "Italic", "Oblique"),
          "the preview prints the edited work info (italic line), summary and notes")
    check(any(b.kind == "anote" and "chapter 1 only" in b.text for b in app.book.chapters[0].blocks) and
          "A note for chapter 1 only." in txt and "A note for chapter 2 only." in txt,
          "each chapter's own notes stay with (and print in) that chapter")
    idle()
    PICK.append(proj)
    app.save_project_dialog()
    PICK.append(proj)
    idle()
    app.open_project_dialog()
    pump(0.3)
    check(bf_matter.override(app.book, "work_info")[0].text.endswith("And Up") and
          bf_matter.heading_override(app.book, "work_notes") == "Notes From The Author" and
          len(app.book.fic_meta) == 5, "fanfic work info, summary and notes: saved and reopened")
    check("Work info" in app.changes_text() and "kept" in app.changes_text(),
          "Re-read source warns that the work info edit would be lost (and that your own pages are kept)")
    ANSWER["askyesno"] = True
    idle()
    app.reread_source()
    idle(60)
    pump(0.2)
    check(not bf_matter.is_edited(app.book, "work_info") and not bf_matter.is_edited(app.book, "work_notes") and
          bf_matter.is_edited(app.book, "title") and bf_matter.is_edited(app.book, "the_end"),
          "after Re-read source: the source's work info and notes again; the title page and THE END edits are kept")
    idle()
    app.edited, app.unsaved = False, False
    ed.select_section("bookplate")
    t.insert("1.0", "Ex Libris: ")
    n = len(SAID)
    ANSWER["askyesno"] = False
    app.new_book()
    check(app.book is not None and any(k == "askyesno" for k, *_x in SAID[n:]) and
          bf_matter.override(app.book, "bookplate") is not None,
          "New book asks first when only a page was edited (typing not yet saved included); No keeps it")
    app._closing = False
    t.insert("1.0", "More ")
    app.unsaved = False
    ANSWER["askyesnocancel"] = None
    n = len(SAID)
    app.on_close()
    check(any(k == "askyesnocancel" for k, *_x in SAID[n:]) and app.winfo_exists(),
          "closing BookForge right after typing on a page asks to save (Cancel keeps working)")
    check(not errors, f"no window errors {errors[:1]}")
    app.destroy()

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
