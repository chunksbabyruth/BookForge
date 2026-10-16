"""Editing regression tests (plain Python - run:  python tests/test_editor.py).

Round-trips real books through the editor model when they are in tests/inputs/ (skipped otherwise), and checks
editor formatting in the PDF, EPUB, HTML and Word outputs, project files, and preview == final with edits.
"""
import os
import shutil
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
INPUTS = os.environ.get("BOOKFORGE_TEST_INPUTS", os.path.join(HERE, "inputs"))

import pdfplumber                                                     # noqa: E402
import bf_impose  # noqa: E402,F401  (registers binding defaults)
from bf_core import DEFAULTS, Book, Chapter, Block, load_book, plain, apply_info   # noqa: E402
from bf_edit import (blocks_to_paras, paras_to_blocks, parse_markup, runs_to_markup,   # noqa: E402
                     save_project, load_project, BR)
from bf_build import build_preview, run_build                         # noqa: E402

FAILS = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


# ---------------------------------------------------------------- markup round trips
for m in ["Plain", "<i>it</i> and <b>bold <i>both</i></b>", "x<super>2</super> H<sub>2</sub>O <strike>old</strike>",
          "<u>under</u> &amp; &lt;tag&gt; caf\u00e9 \u2014 \u201cq\u201d", "line one<br/>line two",
          '<font face="EB Garamond" size="16" color="#aa0000">Big <b>red</b></font> end']:
    back = runs_to_markup(parse_markup(m))
    check(parse_markup(back) == parse_markup(m) and plain(back) == plain(m), f"markup round trip: {m[:40]!r}")
check(BR in "".join(t for t, _ in parse_markup("a<br/>b")), "line breaks become the editor's line-break sign")

# ---------------------------------------------------------------- real books: every chapter survives the editor
for name in ("gutenberg.epub", "ao3_story.epub", "ao3_oneshot.html",
             "standard_ebooks.epub"):
    path = os.path.join(INPUTS, name)
    if not os.path.exists(path):
        print(f"SKIP  {name} (not in {INPUTS})")
        continue
    s = dict(DEFAULTS)
    s["source_path"] = path
    book = load_book(s)
    bad, n_blocks, dropped = [], 0, 0
    for ch in book.chapters:
        new = paras_to_blocks(blocks_to_paras(ch.blocks))
        old = [b for b in ch.blocks if b.kind in ("image", "scene") or plain(b.text or "".join(b.lines)).strip()]
        dropped += len(ch.blocks) - len(old)
        n_blocks += len(old)
        if len(new) != len(old):
            bad.append((ch.title[:30], "count", len(old), len(new)))
            continue
        for a, b in zip(old, new):
            same = a.kind == b.kind and (
                [parse_markup(x) for x in a.lines] == [parse_markup(x) for x in b.lines] if a.kind == "verse"
                else a.image == b.image if a.kind == "image"
                else parse_markup(a.text) == parse_markup(b.text) and a.label == b.label)
            if not same:
                bad.append((ch.title[:30], a.kind, plain(a.text)[:40]))
    check(not bad, f"{name}: all {n_blocks} paragraphs come back unchanged through the editor "
                   f"({dropped} empty dropped) {bad[:2]}")

# ---------------------------------------------------------------- editor formatting reaches the PDF
tmp = tempfile.mkdtemp(prefix="bf_edit_")
blocks = [Block("para", text='Normal start <font face="EB Garamond" size="16" color="#aa0000">BigRed</font> '
                                'then <u>Under</u> <strike>Struck</strike> H<sub>2</sub>O x<super>2</super>.'),
          Block("para", text="CENTREDLINE", align="center"),
          Block("para", text="RIGHTLINE", align="right"),
          Block("para", text="INDENTEDLINE", indent=4),
          Block("head", text="A heading added in the editor"),
          Block("quote", text="A quotation.", space=24)]
book = Book(title="Edited", author="Tester", chapters=[Chapter(title="Edited chapter", blocks=blocks),
                                                     Chapter(title="Second", blocks=[Block("para", text="More.")])])
s = dict(DEFAULTS)
s.update(source_mode="paste", title="Edited", fmt_impose=True, fmt_cover=False, fmt_epub=True, fmt_docx=True,
         fmt_html=True, fmt_txt=True, out_dir=os.path.join(tmp, "out"), drop_cap=False)
res = build_preview(apply_info(book, dict(s)), dict(s), log=lambda m: None)
with pdfplumber.open(res["pdf"]) as d:
    words = [w for p in d.pages for w in p.extract_words(extra_attrs=["fontname", "size", "non_stroking_color"])]
    page_w = d.pages[0].width
w = {x["text"].strip(".,"): x for x in words}
big = w.get("BigRed")
check(big and big["fontname"].split("+")[-1].startswith("EBGaramond") and abs(big["size"] - 16) < 0.6,
      f"editor font + size used in the PDF: {big and (big['fontname'].split('+')[-1], round(big['size'], 1))}")
col = big and big.get("non_stroking_color")
check(col and abs(col[0] - 0xAA / 255) < 0.02 and col[1] < 0.02, f"editor colour used in the PDF: {col}")
for word, test in (("CENTREDLINE", lambda x: abs((x["x0"] + x["x1"]) / 2 - page_w / 2) < 20),
                   ("RIGHTLINE", lambda x: page_w - x["x1"] < 100),
                   ("INDENTEDLINE", lambda x: x["x0"] > w["Normal"]["x0"] + 60)):
    check(word in w and test(w[word]), f"paragraph formatting in the PDF: {word} at x={w.get(word, {}).get('x0')}")
folder, made, notes = run_build(apply_info(book, dict(s)), dict(s), log=lambda m: None)
final = [m for m in made if m.endswith("Print.pdf")][0]
with pdfplumber.open(final) as d1, pdfplumber.open(res["pdf"]) as d2:
    check([p.extract_text() for p in d1.pages] == [p.extract_text() for p in d2.pages],
          "edited book: final Print PDF = preview, page for page")
shutil.rmtree(res["dir"], ignore_errors=True)

# ---------------------------------------------------------------- exports keep the formatting
html_out = open([m for m in made if m.endswith(".html")][0], encoding="utf-8").read()
check("font-family:'EB Garamond'" in html_out and "color:#aa0000" in html_out and "text-align:center" in html_out,
      "HTML keeps editor font, colour and alignment")
epub = [m for m in made if m.endswith(".epub")][0]
with zipfile.ZipFile(epub) as z:
    xhtml = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xhtml"))
check("font-size:16pt" in xhtml and "margin-left:1in" in xhtml, "EPUB keeps editor size and indent")
import docx                                                            # noqa: E402
d = docx.Document([m for m in made if m.endswith(".docx")][0])
runs = [r for p in d.paragraphs for r in p.runs if "BigRed" in r.text]
check(runs and runs[0].font.name == "EB Garamond" and runs[0].font.size.pt == 16 and
      str(runs[0].font.color.rgb) == "AA0000", "Word keeps editor font, size and colour")
paras = {p.text: p for p in d.paragraphs}
check(paras.get("CENTREDLINE") is not None and paras["CENTREDLINE"].alignment == 1, "Word keeps centred alignment")
txt = open([m for m in made if m.endswith(".txt")][0], encoding="utf-8").read()
check("BigRed" in txt and "A heading added in the editor" in txt, "plain text keeps the edited words")

# ---------------------------------------------------------------- projects
img = Block("image", image=b"\x89PNG\r\n\x1a\nfakeimage", image_name="pic.png")
book.chapters[1].blocks.append(img)
pj = os.path.join(tmp, "my book.bookforge")
save_project(pj, book, dict(s, title="Edited"), "test")
book2, s2 = load_project(pj)
check(book2 == book, "project: book with all edits (and pictures) reopens identically")
check(s2.get("title") == "Edited" and s2.get("fmt_docx") is True, "project: settings reopen")
open(os.path.join(tmp, "junk.bookforge"), "w").write("not json")
try:
    load_project(os.path.join(tmp, "junk.bookforge"))
    check(False, "damaged project is refused")
except ValueError as e:
    check("can't be read" in str(e), f"damaged project -> plain message: {e}")
shutil.rmtree(tmp, ignore_errors=True)

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
