"""Regression tests for the 1.5.1 fixes (plain Python - run:  python tests/test_regressions_151.py).

Covers: "Rotate back sides 180" on single-sided printers, the zine warning, bookbinder-js print order and
octavo layout, landscape paper sizes, larger pages on a short small-format book, Word tables / pictures /
tracked changes, the Contents "Chapter N:" prefix in UPPERCASE, running heads on the Notes pages and the
version number. Every file used is generated here, so the checks also run on GitHub Actions.
"""
import io
import os
import re
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix="bf_regress151_")
os.environ["APPDATA"] = os.path.join(TMP, "appdata")          # keep the user's real settings out of it
os.environ["HOME"] = os.path.join(TMP, "home")
os.makedirs(os.environ["APPDATA"], exist_ok=True)
os.makedirs(os.environ["HOME"], exist_ok=True)

import pdfplumber                                                     # noqa: E402
from pypdf import PdfReader                                           # noqa: E402
from reportlab.pdfgen import canvas                                   # noqa: E402
import bf_impose  # noqa: E402,F401  (registers binding defaults)
from bf_impose import impose, IMPOSE_DEFAULTS                         # noqa: E402
from bf_core import DEFAULTS, APP_VERSION, load_book, apply_info, plain   # noqa: E402
from bf_build import run_build                                        # noqa: E402

FAILS = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


def numbered(n, w=5.5 * 72, h=8.5 * 72):
    path = os.path.join(TMP, f"num{n}_{int(w)}x{int(h)}.pdf")
    c = canvas.Canvas(path, pagesize=(w, h))
    for i in range(1, n + 1):
        c.setFont("Helvetica-Bold", 40)
        c.drawCentredString(w / 2, h / 2, f"P{i}")
        c.showPage()
    c.save()
    return path


def imp(n=16, **kw):
    s = dict(IMPOSE_DEFAULTS)
    s.update(kw)
    logs = []
    info = impose(numbered(n), tempfile.mkdtemp(dir=TMP), "t", s, log=logs.append)
    info["log"] = logs
    return info


def labels(page):
    """[(page number, upright?, x0, top)] for the "P<n>" labels on one printed side."""
    out = []
    for w in page.extract_words():
        t = w["text"]
        if re.fullmatch(r"P\d+", t):
            out.append((int(t[1:]), True, w["x0"], w["top"]))
        elif re.fullmatch(r"\d+P", t):                      # upside down: the letters come out reversed
            out.append((int(t[:-1][::-1]), False, w["x0"], w["top"]))
    return out


def label_size(page):
    sizes = [c["size"] for c in page.chars if c["text"] == "P"]
    return round(max(sizes), 1) if sizes else 0


SINGLE = "Single-sided (fronts and backs files)"
REVERSE = "Single-sided, backs in reverse order"

# ---------------------------------------------------------------- 1. "Rotate back sides 180" when single-sided
# Bug: the option only turned the back sides in the duplex file; the single-sided "backs" file ignored it.
for printer in (SINGLE, REVERSE):
    for layout in ("Folio - 2 pages per side", "Small: 18 per side (3-folio strips x 3)"):
        for rb in (False, True):
            info = imp(imp_layout=layout, printer_type=printer, rotate_back=rb)
            fronts = next(f for f in info["files"] if "fronts" in f)
            backs = next(f for f in info["files"] if "backs" in f)
            fr = {int(p.get("/Rotate", 0) or 0) for p in PdfReader(fronts).pages}
            bk = {int(p.get("/Rotate", 0) or 0) for p in PdfReader(backs).pages}
            want = {180} if rb else {0}
            check(fr == {0} and bk == want,
                  f"{printer[:28]}, {layout[:10]}, rotate back={rb}: fronts {sorted(fr)}, backs {sorted(bk)}")

# ---------------------------------------------------------------- 2. the zine says when pages are left out
# Bug: a book longer than 8 pages quietly lost every page after the 8th.
info = imp(12, imp_layout="Single-sheet zine: 8 pages, one side")
check("4 of the 12 pages" in info.get("warning", "") and any(m.startswith("Note:") for m in info["log"]),
      f"zine with 12 pages warns that 4 are left out ({info.get('warning')})")
info = imp(8, imp_layout="Single-sheet zine: 8 pages, one side")
check(not info.get("warning"), "zine with 8 pages: no warning")
s = dict(DEFAULTS)
s.update(out_dir=tempfile.mkdtemp(dir=TMP), source_mode="paste", open_when_done=False, fmt_impose=True,
         imp_layout="Single-sheet zine: 8 pages, one side",
         paste_text="".join(f"\n\nChapter {n}\n\n" + "Words in the chapter. " * 300 for n in (1, 2, 3)))
book = apply_info(load_book(s, s["paste_text"]), s)
_folder, made, notes = run_build(book, s, lambda m: None)
check(any("WARNING" in n and "left out" in n for n in notes), "zine warning is also in the printing notes")

# ---------------------------------------------------------------- 3. print order and octavo as in bookbinder-js
# Each signature's innermost sheet first, inside face first (folio booklet of 8 pages: 4|5, 3|6, 2|7, 1|8).
info = imp(8, imp_layout="Folio - 2 pages per side", sig_format="Booklet (one signature)")
with pdfplumber.open(info["files"][0]) as pdf:
    order = [sorted(n for n, *_ in labels(p)) for p in pdf.pages]
check(order == [[4, 5], [3, 6], [2, 7], [1, 8]], f"folio booklet prints innermost sheet, inside face first {order}")
check("FIRST facing UP" in info["instructions"] and "sheet 1 into the centre of sheet 2" in info["instructions"],
      "folding notes match the new order")

info = imp(16, imp_layout="Octavo - 8 pages per side", sig_format="Perfect bound (every sheet separate)")
with pdfplumber.open(info["files"][0]) as pdf:
    got = []
    for p in pdf.pages:
        cells = labels(p)
        top = [n for n, up, x, t in sorted(cells, key=lambda c: c[2]) if t < p.height / 2]
        bottom = [n for n, up, x, t in sorted(cells, key=lambda c: c[2]) if t >= p.height / 2]
        ups = {(t < p.height / 2): up for n, up, x, t in cells}
        got.append((top, bottom, ups.get(True), ups.get(False)))
want = [([7, 10, 11, 6], [2, 15, 14, 3], False, True), ([5, 12, 9, 8], [4, 13, 16, 1], False, True)]
check(got == want, f"octavo sheet laid out as bookbinder-js lays it out {got}")
check("fold 1: bring the RIGHT edge over to the LEFT edge" in info["instructions"],
      "octavo folding notes start with the right edge over to the left")

# ---------------------------------------------------------------- 4. landscape paper entries (as bookbinder-js)
# Ledger (17 x 11) is listed landscape, so it lays out like Tabloid with "Rotate paper 90".
a = imp(imp_layout="Folio - 2 pages per side", sheet_size="Ledger (17 x 11)")
b = imp(imp_layout="Folio - 2 pages per side", sheet_size="Tabloid / 11x17", rotate_paper=True)
c = imp(imp_layout="Folio - 2 pages per side", sheet_size="Tabloid / 11x17")
check(a["sheet_in"] == b["sheet_in"] == (11.0, 17.0) and c["sheet_in"] == (17.0, 11.0),
      f"Ledger folio = Tabloid turned 90 ({a['sheet_in']}, {b['sheet_in']}, {c['sheet_in']})")
d = imp(imp_layout="Folio - 2 pages per side", sheet_size="Letter")
check(d["sheet_in"] == (11.0, 8.5), f"Letter folio is unchanged ({d['sheet_in']})")
e = imp(imp_layout="Folio - 2 pages per side", sheet_size="Custom", sheet_w=17, sheet_h=11, sheet_unit="in")
check(e["sheet_in"] == (11.0, 17.0), f"custom 17 x 11 lays out like Ledger ({e['sheet_in']})")

# ---------------------------------------------------------------- 5. short book on a small layout fills the sheet
# As bookbinder-js: rows that are empty on both sides of the first sheet are left out, so pages print larger,
# and the fronts and backs keep the same page size (so they still line up).
full = imp(12, imp_layout="Small book: 6 per side (3-folio signature)")
short = imp(4, imp_layout="Small book: 6 per side (3-folio signature)")
one_side_blank = imp(7, imp_layout="Petite: 16 per side (4-folio signatures x 2)")
with pdfplumber.open(full["files"][0]) as f1, pdfplumber.open(short["files"][0]) as f2, \
        pdfplumber.open(one_side_blank["files"][0]) as f3:
    s_full = [label_size(p) for p in f1.pages]
    s_short = [label_size(p) for p in f2.pages]
    s_seven = [label_size(p) for p in f3.pages]
check(s_short[0] > s_full[0] and s_short[0] == s_short[1], f"4 pages print larger than 12, same on both sides "
      f"({s_full} vs {s_short})")
check(len(set(s_seven)) == 1, f"a row blank on one side only is kept, so both sides match ({s_seven})")

# ---------------------------------------------------------------- 6. Word: tables, pictures, tracked changes
# Bug: only top-level paragraphs were read, so table text was lost, pictures were never imported and text in
# tracked insertions (w:ins) disappeared.
import docx                                                        # noqa: E402
from docx.oxml import parse_xml                                    # noqa: E402
from docx.oxml.ns import nsdecls                                   # noqa: E402
from docx.shared import Inches                                     # noqa: E402
from PIL import Image                                              # noqa: E402

d = docx.Document()
d.add_heading("Chapter One", 1)
d.add_paragraph("Before the table.")
t = d.add_table(rows=2, cols=3)
for i, txt in enumerate(("Name", "Role", "Age")):
    t.cell(0, i).text = txt
m = t.cell(1, 0).merge(t.cell(1, 1))
m.text = "Merged"
t.cell(1, 2).text = "34"
d.add_paragraph("After the table.")
p = d.add_paragraph("The cat ")
W = nsdecls("w")
p._p.append(parse_xml(f'<w:ins {W} w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r>'
                      '<w:t xml:space="preserve">quietly </w:t></w:r></w:ins>'))
p._p.append(parse_xml(f'<w:del {W} w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r>'
                      '<w:delText xml:space="preserve">loudly </w:delText></w:r></w:del>'))
p._p.append(parse_xml(f'<w:r {W}><w:t>sat down.</w:t></w:r>'))
img = io.BytesIO()
Image.new("RGB", (40, 30), (200, 40, 40)).save(img, "PNG")
img.seek(0)
p = d.add_paragraph("Words before. ")
p.add_run().add_picture(img, width=Inches(1))
p.add_run(" Words after.")
d.add_heading("Chapter Two", 1)
d.add_paragraph("The end.")
wd = os.path.join(TMP, "tables.docx")
d.save(wd)
for pics in (False, True):
    s = dict(DEFAULTS)
    s.update(source_path=wd, include_images=pics)
    blocks = [(b.kind, plain(b.text)) for c in load_book(s).chapters for b in c.blocks]
    texts = [x for k, x in blocks if k == "para"]
    check("Name Role Age" in texts and "Merged 34" in texts and
          texts.index("Before the table.") < texts.index("Name Role Age") < texts.index("After the table."),
          f"Word table rows imported in place, merged cell once (pictures {pics})")
    check("The cat quietly sat down." in texts and not any("loudly" in x for x in texts),
          "tracked insertion kept, tracked deletion left out")
    kinds = [k for k, x in blocks]
    if pics:
        i = kinds.index("image") if "image" in kinds else -1
        check(i > 0 and blocks[i - 1] == ("para", "Words before.") and blocks[i + 1] == ("para", "Words after."),
              f"picture imported between the words around it {blocks[-4:]}")
    else:
        check("image" not in kinds and ("para", "Words before. Words after.") in blocks,
              "pictures off: no picture, paragraph unchanged")

# ---------------------------------------------------------------- 7. Contents "Chapter N:" prefix and case
# Bug: with Contents case = UPPERCASE the titles were upper-cased but the "Chapter N:" prefix stayed "Chapter".
d = docx.Document()
for title in ("Morning tide", "the lighthouse keeper"):
    d.add_heading(title, 1)
    for _ in range(6):
        d.add_paragraph("Plain words in a chapter. " * 40)
tocdoc = os.path.join(TMP, "toc.docx")
d.save(tocdoc)
for case, want, not_want in (("UPPERCASE", ["CHAPTER 1: MORNING TIDE", "CHAPTER 2: THE LIGHTHOUSE KEEPER"], "Chapter"),
                             ("Title Case", ["Chapter 1: Morning Tide", "Chapter 2: The Lighthouse Keeper"], "CHAPTER"),
                             ("As written", ["Chapter 1: Morning tide", "Chapter 2: the lighthouse keeper"], "CHAPTER")):
    s = dict(DEFAULTS)
    s.update(out_dir=tempfile.mkdtemp(dir=TMP), source_mode="file", source_path=tocdoc, open_when_done=False,
             fmt_impose=False, toc=True, toc_numbers="Chapter N:", toc_case=case)
    book = apply_info(load_book(s), s)
    _f, made, _n = run_build(book, s, lambda m: None)
    pdf = next(f for f in made if f.endswith("Print.pdf"))
    with pdfplumber.open(pdf) as pp:
        toc_page = next((pg.extract_text() or "" for pg in pp.pages[:6]
                         if (pg.extract_text() or "").lstrip().upper().startswith("CONTENTS")), "")
    entries = toc_page.split("\n")[1:3]
    check(all(w in toc_page for w in want) and not any(not_want in e for e in entries),
          f"Contents case {case}: {entries}")

# ---------------------------------------------------------------- 8. running heads on the Notes pages
# Bug: the lined Notes pages at the end carried the last chapter's title as their running head.
NB = "".join(f"\n\nChapter {n}: {t}\n\n" + "Plain words in a chapter. " * 400
             for n, t in ((1, "Morning Tide"), (2, "Lighthouse Keeper")))
s = dict(DEFAULTS)
s.update(out_dir=tempfile.mkdtemp(dir=TMP), source_mode="paste", paste_text=NB, open_when_done=False,
         fmt_impose=False, running_heads=True, notes_pages=2, toc=False)
book = apply_info(load_book(s, NB), s)
_f, made, _n = run_build(book, s, lambda m: None)
pdf = next(f for f in made if f.endswith("Print.pdf"))
with pdfplumber.open(pdf) as pp:
    last = [pg.extract_text() or "" for pg in pp.pages[-2:]]
    body = [pg.extract_text() or "" for pg in pp.pages[:-2]]
check(all("Notes" in x and "Lighthouse Keeper" not in x for x in last),
      f"Notes pages don't carry the last chapter's title {[x[:40] for x in last]}")
check(any("Lighthouse Keeper" in x for x in body[-3:]), "chapter pages still carry the chapter's running head")

# ---------------------------------------------------------------- 9. version number in one place
vi = open(os.path.join(ROOT, "version_info.txt"), encoding="utf-8").read()
parts = tuple(int(x) for x in APP_VERSION.split("."))
nums = re.findall(r"filevers=\((\d+), (\d+), (\d+), (\d+)\)", vi)
check(nums and tuple(int(x) for x in nums[0][:3]) == parts and f"'{APP_VERSION}.0'" in vi,
      f"version_info.txt matches APP_VERSION {APP_VERSION}")

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
