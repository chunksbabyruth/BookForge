"""Checks for complete front & back editing in 1.8.0 (plain Python - run: python tests/test_front_back_180.py).

Every page around the chapters can be edited in Specific Text Edit (title page, bookplate, dedication, extra front
page, work info, summary, author's notes, contents, text before chapter 1, THE END, end pages, Notes pages,
back-cover text). These checks cover:

1. Every page edited (its own words, and a picture on the pages that take pictures) reaches every output that can
   hold it - preview, Print PDF, imposed PDF, cover, EPUB, Word, web page, text - in the same order everywhere.
2. What a format can't hold is said in the printing notes (and the log), never left out silently.
3. The text before chapter 1 (read from the source, or edited) is in every format, after the contents.
4. Word: an edited contents is the edited entries (renamed, hidden, added, typed page numbers; real page numbers
   as fields pointing at the chapter headings). An unedited contents is Word's own contents field, as before.
5. EPUB: the extra front page goes before or after the contents as set (as in the PDF).
6. Save / reopen: every edited page and picture comes back, and makes the same outputs.
7. Making the book (or saving a project) never overwrites the book's source file.
8. An edited contents keeps following the chapters: moving chapters, chapters left out and put back, Undo all.
9. Window: typing on a page is kept when switching pages, chapters and tabs; Preview, Make the book and Save
   project use the latest typing; the "Text before chapter 1" button; Re-read source names the edited pages.
Window checks run only where a display is available. Nothing needs the internet.
"""
import base64
import copy
import hashlib
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="bf_frontback_")
os.environ["APPDATA"] = os.path.join(TMP, "appdata")
os.environ["HOME"] = os.path.join(TMP, "home")
os.environ["BOOKFORGE_NO_UPDATE_CHECK"] = "1"
os.makedirs(os.environ["APPDATA"], exist_ok=True)
os.makedirs(os.path.join(os.environ["HOME"], "Documents"), exist_ok=True)

import docx                                                           # noqa: E402
import pdfplumber                                                     # noqa: E402
from PIL import Image                                                 # noqa: E402

import bf_edit                                                        # noqa: E402
import bf_matter                                                      # noqa: E402
from bf_build import build_preview, run_build, format_limits          # noqa: E402
from bf_core import DEFAULTS, Book, Chapter, Block, apply_info, load_book, not_the_source  # noqa: E402
from bf_pdf import contents_entries, contents_rows, rows_to_edits, apply_contents_edits  # noqa: E402
import bf_pdf                                                         # noqa: E402

FAILS = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg, flush=True)
    if not cond:
        FAILS.append(msg)


def png(w, h, color):
    b = io.BytesIO()
    Image.new("RGB", (w, h), color).save(b, "PNG")
    return b.getvalue()


def md5(path):
    with open(path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


ALL_FORMATS = dict(fmt_pdf=True, fmt_impose=True, fmt_cover=True, fmt_epub=True, fmt_docx=True, fmt_html=True,
                   fmt_txt=True)


def settings(**kw):
    s = copy.deepcopy(DEFAULTS)
    s.update(open_when_done=False, out_dir=tempfile.mkdtemp(dir=TMP), source_mode="paste",
             page_size="Mass Market (4.25 x 6.87 in)", **ALL_FORMATS)
    s.update(kw)
    return s


def chars(path):
    with pdfplumber.open(path) as pdf:
        return [[(c["text"], round(c["x0"], 2), round(c["top"], 2), c["fontname"].split("+")[-1], round(c["size"], 2))
                 for c in pg.chars] + [("IMG", round(i["x0"], 2), round(i["top"], 2)) for i in pg.images]
                for pg in pdf.pages]


def page_texts(path):
    with pdfplumber.open(path) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def outputs(folder):
    """The words (and pictures) of every file made in folder, by format."""
    out = {}
    for f in sorted(os.listdir(folder)):
        p = os.path.join(folder, f)
        if f.endswith("Print.pdf"):
            with pdfplumber.open(p) as pdf:
                out["print"] = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
                out["print_imgs"] = sum(len(pg.images) for pg in pdf.pages)
            out["print"] = re.sub(r"( \.){2,} ", " ", out["print"])        # (the contents' dotted leaders)
            out["print_path"] = p
        elif "Imposed" in f and f.endswith(".pdf"):
            with pdfplumber.open(p) as pdf:
                out["imposed"] = re.sub(r"( \.){2,} ", " ", out.get("imposed", "") + "\n".join(
                    (pg.extract_text() or "") for pg in pdf.pages))
        elif f.endswith("Cover.pdf"):
            with pdfplumber.open(p) as pdf:
                out["cover"] = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
        elif f.endswith(".epub"):
            z = zipfile.ZipFile(p)
            opf = z.read("OEBPS/content.opf").decode()
            order = re.findall(r'itemref idref="([^"]+)"', opf)
            items = dict(re.findall(r'<item id="([^"]+)" href="([^"]+)"', opf))
            out["epub_spine"] = order
            out["epub_xhtml"] = "\n".join(z.read("OEBPS/" + items[i]).decode() for i in order)
            out["epub"] = re.sub(r"<[^>]+>", " ", out["epub_xhtml"])
            out["epub_media"] = [z.read(n) for n in z.namelist() if n.startswith("OEBPS/images/")]
        elif f.endswith(".docx"):
            d = docx.Document(p)
            out["docx_doc"] = d
            out["docx"] = "\n".join(x.text for x in d.paragraphs).replace("\t", " ")
            z = zipfile.ZipFile(p)
            out["docx_xml"] = z.read("word/document.xml").decode()
            out["docx_media"] = [z.read(n) for n in z.namelist() if n.startswith("word/media/")]
            out["docx_path"] = p
        elif f.endswith(".html"):
            t = open(p, encoding="utf-8").read()
            out["html_raw"] = t
            out["html"] = re.sub(r"<[^>]+>", " ", t)
        elif f.endswith(".txt") and not f.startswith("READ ME"):
            out["txt"] = open(p, encoding="utf-8").read()
        elif f.startswith("READ ME"):
            out["readme"] = open(p, encoding="utf-8").read()
    return out


def make(book, s):
    """Make the book (every chosen format) from a copy: (outputs, files made, printing notes)."""
    s = dict(s)
    s["out_dir"] = tempfile.mkdtemp(dir=TMP)
    folder, made, notes = run_build(apply_info(copy.deepcopy(book), dict(s)), s, log=lambda m: None)
    return outputs(folder), made, notes


# ===================================================================== 1. every page edited: every output
PICS = {"note": png(120, 80, (200, 0, 0)), "front_text": png(100, 90, (0, 150, 0)),
        "end_matter": png(90, 100, (0, 0, 200))}
MARK = {"title": "Zqtitleline", "bookplate": "Zqplate", "dedication": "Zqdedic", "note": "Zqextrapage",
        "work_info": "Zqworkinfo", "summary": "Zqsummary", "work_notes": "Zqworknotes", "contents": "Zqtocadded",
        "front_text": "Zqfronttext", "the_end": "Zqtheend", "end_matter": "Zqafterword", "notes_page": "Zqnotestext",
        "blurb": "Zqblurb"}
HEAD = {"work_info": "Zhworkinfo", "summary": "Zhsummary", "work_notes": "Zhworknotes", "end_matter": "Zhafter",
        "notes_page": "Zhnotes"}
ORIGINAL = ["Origfandom", "Origsummary", "Orignotes", "Origfront"]    # what the source said, replaced by the edits


def source_book():
    return Book(title="Front Back Book", subtitle="Sub", author="A. Writer",
                fic_meta=["<b>Rating:</b> G", "<b>Fandom:</b> Origfandom"],
                fic_summary=[Block("para", text="Origsummary words.", label="Summary")],
                fic_notes=[Block("para", text="Orignotes words.", label="Notes")],
                front_blocks=[Block("para", text="Origfront words.")],
                chapters=[Chapter(f"Chapter {i}", blocks=[Block("para", text=f"Text of chapter {i}. " * 60)])
                          for i in range(1, 4)])


S_ALL = settings(title_page=True, bookplate="Separate page", toc=True, the_end=True, notes_pages=1,
                 fic_info="Add a work info page", note_title="Zhextrapage", note_place="After contents",
                 dedication="x", note_text="x", blurb="x", toc_dots=True)


def edited_book():
    book = source_book()
    for key, word in MARK.items():
        if key == "contents":
            continue
        if key == "title":
            blocks = [Block("btitle", "Front Back Book"), Block("bsub", "Sub"), Block("bauth", "A. Writer"),
                      Block("para", f"<i>{word}</i> extra line")]
        elif key == "bookplate":
            blocks = [Block("plate", "This Book Belongs To"), Block("para", f"{word} line")]
        elif key == "the_end":
            blocks = [Block("endline", f"<b>{word}</b>")]
        else:
            kind = {"dedication": "ded", "work_info": "meta"}.get(key, "para")
            blocks = [Block(kind, f"<b>{word}</b> words of the page.")]
        if key in PICS:
            blocks.append(Block("image", image=PICS[key], image_name=key + ".png", img_width=40))
        bf_matter.set_override(book, key, blocks)
        if key in HEAD:
            bf_matter.set_heading(book, key, HEAD[key])
    # contents: chapter 1 gets a typed page number, chapter 2 renamed (bold), chapter 3 hidden, an entry added
    auto = contents_entries(book, S_ALL)
    rows = contents_rows(auto, [])
    rows[0]["page"] = "ix"
    rows[1]["title"] = "Zqtocrenamed"
    rows[1]["markup"] = "<b>Zqtocrenamed</b>"
    rows[2]["show"] = False
    rows.append({"uid": None, "custom": True, "auto": "", "title": "Zqtocadded", "page": "99", "show": True,
                 "level": 0, "auto_level": 0, "markup": None, "space": 0})
    book.toc = rows_to_edits(rows, auto)
    return book


MARK2 = dict(MARK, contents_renamed="Zqtocrenamed", contents_typed="Chapter 1 ix",
             **{"head_" + k: v for k, v in HEAD.items()}, head_note="Zhextrapage")
FORMATS = ["print", "imposed", "epub", "docx", "html", "txt"]
NOT_IN = {"bookplate": {"epub", "html", "txt"},                           # a print page (Word has it too)
          "notes_page": {"epub", "docx", "html", "txt"}, "head_notes_page": {"epub", "docx", "html", "txt"},
          "contents_typed": {"epub", "html", "txt"},                       # no page numbers: the text reflows
          "blurb": set(FORMATS)}                                           # the back cover only

BOOK1 = edited_book()
O1, MADE1, NOTES1 = make(BOOK1, S_ALL)
for key, word in MARK2.items():
    want = [f for f in FORMATS if f not in NOT_IN.get(key, set())]
    got = [f for f in FORMATS if word in O1.get(f, "")]
    check(got == want, f"{key}: the edit is in {', '.join(want) or 'none of the book files'}"
          + (f" (in: {', '.join(got) or 'none'})" if got != want else ""))
check("Zqblurb" in O1.get("cover", ""), "back-cover text: the edit is on the cover")
check(all(w not in O1.get(f, "") for f in FORMATS for w in ORIGINAL),
      "edited work info, summary, author's notes and text before chapter 1 replace what the source said everywhere")
check("Text of chapter 3" in O1.get("print", "") and "Text of chapter 3" in O1.get("epub", ""),
      "a chapter hidden on the contents page is still in the book")
seq = {f: [k for _p, k in sorted((O1.get(f, "").find(w), k) for k, w in MARK2.items() if w in O1.get(f, ""))]
       for f in ("print", "epub", "docx", "html", "txt")}
check(all(seq[f] == [k for k in seq["print"] if k in seq[f]] for f in seq),
      "every format has the pages in the print PDF's order " + " > ".join(seq["print"]))
pics = list(PICS.values())
check(O1.get("print_imgs") == 3 and all(p in O1.get("epub_media", []) for p in pics) and
      all(p in O1.get("docx_media", []) for p in pics) and
      all(base64.b64encode(p).decode() in O1.get("html_raw", "") for p in pics),
      "the pictures on the extra front page, text before chapter 1 and end pages: Print PDF, EPUB, Word, web page")
check("<b>Zqfronttext</b>" in O1.get("epub_xhtml", "") and "<b>Zqfronttext</b>" in O1.get("html_raw", "") and
      any(r.bold and "Zqfronttext" in r.text for p in O1["docx_doc"].paragraphs for r in p.runs),
      "the formatting of the edited text before chapter 1 (bold) is in the EPUB, Word and web page")
res = build_preview(apply_info(copy.deepcopy(BOOK1), dict(S_ALL)), dict(S_ALL), log=lambda m: None)
check(chars(res["pdf"]) == chars(O1["print_path"]), "the preview is the Print PDF, every edited page included")
check(res["sheets"] and all(w in "\n".join(t for f in res["sheets"] for t in page_texts(f))
                            for w in ("Zqfronttext", "Zqafterword", "Zqtocrenamed")),
      "the preview's imposed sheets have the edited pages")

# ===================================================================== 2. what a format can't hold is said
LIMITS = ["Not in the EPUB, web page and text file: the bookplate page you edited",
          "Not in the EPUB, Word document, web page and text file: the Notes pages you edited",
          "page numbers typed in Edit contents are only in the print PDF and the Word document",
          "fonts, bold, colours and space above are in the print PDF and the Word document",
          "The text file has no pictures"]
check(all(any(x in n for n in NOTES1) for x in LIMITS) and all(x in O1.get("readme", "") for x in LIMITS),
      "the printing notes say what the EPUB, web page, text file and Word can't hold")
check(format_limits(source_book(), S_ALL) == [], "nothing is said for a book whose pages aren't edited")
only_pdf = dict(S_ALL, fmt_epub=False, fmt_docx=False, fmt_html=False, fmt_txt=False)
check([n for n in format_limits(BOOK1, only_pdf)] == [], "nothing is said when only the PDFs are made")
o, _m, notes = make(BOOK1, dict(only_pdf, fmt_cover=False, fmt_impose=False))
check(any("back-cover text you edited is printed on the cover only" in n for n in notes),
      "edited back-cover text with the cover not being made: the notes say so")

# ===================================================================== 3. text before chapter 1 in every format
for label, kw in (("contents on", {}), ("no title page or contents", dict(title_page=False, toc=False))):
    o, _m, _n = make(source_book(), settings(**kw))
    pos = {f: (o.get(f, "").find("Origfront"), o.get(f, "").find("Text of chapter 1")) for f in FORMATS
           if f != "imposed"}                        # (the imposed sheets aren't in reading order)
    check(all(0 <= a < b for a, b in pos.values()) and "Origfront" in o.get("imposed", ""),
          f"text before chapter 1 read from the source: in every format, before chapter 1 ({label})")
check("front" in O1["epub_spine"] and O1["epub_spine"].index("nav") < O1["epub_spine"].index("front") <
      O1["epub_spine"].index("c1"), "EPUB: the text before chapter 1 is after the contents, before chapter 1")
b3 = source_book()
bf_matter.set_override(b3, "front_text", [])                      # edited to nothing
o, _m, _n = make(b3, settings())
check(all("Origfront" not in o.get(f, "") for f in FORMATS) and "front" not in o["epub_spine"],
      "text before chapter 1 emptied: in no format")
b3 = source_book()
bf_matter.set_override(b3, "front_text", [Block("image", image=PICS["front_text"], image_name="p.png")])
o, _m, _n = make(b3, settings())
check(o.get("print_imgs") == 1 and PICS["front_text"] in o["epub_media"] and PICS["front_text"] in o["docx_media"]
      and "Origfront" not in o.get("txt", ""), "text before chapter 1 made a picture only: the picture, everywhere")

# ===================================================================== 4. Word contents
o, _m, _n = make(source_book(), settings(fmt_pdf=False, fmt_impose=False, fmt_cover=False))
check('TOC \\o "1-2"' in o["docx_xml"] and "PAGEREF" not in o["docx_xml"] and "bf_" not in o["docx_xml"],
      "Word, contents not edited: Word's own contents field, as before")
paras = O1["docx_doc"].paragraphs
ci = next(n for n, p in enumerate(paras) if p.text.strip() == "CONTENTS")
entries = [p.text for p in paras[ci + 1:ci + 4]]
uid2 = BOOK1.chapters[1].uid
check(entries == ["Chapter 1\tix", "Zqtocrenamed\t", "Zqtocadded\t99"] and 'TOC \\o' not in O1["docx_xml"],
      f"Word, contents edited: the edited entries in order, typed numbers as typed, the hidden one left out {entries}")
check(f"PAGEREF bf_{uid2}" in O1["docx_xml"] and
      re.search(rf'<w:bookmarkStart w:id="\d+" w:name="bf_{uid2}"/>.*?Chapter 2.*?<w:bookmarkEnd', O1["docx_xml"]),
      "Word: the renamed entry's page number is a field that points at chapter 2's heading")
rp = paras[ci + 2]
check(any(r.bold and r.text == "Zqtocrenamed" for r in rp.runs) and rp.paragraph_format.tab_stops[0].leader == 1,
      "Word: an entry's formatting (bold) and the dotted leader")
if shutil.which("soffice"):
    lo = tempfile.mkdtemp(dir=TMP)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", lo, O1["docx_path"]],
                       capture_output=True, timeout=180)
        pdfs = [os.path.join(lo, f) for f in os.listdir(lo) if f.endswith(".pdf")]
    except (OSError, subprocess.SubprocessError):
        pdfs = []
    if pdfs:
        pts = page_texts(pdfs[0])
        cp = next((t for t in pts if "Zqtocrenamed" in t), "")
        start = next((n for n, t in enumerate(pts, 1) if "Text of chapter 2" in t), None)
        m = re.search(r"Zqtocrenamed\D*(\d+)", cp)
        check(m and int(m.group(1)) == start, f"Word (opened in LibreOffice): the renamed entry shows chapter 2's "
              f"real page ({m.group(1) if m else None} = {start})")
    else:
        print("SKIP  Word contents opened in LibreOffice (it couldn't convert the file here)")
else:
    print("SKIP  Word contents opened in LibreOffice (LibreOffice isn't installed here)")

# ===================================================================== 5. EPUB extra front page placement
for place in ("Before contents", "After contents"):
    o, _m, _n = make(source_book(), settings(note_title="Zhplace", note_text="Place words.", note_place=place,
                                             fmt_impose=False, fmt_cover=False, fmt_docx=False))
    sp = o["epub_spine"]
    pt = o["print"]
    pdf_before = pt.find("Zhplace") < pt.find("CONTENTS")
    check((sp.index("note") < sp.index("nav")) == (place == "Before contents") == pdf_before and
          (o["html"].find("Zhplace") < o["html"].find("CONTENTS")) == pdf_before and
          (o["txt"].find("Zhplace") < o["txt"].find("CONTENTS")) == pdf_before,
          f"extra front page {place.lower()}: EPUB, web page and text file as in the PDF")

# ===================================================================== 6. save / reopen
proj = os.path.join(TMP, "front back.bookforge")
bf_edit.save_project(proj, BOOK1, S_ALL, "1.8.0")
b6, s6 = bf_edit.load_project(proj)
sig = lambda b: {k: [(x.kind, x.text, x.image, x.lines) for x in (bf_matter.override(b, k) or [])]   # noqa: E731
                 for k in bf_matter.SECTIONS}
check(sig(b6) == sig(BOOK1) and all(bf_matter.heading_override(b6, k) == bf_matter.heading_override(BOOK1, k)
                                    for k in bf_matter.SECTIONS) and b6.toc == BOOK1.toc,
      "save / reopen: every edited page, picture, heading and the contents edits come back")
o6, _m, _n = make(b6, s6)
check(chars(o6["print_path"]) == chars(O1["print_path"]) and
      all(o6.get(f) == O1.get(f) for f in ("imposed", "cover", "epub", "docx", "html", "txt")) and
      o6["epub_media"] == O1["epub_media"] and o6["docx_media"] == O1["docx_media"],
      "the reopened project makes the same Print PDF, imposed PDF, cover and exports")

# ===================================================================== 7. the source is never overwritten
check(not_the_source("x.txt", {"source_mode": "paste"}) == "x.txt", "pasted text: files keep their names")
books = os.path.join(TMP, "Books")
os.makedirs(os.path.join(books, "My Book"))
src = os.path.join(books, "My Book", "My Book.txt")
with open(src, "w", encoding="utf-8") as f:
    f.write("Chapter 1\n\n" + "The original words of my book. " * 30 + "\n\nChapter 2\n\nMore original words.\n")
s7 = settings(source_mode="file", source_path=src, out_dir=books, title="My Book", fmt_impose=False, fmt_cover=False)
h0 = md5(src)
check(not_the_source(os.path.join(books, "My Book", ".", "My Book.txt"), s7) != src and
      not_the_source(os.path.join(books, "My Book", "Other.txt"), s7).endswith("Other.txt"),
      "the source's own path (however written) is never used for a file BookForge writes")
b7 = apply_info(load_book(s7), s7)
bf_matter.set_override(b7, "end_matter", [Block("para", "Words I added.")])
folder, made, notes = run_build(b7, dict(s7), log=lambda m: None)
alt = os.path.join(books, "My Book", "My Book (BookForge).txt")
check(md5(src) == h0 and alt in made and os.path.isfile(alt) and "Words I added." in open(alt, encoding="utf-8").read(),
      "a text source kept in the book's folder: unchanged; the text file is saved beside it")
check(any("is the book's source file" in n for n in notes), "the printing notes say why the text file was renamed")
src2 = os.path.join(books, "My Book", "My Book - Print.pdf")
h2 = md5(src2)
s7b = dict(s7, source_path=src2, pdf_mode="Keep the pages as they are", fmt_txt=False, fmt_epub=False,
           fmt_docx=False, fmt_html=False, fmt_impose=True)
b7b = apply_info(load_book(s7b), s7b)
folder, made, notes = run_build(b7b, dict(s7b), log=lambda m: None)
check(md5(src2) == h2 and any(m.endswith("Print (BookForge).pdf") for m in made) and
      any("Imposed" in os.path.basename(m) for m in made),
      "BookForge's own Print PDF used again as the source: unchanged; the new files are saved beside it")

# ===================================================================== 8. contents follow the chapters
e = [{"uid": "a", "page": "3"}, {"uid": None, "custom": True, "title": "Map"}, {"uid": "b", "title": "B!"}]
sw = bf_pdf.swap_contents_entries(e, "a", "b")
check([x.get("uid") for x in sw] == ["b", None, "a"] and sw[2]["page"] == "3" and e[0]["uid"] == "a",
      "chapters swapped: their contents entries swap, an added entry stays where it is")
bk = Book(title="K", chapters=[Chapter(t) for t in "ABC"])
ua, ub, uc = (c.uid for c in bk.chapters)
bk.chapters[1].include = False
old = [{"uid": ua}, {"uid": ub, "page": "7"}, {"uid": uc, "title": "Cee"}]
kept = bf_pdf.keep_left_out([{"uid": ua, "title": "Aye"}, {"uid": uc, "title": "Cee"}], old, bk)
check([x["uid"] for x in kept] == [ua, ub, uc] and kept[1]["page"] == "7" and kept[0]["title"] == "Aye",
      "contents edited while a chapter is left out: that chapter's entry (typed page number) is kept in its place")

# ===================================================================== 9. window checks
try:
    import tkinter as tk
    _probe = tk.Tk()
    _probe.destroy()
    have_display = True
except Exception as ex:                    # noqa: BLE001
    have_display = False
    print(f"SKIP  window checks (no display: {ex})")

if have_display:
    from tkinter import messagebox, filedialog
    SAID = []
    ANSWER = {}
    for k in ("askyesno", "askyesnocancel", "showinfo", "showerror", "showwarning", "askokcancel"):
        setattr(messagebox, k, (lambda kind: (lambda title="", message="", **kw: SAID.append((kind, title, message))
                                              or ANSWER.get(kind, True)))(k))
    PICK = []
    filedialog.askopenfilename = lambda **kw: PICK.pop(0) if PICK else ""
    filedialog.asksaveasfilename = lambda **kw: PICK.pop(0) if PICK else ""
    import BookForge as BF
    import bf_theme as th
    app = BF.App()
    app._no_save = True
    errors = []
    app.report_callback_exception = lambda *a: errors.append("".join(__import__("traceback").format_exception(*a)))
    ed, t = app.editor, app.editor.txt

    def pump(sec=0.2):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            time.sleep(0.01)

    def wait(cond, sec=180):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            if cond():
                return True
            time.sleep(0.02)
        return False

    def idle(sec=240):
        return wait(lambda: not app.busy, sec)

    def tab(key):
        app.shelf.select(th.tab_index(key))
        pump(0.15)

    def preview():
        idle()
        app.preview_book()
        idle()
        pump(0.2)
        pw = getattr(app, "preview_win", None)
        return pw.result if pw is not None and pw.alive() else None

    def shown():
        return t.get("1.0", "end-1c")

    def typing(text):
        """Type at the insert point as a person does (key events)."""
        t.focus_force()
        for ch in text:
            ks = {" ": "space", ".": "period"}.get(ch, ch)
            t.event_generate("<KeyPress>", keysym=ks)
            t.event_generate("<KeyRelease>", keysym=ks)
        pump(0.05)

    def words_of(key):
        return " ".join(b.text or "" for b in (bf_matter.override(app.book, key) or []))

    def pdf_text(path):
        return "\n".join(page_texts(path)) if path else ""

    TEXT = ("A few words before the story starts.\n\n" +
            "".join(f"Chapter {i}\n\n" + "\n\n".join(f"Paragraph {i}.{j}: " + "words of text here " * 10
                                                       for j in range(1, 3)) + "\n\n" for i in range(1, 5)))
    srcfile = os.path.join(TMP, "window source.txt")
    with open(srcfile, "w", encoding="utf-8") as f:
        f.write(TEXT)
    src_hash = md5(srcfile)
    app.var("source_mode").set("file")
    app._mode_changed()
    app.var("source_path").set(srcfile)
    app.var("title").set("Window Book")
    app.var("keep_front_text").set(True)
    app.var("toc").set(True)
    app.var("page_size").set("Mass Market (4.25 x 6.87 in)")
    idle()
    app.scan(force=True)
    wait(lambda: not app.busy and app.book is not None, 120)
    pump(0.2)
    check(app.book is not None and app.book.front_blocks and len(app.book.chapters) == 4,
          "a source with text before chapter 1 is read")

    # ---- the Front & back button for the text before chapter 1
    tab("front")
    btn = app.__dict__.get("matter_states", {}).get("front_text")
    check(btn is not None and str(btn.cget("text")).startswith("Text before chapter 1"),
          "Front & back has a “Text before chapter 1” button with the other pages")
    if btn is not None:
        btn.invoke()
    else:                                           # (so the checks below still run)
        app.edit_section("front_text")
    pump(0.3)
    check(th.TABS[app.cur_page][0] == "edit" and ed.sec == "front_text" and
          shown().startswith("A few words before the story starts."),
          "it opens Specific Text Edit at that page, showing the words read from the source")

    # ---- typing is kept when switching pages, chapters and tabs (each in the right place)
    t.mark_set("insert", "1.end")
    typing(" Typed here.")
    ed.select_section("end_matter")
    pump(0.05)
    check("Typed here." in words_of("front_text") and bf_matter.override(app.book, "end_matter") is None and
          ed.sec == "end_matter" and "Typed here." not in shown(),
          "typing on a page, then opening another page: kept on the first page only")
    t.insert("1.0", "An afterword typed in the window.")
    first_ch = [b.text for b in app.book.chapters[0].blocks]
    ed.select(0)
    pump(0.05)
    em = bf_matter.override(app.book, "end_matter")
    check(em and em[0].text == "An afterword typed in the window." and
          [b.text for b in app.book.chapters[0].blocks] == first_ch,
          "typing on a page, then opening a chapter: kept on the page; the chapter is unchanged")
    ed.select_section("the_end")
    pump(0.05)
    t.delete("1.0", "end")
    t.insert("1.0", "Closing words")
    tab("out")
    check(words_of("the_end") == "Closing words",
          "typing on a page, then going to another step: kept")
    tab("edit")
    check(ed.sec == "the_end" and shown() == "Closing words", "back in Specific Text Edit: the same page and words")
    for k in ("src", "info", "front", "chap", "lay", "out", "edit"):
        tab(k)
    check(words_of("the_end") == "Closing words" and "Typed here." in words_of("front_text"),
          "the edited pages are kept through every step")

    # ---- Preview / Make the book use what was just typed (not yet saved by switching)
    ed.select_section("front_text")
    pump(0.05)
    t.mark_set("insert", "1.end")
    typing(" Zqseen.")
    res = preview()
    check(res and "Zqseen" in pdf_text(res["pdf"]), "Preview right after typing: the typed words are in it")
    pic = os.path.join(TMP, "green.png")
    Image.new("RGB", (160, 100), (20, 160, 40)).save(pic)
    t.mark_set("insert", "end")
    PICK.append(pic)
    ed.insert_picture()
    ed.select_section("end_matter")
    pump(0.05)
    t.insert("end", "\n\nZqmade")
    outdir = tempfile.mkdtemp(dir=TMP)
    app.var("out_dir").set(outdir)
    app.var("open_when_done").set(False)
    for k, v in dict(ALL_FORMATS, fmt_cover=False).items():
        app.var(k).set(v)
    idle()
    app.build()
    idle()
    pump(0.2)
    made = [r for r, _d, fs in os.walk(outdir) for f in fs if f.endswith("Print.pdf")]
    o9 = outputs(made[0]) if made else {}
    check(o9 and all("Zqmade" in o9.get(f, "") and "Zqseen" in o9.get(f, "")
                     for f in FORMATS) and open(pic, "rb").read() in o9.get("epub_media", []),
          "Make the book right after typing: the typed words in every file, the picture in the EPUB")
    check(md5(srcfile) == src_hash, "the source file is unchanged after making the book")

    # ---- Save project right after typing; reopen
    t.insert("end", " Zqsaved")
    proj9 = os.path.join(TMP, "window.bookforge")
    PICK.append(srcfile)                            # the source file picked by mistake: refused
    n = len(SAID)
    app.save_project_dialog()
    check(md5(srcfile) == src_hash and any(k == "showerror" and "source" in m for k, _t, m in SAID[n:]),
          "Save project onto the source file: refused (the source is never replaced)")
    PICK.append(proj9)
    idle()
    app.save_project_dialog()
    check(not app.unsaved, "the project is saved")
    want = sig(app.book)
    PICK.append(proj9)
    ANSWER["askyesno"] = True
    idle()
    app.open_project_dialog()
    pump(0.3)
    em = bf_matter.override(app.book, "end_matter") or []
    ft = bf_matter.override(app.book, "front_text") or []
    check(sig(app.book) == want and any("Zqsaved" in (b.text or "") for b in em) and
          any(b.kind == "image" and b.image == open(pic, "rb").read() for b in ft),
          "reopened: the pages typed on (and the picture) are back")
    res = preview()
    check(res and "Zqsaved" in pdf_text(res["pdf"]) and "Zqseen" in pdf_text(res["pdf"])
          and sum(len(p.images) for p in pdfplumber.open(res["pdf"]).pages) == 1,
          "the reopened project previews with the edited pages and the picture")

    # ---- the contents follow the chapters on the Chapters step
    tab("chap")
    app.edit_contents()
    pump(0.1)
    cw = app.contents_win
    cw.rows[1]["title"] = "Renamed Two"
    cw.rows[2]["page"] = "33"
    cw.ok()
    pump(0.1)

    def order():
        return [x["text"] for x in apply_contents_edits(contents_entries(app.book, app.collect()), app.book.toc)]

    def typed(title):
        uid = next(c.uid for c in app.book.chapters if c.title == title)
        return next((x.get("page") for x in app.book.toc if x.get("uid") == uid), None)
    app.fill_list()
    app.ctree.selection_set("3")
    app.move_ch(-1)                                 # Chapter 4 above Chapter 3
    pump(0.1)
    check([c.title for c in app.book.chapters] == ["Chapter 1", "Chapter 2", "Chapter 4", "Chapter 3"] and
          order() == ["Chapter 1", "Renamed Two", "Chapter 4", "Chapter 3"] and typed("Chapter 3") == "33",
          "moving a chapter: the edited contents follow the new order (renamed words and typed numbers stay)")
    app.edit_contents()
    pump(0.1)
    cw = app.contents_win
    u4 = next(c.uid for c in app.book.chapters if c.title == "Chapter 4")
    next(r for r in cw.rows if r["uid"] == u4)["title"] = "Renamed Four"     # typed, not OK'd yet
    app.fill_list()
    app.ctree.selection_set("2")
    app.move_ch(1)                                  # Chapter 4 back below Chapter 3
    pump(0.1)
    in_window = [r["title"] or r["auto"] for r in cw.rows]
    cw.ok()
    pump(0.1)
    check([c.title for c in app.book.chapters] == ["Chapter 1", "Chapter 2", "Chapter 3", "Chapter 4"] and
          order() == ["Chapter 1", "Renamed Two", "Chapter 3", "Renamed Four"] and in_window == order() and
          typed("Chapter 3") == "33",
          "moving a chapter while Edit contents is open: the window and the contents follow, its typed change kept")
    tab("edit")
    ed.select(3)                                    # Chapter 4 in Specific Text Edit...
    ed.move_chapter(-1)                             # ...Move up
    pump(0.1)
    check([c.title for c in app.book.chapters] == ["Chapter 1", "Chapter 2", "Chapter 4", "Chapter 3"] and
          order() == ["Chapter 1", "Renamed Two", "Renamed Four", "Chapter 3"],
          "Move up on a chapter in Specific Text Edit: the edited contents follow too")
    i3 = next(k for k, c in enumerate(app.book.chapters) if c.title == "Chapter 3")
    app.fill_list()
    app.ctree.selection_set(str(i3))
    app.toggle_ch()                                 # left out for now
    pump(0.1)
    app.edit_contents()
    pump(0.1)
    cw = app.contents_win
    cw.rows[0]["title"] = "Renamed One"
    cw.ok()
    pump(0.1)
    tab("edit")
    ed.select_section("contents")
    pump(0.05)
    t.insert("1.0", "Opening: ")
    ed.commit()
    app.fill_list()
    app.ctree.selection_set(str(i3))
    app.toggle_ch()                                 # put back
    pump(0.1)
    check(typed("Chapter 3") == "33" and order()[0] == "Opening: Renamed One",
          "a chapter left out while the contents are edited (Edit contents and the text) keeps its typed number")
    res = preview()
    cp = next((p for p in page_texts(res["pdf"]) if "Renamed Two" in p), "") if res else ""
    check(re.search(r"Chapter 3[ .]*33", cp), "the preview prints it")
    app.edit_contents()
    pump(0.1)
    cw = app.contents_win
    ANSWER["askyesno"] = True
    cw.reset()
    cw.ok()
    pump(0.1)
    check(app.book.toc == [], "Undo all changes: the automatic contents again")

    # ---- Re-read source names the edited pages; the user's own pages stay
    idle()
    lines = app.changes_text()
    check("Text before chapter 1" in lines and "kept" in lines,
          "Re-read source warns that the edited text before chapter 1 would be replaced")
    ANSWER["askyesno"] = True
    app.reread_source()
    idle(60)
    pump(0.2)
    check(not bf_matter.is_edited(app.book, "front_text") and bf_matter.is_edited(app.book, "end_matter") and
          bf_matter.is_edited(app.book, "the_end") and md5(srcfile) == src_hash,
          "after Re-read source: the source's text before chapter 1 again; the end pages and THE END are kept")
    check(not errors, f"no window errors {errors[:1]}")
    app.destroy()

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
