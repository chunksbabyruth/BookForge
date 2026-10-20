"""Checks for BookForge 1.6.0 (plain Python - run:  python tests/test_regressions_160.py).

Covers: page numbers on every contents page and automatic contents spacing, editing the contents,
pictures (size, alignment, own page, full page) and cover pictures, PDF files kept as they are or rebuilt
as editable books, keeping the source's own look, the dark-mode dropdown colours and the update check.
Every file used is generated here and nothing is fetched from the internet, so this also runs on GitHub
Actions.
"""
import difflib
import io
import json
import os
import re
import sys
import tempfile
import threading
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix="bf_regress160_")
os.environ["APPDATA"] = os.path.join(TMP, "appdata")          # keep the user's real settings out of it
os.environ["HOME"] = os.path.join(TMP, "home")
os.environ["BOOKFORGE_NO_UPDATE_CHECK"] = "1"
os.makedirs(os.environ["APPDATA"], exist_ok=True)
os.makedirs(os.environ["HOME"], exist_ok=True)

import pdfplumber                                                     # noqa: E402
from PIL import Image                                                 # noqa: E402
from reportlab.pdfgen import canvas                                   # noqa: E402
import bf_impose  # noqa: E402,F401  (registers binding defaults)
from bf_core import DEFAULTS, APP_VERSION, Block, load_book, apply_info, plain   # noqa: E402
from bf_pdf import build_pdf, contents_entries, contents_rows, rows_to_edits   # noqa: E402
from bf_export import build_epub, build_html, build_txt, build_docx, build_cover   # noqa: E402
from bf_build import run_build                                        # noqa: E402
import bf_edit                                                        # noqa: E402
import bf_pictures                                                    # noqa: E402
import bf_update                                                      # noqa: E402

FAILS = []
HALF = "Half Letter / Digest (5.5 x 8.5 in)"


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


def settings(**kw):
    s = dict(DEFAULTS)
    s.update(open_when_done=False, out_dir=tempfile.mkdtemp(dir=TMP))
    s.update(kw)
    return s


def chapters_text(n, words=40):
    return "".join(f"\n\nChapter {k}\n\n" + ("Words in this chapter go on and on for a while. " * (words + (k * 37) % 300))
                   for k in range(1, n + 1))


def pdf_pages(path):
    with pdfplumber.open(path) as pdf:
        out = []
        for pg in pdf.pages:
            words = pg.extract_words()
            foot = [w["text"] for w in words if w["top"] > pg.height - 60]
            out.append({"text": pg.extract_text() or "", "foot": foot[-1] if foot else "", "images": pg.images,
                        "w": pg.width, "h": pg.height, "chars": pg.chars})
        return out


def picture(w, h, color=(40, 90, 160), fmt="PNG", **save):
    im = Image.new("RGB", (w, h), color)
    b = io.BytesIO()
    im.save(b, fmt, **save)
    return b.getvalue()


# ---------------------------------------------------------------- 1. contents: page numbers on every page, spacing
# Bug: on contents pages after the first, the page number was missing (those pages were treated as blank),
# and 41 chapters at Half Letter left the last few entries alone on a third page.
def contents_run(n, **kw):
    s = settings(source_mode="paste", toc=True, toc_numbers="Number (1.)", page_size=HALF, fmt_impose=False, **kw)
    book = apply_info(load_book(s, chapters_text(n)), s)
    out = os.path.join(TMP, f"toc{n}_{len(os.listdir(TMP))}.pdf")
    logs = []
    build_pdf(book, s, out, log=logs.append)
    pages = pdf_pages(out)
    toc = [i for i, p in enumerate(pages) if re.search(r"^\d+\. Chapter \d+", p["text"], re.M) and
           not re.match(r"^Chapter \d+$", p["text"].split("\n")[0].strip())]
    counts = [len(re.findall(r"^\d+\. Chapter \d+", pages[i]["text"], re.M)) for i in toc]
    listed = {}
    for i in toc:
        for m in re.finditer(r"^(\d+)\. Chapter \d+.*?(\S+)$", pages[i]["text"], re.M):
            listed[int(m.group(1))] = m.group(2)
    starts = {}
    for i, p in enumerate(pages):
        if i in toc:
            continue
        m = re.search(r"^Chapter (\d+)$", p["text"], re.M)
        if m and int(m.group(1)) not in starts:
            starts[int(m.group(1))] = p["foot"]
    return book, pages, toc, counts, listed, starts, logs


book41, pages, toc, counts, listed, starts, logs = contents_run(41)
check(len(toc) == 2 and min(counts) > 3, f"41 chapters: contents fits 2 pages, no nearly empty last page {counts}")
check(all(pages[i]["foot"] for i in toc), f"every contents page has its page number {[pages[i]['foot'] for i in toc]}")
check(listed and all(listed.get(k) == starts.get(k) for k in starts),
      "contents page numbers = the numbers printed on each chapter's first page")
check(any("Contents: spacing tightened" in m for m in logs), "the log says the spacing was tightened")
_b, pages, toc, counts, _l, _s, _g = contents_run(41, toc_spacing="Normal")
check(len(toc) == 3 and all(pages[i]["foot"] for i in toc),
      f"Normal spacing: 3 contents pages, all numbered {[pages[i]['foot'] for i in toc]}")
_b, pages, toc, counts, listed, starts, _g = contents_run(80)
check(all(pages[i]["foot"] for i in toc) and all(listed.get(k) == starts.get(k) for k in starts),
      f"80 chapters: {len(toc)} contents pages, all numbered and correct")

# ---------------------------------------------------------------- 2. editing the contents
s = settings(source_mode="paste", toc=True, toc_numbers="Number (1.)", page_size=HALF, fmt_impose=False)
book = apply_info(load_book(s, chapters_text(6)), s)
auto = contents_entries(book, s)
rows = contents_rows(auto, [])
rows[0]["title"] = "Opening Night"
rows[2]["show"] = False
rows[4], rows[5] = rows[5], rows[4]
rows[3]["page"] = "99"
rows.append({"uid": None, "custom": True, "auto": "", "title": "Map of the island", "page": "", "show": True,
             "level": 0, "auto_level": 0})
book.toc = rows_to_edits(rows, auto)
check(rows_to_edits(contents_rows(auto, []), auto) == [], "no changes = automatic contents (nothing saved)")
out = os.path.join(TMP, "edited.pdf")
logs = []
build_pdf(book, s, out, log=logs.append)
txt = "\n".join(p["text"] for p in pdf_pages(out)[:3])
check("Opening Night" in txt and "Map of the island" in txt and not re.search(r"^3\. Chapter 3\b", txt, re.M),
      "renamed, added and hidden entries printed as edited")
check(re.search(r"^4\. Chapter 4 99$", txt, re.M) and txt.index("6. Chapter 6") < txt.index("5. Chapter 5"),
      "typed page number printed; moved entry in its new place")
check(any("you typed page 99" in m for m in logs), "a typed page number that doesn't match is reported in the log")
pj = os.path.join(TMP, "toc.bookforge")
bf_edit.save_project(pj, book, s, APP_VERSION)
b2, _s2 = bf_edit.load_project(pj)
check(b2.toc == book.toc and [c.uid for c in b2.chapters] == [c.uid for c in book.chapters],
      "contents changes and chapter ids are saved in projects")
build_epub(book, s, os.path.join(TMP, "e.epub"))
nav = zipfile.ZipFile(os.path.join(TMP, "e.epub")).read("OEBPS/nav.xhtml").decode()
build_txt(book, s, os.path.join(TMP, "e.txt"))
tx = open(os.path.join(TMP, "e.txt"), encoding="utf-8").read()
check("Opening Night" in nav and "Map of the island" in nav and ">Chapter 3<" not in nav, "EPUB contents follows the edits")
check("Opening Night" in tx and "  Chapter 3\n" not in tx, "plain-text contents follows the edits")

# ---------------------------------------------------------------- 3. pictures in the book
P_WIDE, P_TALL = picture(800, 400), picture(1200, 1800, (200, 60, 40), "JPEG", quality=90)
s = settings(source_mode="paste", toc=False, title_page=False, bookplate="None", fmt_impose=False)
book = apply_info(load_book(s, chapters_text(1, 400)), s)
blocks = book.chapters[0].blocks
blocks.insert(1, Block("image", image=P_WIDE, image_name="wide.png", img_width=40, align="right"))
blocks.insert(2, Block("image", image=P_TALL, image_name="tall.jpg", img_place="page"))
blocks.insert(3, Block("image", image=P_TALL, image_name="tall.jpg", img_place="full"))
out = os.path.join(TMP, "pics.pdf")
build_pdf(book, s, out, log=lambda m: None)
pp = pdf_pages(out)
imgs = [(i, im) for i, p in enumerate(pp) for im in p["images"]]
W = pp[0]["w"]
fw = W - 72 * float(s["margin_inner"]) - 72 * float(s["margin_outer"])
small = [(i, im) for i, im in imgs if abs((im["x1"] - im["x0"]) - fw * 0.4) < 2]
check(small and small[0][1]["x1"] > W - 72 * 1.2, "40% picture: 40% of the text width, on the right")
own = [i for i, im in imgs if (im["x1"] - im["x0"]) < W - 10 and (im["bottom"] - im["top"]) > pp[i]["h"] * 0.6]
full = [i for i, im in imgs if im["x0"] <= 0.5 and im["x1"] >= W - 0.5]
check(own and pp[own[0]]["text"].strip().isdigit(), "picture on its own page: only the page number with it")
check(full and not pp[full[0]]["text"].strip(), "full-page picture: covers the page, no page number or heading")
check(own and full and own[0] < full[0], "pictures stay in their places in the book")
html_out = os.path.join(TMP, "pics.html")
build_html(book, s, html_out)
h = open(html_out, encoding="utf-8").read()
check('style="text-align:right"' in h and "width:40%" in h and "page-break-before:always" in h,
      "web page keeps picture size, alignment and placement")
dx = os.path.join(TMP, "pics.docx")
build_docx(book, s, dx)
import docx                                                           # noqa: E402
widths = [round(sh.width.inches, 2) for sh in docx.Document(dx).inline_shapes]
check(widths and abs(widths[0] - (8.5 - 2) * 0.4) < 0.05, f"Word file keeps the 40% width {widths}")
pj = os.path.join(TMP, "pics.bookforge")
bf_edit.save_project(pj, book, s, APP_VERSION)
b2, _ = bf_edit.load_project(pj)
kept = [(b.img_width, b.img_place, b.align, b.image == P_WIDE or b.image == P_TALL)
        for b in b2.chapters[0].blocks if b.kind == "image"]
check(kept == [(40, "", "right", True), (0, "page", "", True), (0, "full", "", True)],
      f"projects keep pictures and their settings {kept}")
# picture files: other formats are converted, phone photos turned upright
bmp = os.path.join(TMP, "p.bmp")
Image.new("RGB", (300, 200), (10, 120, 40)).save(bmp)
d, n = bf_pictures.read_picture(bmp)
check(bf_pictures.kind(d) == "png" and n == "p.png", "BMP becomes PNG")
ex = Image.Exif()
ex[0x0112] = 6                                           # "turn 90 degrees to view"
jp = os.path.join(TMP, "phone.jpg")
Image.new("RGB", (600, 300), (90, 90, 200)).save(jp, exif=ex.tobytes())
d, n = bf_pictures.read_picture(jp)
check(bf_pictures.size(d) == (300, 600), "a phone photo's rotation is applied (no sideways pictures)")
try:
    bf_pictures.read_picture(os.path.join(ROOT, "README.md"))
    check(False, "a non-picture file is refused")
except ValueError as e:
    check("isn't a picture" in str(e), f"a non-picture file is refused: {e}")

# ---------------------------------------------------------------- 4. cover pictures and spine width
front, back = os.path.join(TMP, "front.jpg"), os.path.join(TMP, "back.png")
Image.new("RGB", (1200, 1800), (30, 60, 150)).save(front, quality=90)
Image.new("RGB", (900, 1400), (150, 30, 30)).save(back)
s = settings(source_mode="paste", fmt_impose=False, page_size=HALF, cover_front_image=front, cover_back_image=back,
             spine_width=0.75)
book = apply_info(load_book(s, chapters_text(3)), s)
book.title = "Blue Cover Book"
cv = os.path.join(TMP, "cover.pdf")
info = build_cover(book, s, cv, 120)
cp = pdf_pages(cv)[0]
check(abs(info["width"] - (5.5 * 2 + 0.75)) < 0.01 and info["measured"], "measured spine width is used")
check(len(cp["images"]) == 2, "front and back pictures on the cover")
panel = (5.5 + 0.75) * 72
fx = [im for im in cp["images"] if im["x1"] > panel + 10]
check(fx and fx[0]["x0"] <= panel + 1 and fx[0]["x1"] >= cp["w"] - 1 and fx[0]["top"] <= 1 and
      fx[0]["bottom"] >= cp["h"] - 1, "front picture fills the whole front panel")
on_front = lambda pg: [ch for ch in pg["chars"] if ch["x0"] > panel + 2 and ch["text"].strip()]   # noqa: E731
check(not on_front(cp), "no title printed over the front picture (it usually has its own)")
s2 = dict(s, cover_text_over_art=True)
build_cover(book, s2, cv, 120)
check(len(on_front(pdf_pages(cv)[0])) >= len("BlueCoverBook"), "title printed over the picture when asked")
s3 = dict(s, cover_art_in_book=True, title_page=True)
out = os.path.join(TMP, "artbook.pdf")
build_pdf(book, s3, out, log=lambda m: None)
pp = pdf_pages(out)
check(pp[0]["images"] and not pp[0]["text"].strip() and not pp[1]["text"].strip() and "Blue Cover Book" in pp[2]["text"],
      "front picture as the first page, its back blank, then the title page")
book.cover_front = open(front, "rb").read()
s4 = dict(s, cover_front_image=os.path.join(TMP, "moved-away.jpg"))
check(bf_pictures.cover_art(book, s4, "front") == book.cover_front, "a moved cover file falls back to the project's copy")

# ---------------------------------------------------------------- 5. PDF: keep the pages as they are
from pypdf import PdfReader, PdfWriter                                # noqa: E402
from pypdf.generic import RectangleObject                             # noqa: E402
raw = os.path.join(TMP, "raw.pdf")
c = canvas.Canvas(raw)
for i, (w, h) in enumerate([(432, 648), (432, 648), (612, 792), (648, 432), (432, 648), (432, 648)]):
    c.setPageSize((w, h))
    c.setFont("Helvetica-Bold", 28)
    c.drawString(40, h - 80, f"PAGE {i + 1}")
    c.showPage()
c.save()
r = PdfReader(raw)
wr = PdfWriter()
for i, pg in enumerate(r.pages):
    if i == 1:
        pg.rotate(90)
    if i == 4:
        pg.cropbox = RectangleObject([36, 36, 396, 612])
    wr.add_page(pg)
tricky = os.path.join(TMP, "tricky.pdf")
wr.write(tricky)
s = settings(source_mode="file", source_path=tricky, fmt_cover=True, fmt_epub=True)
check(s["pdf_mode"] == "Keep the pages as they are", "PDFs keep their pages by default")
book = apply_info(load_book(s), s)
check(book.fixed_pdf and book.fixed_pages == 6 and not book.chapters, "a PDF kept as it is has no chapters")
logs = []
folder, made, _n = run_build(book, s, logs.append)
names = [os.path.basename(m) for m in made]
check(any(n.endswith("Print.pdf") for n in names) and any("Imposed" in n for n in names) and
      any("Cover" in n for n in names) and not any(n.endswith(".epub") for n in names),
      f"print, imposed and cover PDFs made; EPUB skipped {names}")
pp = pdf_pages(next(m for m in made if m.endswith("Print.pdf")))
check(len(pp) == 6 and all((p["w"], p["h"]) == (432, 648) for p in pp), "all pages one size, nothing lost")
import pypdfium2 as pdfium                                            # noqa: E402
pdoc = pdfium.PdfDocument(next(m for m in made if m.endswith("Print.pdf")))
seen = [re.sub(r"\s+", " ", pdoc[i].get_textpage().get_text_range()).strip() for i in range(len(pdoc))]
check(seen == [f"PAGE {i}" for i in range(1, 7)], f"the pages keep their content and order {seen}")
pdoc.close()
check(any("scaled to fit" in m for m in logs) and any("skipped" in m for m in logs),
      "the log explains refitted pages and the skipped EPUB")
scan = os.path.join(TMP, "scan.pdf")
ims = [Image.new("RGB", (510, 660), (250, 250, 245)) for _ in range(3)]
ims[0].save(scan, save_all=True, append_images=ims[1:])
check(load_book(settings(source_mode="file", source_path=scan)).fixed_pages == 3, "a scanned PDF can be kept as it is")
try:
    load_book(settings(source_mode="file", source_path=scan, pdf_mode="Rebuild as an editable book"))
    check(False, "a scanned PDF can't be rebuilt")
except ValueError as e:
    check("scanned" in str(e) and "OCR" in str(e), "rebuilding a scanned PDF explains why not and what to do")
locked = os.path.join(TMP, "locked.pdf")
wr = PdfWriter(clone_from=tricky)
wr.encrypt(user_password="secret", owner_password="o")
wr.write(locked)
try:
    load_book(settings(source_mode="file", source_path=locked))
    check(False, "a password-protected PDF is refused with a message")
except ValueError as e:
    check("password" in str(e), "a password-protected PDF is refused with a message")

# ---------------------------------------------------------------- 6. PDF: rebuild as an editable book
s = settings(source_mode="paste", toc=True, running_heads=True, fmt_impose=False)
src_book = apply_info(load_book(s, chapters_text(5, 120)), s)
typeset = os.path.join(TMP, "typeset.pdf")
build_pdf(src_book, s, typeset, log=lambda m: None)
s = settings(source_mode="file", source_path=typeset, pdf_mode="Rebuild as an editable book", skip_contents=True)
rb = load_book(s, log=lambda m: None)


def words(b):
    return re.findall(r"\w+", " ".join([c.title for c in b.chapters] +
                                         [plain(x.text) for c in b.chapters for x in c.blocks]).lower())


check([c.title for c in rb.chapters] == [c.title for c in src_book.chapters],
      f"rebuilt chapters = the original chapters {[c.title for c in rb.chapters]}")
ratio = difflib.SequenceMatcher(None, words(src_book), words(rb), autojunk=False).ratio()
check(ratio > 0.995, f"rebuilt text matches the original word for word ({ratio:.4f}): hyphenated words "
                     "joined, page numbers and running heads left out")

# ---------------------------------------------------------------- 7. keeping the source's look
d = docx.Document()
d.styles["Normal"].font.size = docx.shared.Pt(11)
d.add_heading("Chapter 1", 1)
p = d.add_paragraph("Normal start, ")
run = p.add_run("big red words")
run.font.size = docx.shared.Pt(16)
run.font.color.rgb = docx.shared.RGBColor(0xC0, 0, 0)
p.add_run(" then normal.")
p = d.add_paragraph("A centred line.")
p.alignment = 1
p = d.add_paragraph("An indented letter.")
p.paragraph_format.left_indent = docx.shared.Inches(0.75)
d.add_heading("Chapter 2", 1)
d.add_paragraph("More words.")
fmt_doc = os.path.join(TMP, "fmt.docx")
d.save(fmt_doc)
plain_b = load_book(settings(source_mode="file", source_path=fmt_doc))
kept_b = load_book(settings(source_mode="file", source_path=fmt_doc, keep_format=True))
pb = plain_b.chapters[0].blocks
kb = kept_b.chapters[0].blocks
check("<font" not in pb[0].text and not pb[1].align and not pb[2].indent, "option off: the book's own look only")
check('<font size="145%" color="#c00000">big red words</font>' in kb[0].text and kb[1].align == "center" and
      kb[2].indent == 3, f"option on: size, colour, centring and indent kept {kb[0].text!r}")
s = settings(source_mode="file", source_path=fmt_doc, keep_format=True, toc=False, title_page=False,
             bookplate="None", fmt_impose=False)
out = os.path.join(TMP, "fmt.pdf")
build_pdf(apply_info(kept_b, s), s, out, log=lambda m: None)
chars = [ch for p in pdf_pages(out) for ch in p["chars"]]
red = [ch for ch in chars if ch["text"] == "b" and ch.get("non_stroking_color") and
       tuple(round(v, 2) for v in ch["non_stroking_color"])[:1] == (0.75,)]
body = float(s["body_size"])
check(red and abs(red[0]["size"] - body * 1.45) < 0.2, "kept words print bigger and red in the PDF")
html_doc = os.path.join(TMP, "fmt.html")
open(html_doc, "w", encoding="utf-8").write(
    "<html><head><title>T</title><style>.c{text-align:center} .big{font-size:1.5em;color:#a00000}</style></head>"
    "<body><h1>One</h1><p>Start <span class='big'>big</span> and <span class='italic'>slanted</span>.</p>"
    "<p class='c'>Centred.</p><h1>Two</h1><p>More.</p></body></html>")
hb = load_book(settings(source_mode="file", source_path=html_doc, keep_format=True)).chapters[0].blocks
check('<font size="150%" color="#a00000">big</font>' in hb[0].text and "<i>slanted</i>" in hb[0].text and
      hb[1].align == "center", "web pages / EPUB: style sheet sizes, colours and centring kept")

# ---------------------------------------------------------------- 8. dark mode: dropdowns stay dark after a choice
try:
    import tkinter as tk
    root = tk.Tk()
    have_display = True
except Exception as e:                     # noqa: BLE001
    have_display = False
    print(f"SKIP  dropdown colour checks (no display: {e})")
if have_display:
    import bf_theme as th
    from tkinter import ttk
    th.set_mode(True)
    th.apply(root)
    st = ttk.Style(root)
    sel = st.lookup(".", "selectbackground")
    unfocused = [v for state, v in [(tuple(x[:-1]), x[-1]) for x in st.map("src.TCombobox", "selectbackground")]
                 if "!focus" in state]
    lum = lambda c: sum(int(c[i:i + 2], 16) for i in (1, 3, 5)) / 3      # noqa: E731
    check(lum(sel) < 128, f"dark mode: selected text background is dark ({sel})")
    check(unfocused and lum(str(unfocused[0])) < 128,
          f"dark mode: a chosen dropdown value stays dark after clicking away ({unfocused})")
    root.destroy()

# ---------------------------------------------------------------- 9. update check (a local stand-in for GitHub)
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer    # noqa: E402
ANSWERS = {
    "/newer": (200, {"tag_name": "v9.0.1", "name": "BookForge 9.0.1", "body": "## New\n- **Faster**",
                     "html_url": "https://github.com/chunksbabyruth/BookForge/releases/tag/v9.0.1"}),
    "/same": (200, {"tag_name": f"v{APP_VERSION}"}),
    "/odd": (200, {"tag_name": "v0.1", "html_url": "https://example.com/not-github"}),
    "/none": (404, {}), "/limit": (403, {}), "/junk": (200, None),
}


class H(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        code, body = ANSWERS[self.path]
        self.send_response(code)
        self.end_headers()
        self.wfile.write(b"<html>" if body is None else json.dumps(body).encode())


for k in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy", "ALL_PROXY", "all_proxy"):
    os.environ.pop(k, None)
srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
threading.Thread(target=srv.serve_forever, daemon=True).start()
base = f"http://127.0.0.1:{srv.server_address[1]}"
newer, info = bf_update.check(APP_VERSION, base + "/newer")
check(newer and info["version"] == "9.0.1" and "Faster" in bf_update.plain_notes(info["notes"]),
      "a newer release is found, with readable notes")
check(not bf_update.check(APP_VERSION, base + "/same")[0], "the same version is up to date")
newer, info = bf_update.check(APP_VERSION, base + "/odd")
check(not newer and info["url"] == bf_update.RELEASES_PAGE, "only GitHub pages are ever opened")
check(bf_update.is_newer("v1.10.0", "1.9.9") and not bf_update.is_newer("v1.5", "1.6.0"), "versions compare as numbers")
for path, word in (("/none", "No BookForge releases"), ("/limit", "limiting"), ("/junk", "couldn't be read")):
    try:
        bf_update.check(APP_VERSION, base + path)
        check(False, f"{path} gives a message")
    except bf_update.UpdateError as e:
        check(word in str(e), f"{path}: {e}")
try:
    bf_update.check(APP_VERSION, "http://127.0.0.1:9/", timeout=2)
    check(False, "offline gives a message")
except bf_update.UpdateError as e:
    check("couldn't be reached" in str(e), f"offline: {e}")
srv.shutdown()
src = open(os.path.join(ROOT, "bf_update.py"), encoding="utf-8").read()
check(not re.search(r"urlretrieve|subprocess|os\.system|startfile|exec\(", src),
      "the update check never downloads or runs anything")

# ---------------------------------------------------------------- 10. version
check(tuple(int(x) for x in APP_VERSION.split(".")) >= (1, 6, 0), f"version is 1.6.0 or later ({APP_VERSION})")
vi = open(os.path.join(ROOT, "version_info.txt"), encoding="utf-8").read()
check(f"'{APP_VERSION}.0'" in vi, f"version_info.txt says {APP_VERSION}.0")

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
