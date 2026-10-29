"""Checks that the title settings reach every format in 1.8.0 (plain Python - run: python tests/test_title_exports_180.py).

Layout -> "Titles: alignment & position" (title page, chapter and part titles, the headings BookForge adds) and
Output -> Cover ("Cover title alignment", "Cover title position"). These checks read the files themselves:

1. PDF: the title page, chapter and part titles where the settings put them; preview = Print PDF = imposed sheets.
2. EPUB: the stylesheet and the markup inside the .epub, worked out as a reader applies them (cascade,
   specificity, inline styles, inheritance): which alignment, size and space each title really gets.
3. Word: the Heading styles and the paragraphs (alignment, space above = the PDF's place on the page, sizes);
   where LibreOffice is installed, the Word file is laid out by it and the titles are measured on its pages.
4. Web page: the same as the EPUB, from the .html file.
5. Cover: the front cover title's alignment and position, on their own: the chapter / part settings don't move
   it; "Same as title page" follows the title page alignment; the standard values draw it where it always was.
6. Things that keep their own look: headings typed in the text, the dedication, the bookplate, THE END,
   section breaks; the plain text file has no layout and is unchanged.
7. Saved in projects and back after reopening; an older project without these settings gets the standard
   values (the PDF exactly as before).
Window checks run only where a display is available. Nothing needs the internet.
"""
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from html.parser import HTMLParser

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="bf_titlex_")
os.environ["APPDATA"] = os.path.join(TMP, "appdata")
os.environ["HOME"] = os.path.join(TMP, "home")
os.environ["BOOKFORGE_NO_UPDATE_CHECK"] = "1"
os.makedirs(os.environ["APPDATA"], exist_ok=True)
os.makedirs(os.path.join(os.environ["HOME"], "Documents"), exist_ok=True)

import docx                                                           # noqa: E402
import pdfplumber                                                     # noqa: E402
import pypdfium2 as pdfium                                            # noqa: E402

import bf_edit                                                        # noqa: E402
import bf_export                                                      # noqa: E402
import bf_matter                                                      # noqa: E402
from bf_build import build_preview, run_build                         # noqa: E402
from bf_core import DEFAULTS, Book, Chapter, Block, apply_info, page_size_inches, TITLE_LAYOUT_KEYS  # noqa: E402

FAILS = []


def check(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg, flush=True)
    if not cond:
        FAILS.append(msg)


DIGEST = "Half Letter / Digest (5.5 x 8.5 in)"
ALL = dict(fmt_pdf=True, fmt_impose=True, fmt_cover=True, fmt_epub=True, fmt_docx=True, fmt_html=True,
           fmt_txt=True)
CHOSEN = dict(tp_align="Left", tp_pos=40.0, tp_title_size=40.0, ch_align="Right", chapter_sink=1.0,
              part_align="Left", part_pos=50.0, part_size=30.0, cover_align="Right", cover_pos=50.0)


def settings(**kw):
    s = copy.deepcopy(DEFAULTS)
    s.update(open_when_done=False, out_dir=tempfile.mkdtemp(dir=TMP), source_mode="paste", page_size=DIGEST,
             margin_top=0.7, margin_bottom=0.7, margin_inner=0.7, margin_outer=0.7, chapter_label="Chapter One",
             toc=True, fic_info="Add a work info page", dedication="For the reader", note_title="Zhnote",
             note_text="Note words.", the_end=True, bookplate="On title page", imp_files="Aggregated file only",
             **ALL)
    s.update(kw)
    return s


def book():
    b = Book(title="Zqtitle", subtitle="Zqsubtitle", author="Zqauthor", credit="Zqcredit",
             fic_meta=["<b>Rating:</b> G"], fic_summary=[Block("para", text="Summary words.", label="Summary")],
             chapters=[Chapter("Zqpart", label="Part One", level=0, blocks=[Block("para", text="Part words.")]),
                       Chapter("Zqarrival", blocks=[Block("subtitle", text="Zqchsub"),
                                                    Block("para", text="Arrival words go here. " * 40),
                                                    Block("head", text="Zqinner"),
                                                    Block("scene", text="* * *"),
                                                    Block("para", text="More words. " * 60)]),
                       Chapter("Zqdeparture", blocks=[Block("para", text="Departure words. " * 150)])])
    bf_matter.set_override(b, "end_matter", [Block("para", text="Afterword words.")])
    bf_matter.set_heading(b, "end_matter", "Zhafter")
    return b


def make(b, s):
    s = dict(s)
    s["out_dir"] = tempfile.mkdtemp(dir=TMP)
    folder, made, notes = run_build(apply_info(copy.deepcopy(b), dict(s)), s, log=lambda m: None)
    get = lambda end: next((m for m in made if m.endswith(end)), None)     # noqa: E731
    return {"pdf": get("Print.pdf"), "imposed": get("Imposed for binding.pdf"), "cover": get("Cover.pdf"),
            "epub": get(".epub"), "docx": get(".docx"), "html": get(".html"),
            "txt": next((m for m in made if m.endswith(".txt") and "READ ME" not in m), None)}


def words(path):
    out = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            for w in pg.extract_words(extra_attrs=["size"]):
                out.append((pg.page_number, w, pg.width, pg.height))
    return out


def word(path, text):
    """(page, word, page width, page height) of the biggest word starting with text (a title, not its contents
    entry); the first of those."""
    found = [(p, w, W, H) for p, w, W, H in words(path) if w["text"].startswith(text)]
    big = max((w["size"] for _p, w, _W, _H in found), default=0)
    return next((f for f in found if f[1]["size"] == big), None)


def chars(path):
    with pdfplumber.open(path) as pdf:
        return [[(c["text"], round(c["x0"], 2), round(c["top"], 2), c["fontname"].split("+")[-1], round(c["size"], 2))
                 for c in pg.chars] for pg in pdf.pages]


def renders(path):
    doc = pdfium.PdfDocument(path)
    out = [doc[i].render(scale=0.25).to_pil().tobytes() for i in range(len(doc))]
    doc.close()
    return out


# ---- reading HTML / XHTML as a reader does: which rules apply to an element (a small CSS cascade)
class Paths(HTMLParser):
    """For each piece of text: the elements it is in, from the outside in: (tag, classes, inline style)."""
    VOID = {"br", "img", "hr", "meta", "link", "input"}

    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.stack, self.texts, self.css = [], [], []
        self.in_style = False
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        if tag == "style":
            self.in_style = True
        if tag in self.VOID:
            return
        a = dict(attrs)
        self.stack.append((tag, set((a.get("class") or "").split()), a.get("style") or ""))

    def handle_endtag(self, tag):
        if tag == "style":
            self.in_style = False
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        if self.in_style:
            self.css.append(data)
        elif data.strip() and not any(e[0] in ("head", "nav") for e in self.stack):
            self.texts.append((data.strip(), list(self.stack)))   # (not the page title or the contents list)

    def path(self, text):
        return next((p for t, p in self.texts if t == text), None)


def css_rules(css):
    out = []
    for n, m in enumerate(re.finditer(r"([^{}]+)\{([^{}]*)\}", css)):
        decls = dict((k.strip(), v.strip()) for k, v in (d.split(":", 1) for d in m.group(2).split(";") if ":" in d))
        for sel in m.group(1).split(","):
            out.append((sel.strip(), decls, n))
    return out


def _compound(comp, el):
    m = re.fullmatch(r"([a-z0-9]*)((?:\.[\w-]+)*)", comp)
    if not m:
        return False
    tag, classes = m.group(1), m.group(2).split(".")[1:]
    return (not tag or tag == el[0]) and all(c in el[1] for c in classes)


def _matches(sel, path):
    comps = sel.split()
    if not comps or not _compound(comps[-1], path[-1]):
        return False
    i = len(path) - 2
    for comp in reversed(comps[:-1]):
        while i >= 0 and not _compound(comp, path[i]):
            i -= 1
        if i < 0:
            return False
        i -= 1
    return True


def _spec(sel):
    comps = sel.split()
    return (sum(c.count(".") for c in comps), sum(1 for c in comps if re.match(r"[a-z]", c)))


INHERITED = {"text-align", "font-size", "text-indent"}


def computed(css, path, prop):
    """The value of prop for the innermost element of path: its inline style, else the rule that wins (highest
    specificity, then the last one), else (inherited properties) its parent's value."""
    rules = css_rules(css)
    for depth in range(len(path), 0, -1):
        sub = path[:depth]
        inline = dict((k.strip(), v.strip()) for k, v in (d.split(":", 1) for d in sub[-1][2].split(";")
                                                          if ":" in d))
        if prop in inline:
            return inline[prop]
        best = None
        for sel, decls, n in rules:
            if prop in decls and _matches(sel, sub):
                key = (_spec(sel), n)
                if best is None or key >= best[0]:
                    best = (key, decls[prop])
        if best:
            return best[1]
        if prop not in INHERITED:
            return None
    return None


def epub_docs(path):
    """{text file name: Paths}, and the stylesheet."""
    z = zipfile.ZipFile(path)
    css = z.read("OEBPS/style.css").decode()
    return {n.split("/")[-1]: Paths(z.read(n).decode()) for n in z.namelist() if n.endswith(".xhtml")}, css


def epub_look(path, text, prop, nav=False):
    """(value of prop for the element holding exactly text, its path); nav: look in the contents page."""
    docs, css = epub_docs(path)
    if nav:
        z = zipfile.ZipFile(path)
        h = z.read("OEBPS/nav.xhtml").decode()
        m = re.search(r"<nav[^>]*><h2>([^<]*)</h2>", h)
        p = [("html", set(), ""), ("body", set(), ""), ("nav", set(), ""), ("h2", set(), "")]
        return (computed(css, p, prop), p) if m and m.group(1) == text else (None, None)
    for d in docs.values():
        p = d.path(text)
        if p:
            return computed(css, p, prop), p
    return None, None


def html_look(path, text, prop):
    d = Paths(open(path, encoding="utf-8").read())
    p = d.path(text)
    return (computed("".join(d.css), p, prop), p) if p else (None, None)


S0 = settings()
S1 = settings(**CHOSEN)
B = book()
F0 = make(B, S0)
F1 = make(B, S1)
pw, ph = page_size_inches(S1)
MT, FH = 0.7 * 72, (ph - 1.4) * 72                  # text area: top margin, height (pt)
LEFT, RIGHT = 0.7 * 72, (pw - 0.7) * 72

# ============================================================================ 1. PDF (and preview, imposed)
t1, t0 = word(F1["pdf"], "Zqtitle"), word(F0["pdf"], "Zqtitle")
check(t1 and abs(t1[1]["x0"] - LEFT) < 1.5 and abs(t1[1]["size"] - 40) < 0.1 and
      abs(t1[1]["top"] - (MT + 0.4 * FH)) < 12 and t0 and abs(t0[1]["top"] - (MT + 0.2 * FH)) < 12 and
      abs((t0[1]["x0"] + t0[1]["x1"]) / 2 - (LEFT + RIGHT) / 2) < 1.5,
      "PDF: book title left, 40 pt, 40 % down (standard: centred, 20 % down)")
c1, c0 = word(F1["pdf"], "Zqarrival"), word(F0["pdf"], "Zqarrival")
check(c1 and abs(c1[1]["x1"] - RIGHT) < 1.5 and c0 and 70 < c1[1]["top"] - c0[1]["top"] < 80,
      "PDF: chapter title right, dropped 1 in")
p1 = word(F1["pdf"], "Zqpart")
check(p1 and abs(p1[1]["x0"] - LEFT) < 1.5 and abs(p1[1]["size"] - 30) < 0.1 and
      MT + 0.5 * FH - 30 < p1[1]["top"] < MT + 0.5 * FH + 30, "PDF: part title left, 30 pt, about 50 % down")
res = build_preview(apply_info(copy.deepcopy(B), dict(S1)), dict(S1), log=lambda m: None)
check(chars(res["pdf"]) == chars(F1["pdf"]) and len(res["sheets"]) == 1 and
      renders(res["sheets"][0]) == renders(F1["imposed"]),
      "preview = Print PDF = imposed sheets (rendered side by side) with the title settings")

# ============================================================================ 2. EPUB
A = {}                                          # what a reader applies: text-align of each title
for name, f in (("chosen", F1), ("standard", F0)):
    A[name] = {k: epub_look(f["epub"], txt, "text-align")[0]
               for k, txt in (("title", "Zqtitle"), ("subtitle", "Zqsubtitle"), ("author", "Zqauthor"),
                              ("label", "Chapter One"), ("chapter", "Zqarrival"), ("chsub", "Zqchsub"),
                              ("part", "Zqpart"), ("partlabel", "Part One"),
                              ("workinfo", "Work Info"), ("note", "Zhnote"), ("after", "Zhafter"),
                              ("inner", "Zqinner"), ("ded", "For the reader"), ("end", "THE END"))}
    A[name]["contents"] = epub_look(f["epub"], "CONTENTS", "text-align", nav=True)[0]
want = dict(title="left", subtitle="left", author="left", label="right", chapter="right", chsub="right",
            part="left", partlabel="left", contents="right", workinfo="right", note="right", after="right",
            inner="center", ded="center", end="center")
check(A["chosen"] == want, f"EPUB: each title's alignment as a reader applies the stylesheet {A['chosen']}")
check(all(v == "center" for v in A["standard"].values()), "EPUB, standard settings: every title centred, as before")
_d, css1 = epub_docs(F1["epub"])
docs0, css0 = epub_docs(F0["epub"])
size = lambda f, t: epub_look(f, t, "font-size")[0]                        # noqa: E731
check(size(F1["epub"], "Zqtitle") == f"{40 / 13:.4g}em" and size(F0["epub"], "Zqtitle") == "2.2em" and
      size(F1["epub"], "Zqpart") == f"{30 / 13:.4g}em" and size(F0["epub"], "Zqpart") is None,
      "EPUB: book title and part title sizes in ems of the text (standard: the sizes as before)")
mt = lambda f, t, prop: computed(epub_docs(f)[1], epub_look(f, t, prop)[1][:-1], prop)   # noqa: E731
check(mt(F1["epub"], "Zqtitle", "margin-top") == "50%" and mt(F0["epub"], "Zqtitle", "margin-top") == "25%" and
      mt(F1["epub"], "Zqpart", "margin-top") == "50%" and mt(F0["epub"], "Zqpart", "margin-top") == "30%" and
      mt(F1["epub"], "Zqarrival", "padding-top") == f"{72 / 13:.4g}em" and
      mt(F0["epub"], "Zqarrival", "padding-top") is None,
      "EPUB: positions as space above (title page 40 % -> 50 %, standard 25 %; part 50 %, standard 30 %; "
      "chapter drop 1 in -> 5.5 em)")
check(epub_look(F1["epub"], "Chapter One", "text-indent")[0] == "0" and
      epub_look(F1["epub"], "Part One", "text-indent")[0] == "0" and
      html_look(F1["html"], "Chapter One", "text-indent")[0] == "0",
      "EPUB / web page: the number lines have no first-line indent, so they line up with their titles")
check(re.search(r"<div class=\"chead\">\s*<p class=\"label\">Chapter One</p>\s*<h2>Zqarrival</h2>\s*</div>",
                zipfile.ZipFile(F1["epub"]).read("OEBPS/ch002.xhtml").decode()) is not None,
      "EPUB: the chapter's number line and title are one block (the drop space goes above both)")

# ============================================================================ 3. Word
D1, D0 = docx.Document(F1["docx"]), docx.Document(F0["docx"])


def para(d, text):
    return next((p for p in d.paragraphs if p.text.strip().startswith(text) or
                 p.text.replace("\n", " ").strip().endswith(text)), None)


AL = {0: "left", 1: "center", 2: "right", None: None}


def walign(d, p):
    a = p.alignment if p.alignment is not None else p.style.paragraph_format.alignment
    return AL.get(int(a) if a is not None else None)


W1 = {k: walign(D1, para(D1, t)) for k, t in (("title", "Zqtitle"), ("subtitle", "Zqsubtitle"), ("author", "Zqauthor"),
                                              ("chapter", "Zqarrival"), ("chsub", "Zqchsub"), ("part", "Zqpart"),
                                              ("contents", "CONTENTS"), ("workinfo", "Work Info"),
                                              ("note", "Zhnote"), ("after", "Zhafter"), ("inner", "Zqinner"),
                                              ("ded", "For the reader"), ("plate", "This Book Belongs To"),
                                              ("end", "THE END"))}
check(W1 == dict(title="left", subtitle="left", author="left", chapter="right", chsub="right", part="left",
                 contents="right", workinfo="right", note="right", after="right", inner="center", ded="center",
                 plate="center", end="center"), f"Word: each title paragraph's alignment {W1}")
h1, h2 = D1.styles["Heading 1"], D1.styles["Heading 2"]
check(int(h1.paragraph_format.alignment) == 0 and int(h2.paragraph_format.alignment) == 2 and
      h1.font.size.pt == 30 and h2.font.size.pt == 23 and
      int(D0.styles["Heading 1"].paragraph_format.alignment) == 1 and
      int(D0.styles["Heading 2"].paragraph_format.alignment) == 1 and
      abs(D0.styles["Heading 1"].font.size.pt - 23 * 1.2) <= 0.25,           # (Word keeps half points)
      "Word: the Heading styles carry it (parts = Heading 1: left, 30 pt; chapters = Heading 2: right; standard: "
      "centred, part titles 1.2 x the chapter title size as in the PDF)")
sb = lambda d, t: (para(d, t).paragraph_format.space_before or 0) / 12700           # noqa: E731
own = (D1.styles["Heading 2"].paragraph_format.space_before or 0) / 12700
check(abs(sb(D1, "Zqtitle") - 0.4 * FH) < 0.5 and abs(sb(D0, "Zqtitle") - 0.2 * FH) < 0.5 and
      abs(sb(D1, "Zqpart") - 0.5 * FH) < 0.5 and abs(sb(D0, "Zqpart") - 0.3 * FH) < 0.5 and
      abs(sb(D1, "Zqarrival") - (72 + own)) < 0.5 and para(D0, "Zqarrival").paragraph_format.space_before is None,
      "Word: positions as the PDF's on the page (title 40 % / part 50 % of the text area, chapter +1 in)")
check(any(r.font.size and r.font.size.pt == 40 for r in para(D1, "Zqtitle").runs) and
      any(r.font.size and r.font.size.pt == max(30, 23 * 1.5) for r in para(D0, "Zqtitle").runs),
      "Word: book title size 40 pt (standard: automatic, as in the PDF)")
if shutil.which("soffice"):
    lo = tempfile.mkdtemp(dir=TMP)
    lo1, lo0 = os.path.join(lo, "chosen.pdf"), os.path.join(lo, "standard.pdf")
    try:
        shutil.copy(F1["docx"], os.path.join(lo, "chosen.docx"))
        shutil.copy(F0["docx"], os.path.join(lo, "standard.docx"))
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", lo,
                        os.path.join(lo, "chosen.docx"), os.path.join(lo, "standard.docx")],
                       capture_output=True, timeout=240)
    except (OSError, subprocess.SubprocessError):
        pass
    if os.path.isfile(lo1) and os.path.isfile(lo0):
        a, b0 = word(lo1, "Zqtitle"), word(lo0, "Zqtitle")
        c, c0w = word(lo1, "Zqarrival"), word(lo0, "Zqarrival")
        pt = word(lo1, "Zqpart")
        ok = bool(a and b0 and c and c0w and pt)
        check(ok and abs(a[1]["x0"] - LEFT) < 3 and abs(a[1]["top"] - t1[1]["top"]) < 12 and
              abs(b0[1]["top"] - t0[1]["top"]) < 12 and abs(c[1]["x1"] - RIGHT) < 3 and
              66 < c[1]["top"] - c0w[1]["top"] < 78 and abs(pt[1]["top"] - p1[1]["top"]) < 30,
              "Word laid out by LibreOffice: title left at the PDF's height, chapter title right and 1 in lower, "
              "part title at the PDF's height" + (f" ({a[1]['top']:.0f}/{t1[1]['top']:.0f}, "
                                                  f"{c[1]['top'] - c0w[1]['top']:.0f}, "
                                                  f"{pt[1]['top']:.0f}/{p1[1]['top']:.0f})" if ok else ""))
    else:
        print("SKIP  Word laid out by LibreOffice (it couldn't convert the files here)")
else:
    print("SKIP  Word laid out by LibreOffice (LibreOffice isn't installed here)")

# ============================================================================ 4. web page
H1 = {k: html_look(F1["html"], t, "text-align")[0]
      for k, t in (("title", "Zqtitle"), ("author", "Zqauthor"), ("label", "Chapter One"), ("chapter", "Zqarrival"),
                   ("chsub", "Zqchsub"), ("part", "Zqpart"), ("partlabel", "Part One"), ("contents", "CONTENTS"),
                   ("workinfo", "Work Info"), ("note", "Zhnote"), ("after", "Zhafter"), ("inner", "Zqinner"),
                   ("ded", "For the reader"), ("end", "THE END"))}
check(H1 == dict(title="left", author="left", label="right", chapter="right", chsub="right", part="left",
                 partlabel="left", contents="right", workinfo="right", note="right", after="right", inner="center",
                 ded="center", end="center"), f"web page: each title's alignment as a browser applies the CSS {H1}")
H0 = [html_look(F0["html"], t, "text-align")[0] for t in ("Zqtitle", "Zqarrival", "Zqpart", "CONTENTS")]
hp = Paths(open(F1["html"], encoding="utf-8").read())
check(H0 == ["center"] * 4 and html_look(F1["html"], "Zqtitle", "font-size")[0] == f"{40 / 13:.4g}em" and
      computed("".join(hp.css), hp.path("Zqpart")[:-1], "margin-top") == "50%" and
      computed("".join(hp.css), hp.path("Zqarrival")[:-1], "padding-top") == f"{72 / 13:.4g}em",
      "web page: standard = centred; sizes and positions as in the EPUB")

# ============================================================================ 5. cover
cv1, cv0 = word(F1["cover"], "Zqtitle"), word(F0["cover"], "Zqtitle")
with pdfplumber.open(F1["cover"]) as pdf:
    CW, CH = pdf.pages[0].width, pdf.pages[0].height
pad = 0.6 * 72
check(cv1 and abs(cv1[1]["x1"] - (CW - pad)) < 2 and abs(cv1[1]["bottom"] - 0.5 * CH) < cv1[1]["size"] * 0.4 and
      word(F1["cover"], "Zqauthor") and abs(word(F1["cover"], "Zqauthor")[1]["x1"] - (CW - pad)) < 2,
      "cover: title and author lines right-aligned on the front, the title line 50 % down")
cover_std = dict(S1, cover_align="Same as title page", cover_pos=32.0, tp_align="Centered")
c_std = os.path.join(TMP, "cover_std.pdf")
c_ref = os.path.join(TMP, "cover_ref.pdf")
bk = apply_info(copy.deepcopy(B), dict(S0))
bf_export.build_cover(bk, dict(cover_std, fmt_impose=False), c_std, 40)
bf_export.build_cover(bk, dict(S0, fmt_impose=False), c_ref, 40)
cs = word(c_std, "Zqtitle")
with pdfplumber.open(c_std) as pdf:
    W2, H2 = pdf.pages[0].width, pdf.pages[0].height
tw = (W2 - bf_export.estimated_spine(S0, 40) * 72) / 2
check(cv0 and cs and chars(c_std) == chars(c_ref) and abs((cs[1]["x0"] + cs[1]["x1"]) / 2 - (W2 - tw / 2)) < 1.5 and
      abs(cs[1]["bottom"] - 0.32 * H2) < cs[1]["size"] * 0.4,
      "cover, standard values: centred, 32 % down, whatever the chapter, part and title page sizes/positions")
c_left = os.path.join(TMP, "cover_left.pdf")
bf_export.build_cover(bk, dict(S0, fmt_impose=False, tp_align="Left"), c_left, 40)
c_cent = os.path.join(TMP, "cover_cent.pdf")
bf_export.build_cover(bk, dict(S0, fmt_impose=False, tp_align="Left", cover_align="Centered"), c_cent, 40)
cl, cc = word(c_left, "Zqtitle"), word(c_cent, "Zqtitle")
check(cl and abs(cl[1]["x0"] - (W2 - tw + pad)) < 2 and cc and chars(c_cent) == chars(c_ref),
      "cover: “Same as title page” follows the title page alignment; “Centered” keeps it centred")
c_hi = os.path.join(TMP, "cover_hi.pdf")
bf_export.build_cover(bk, dict(S0, fmt_impose=False, cover_pos=5.0), c_hi, 40)
c_lo = os.path.join(TMP, "cover_lo.pdf")
bf_export.build_cover(bk, dict(S0, fmt_impose=False, cover_pos=85.0), c_lo, 40)
hi, lo_, au = word(c_hi, "Zqtitle"), word(c_lo, "Zqsubtitle"), word(c_lo, "Zqauthor")
check(hi and hi[1]["top"] >= pad - 1 and lo_ and au and lo_[1]["bottom"] < au[1]["top"],
      "cover: the title stays on the cover at the top, and the subtitle never runs into the author line")

# ============================================================================ 6. what keeps its own look; text file
check(open(F1["txt"], encoding="utf-8").read() == open(F0["txt"], encoding="utf-8").read(),
      "plain text file: no layout, the same with any title settings")
check(epub_look(F1["epub"], "Zqinner", "text-align")[0] == "center" and W1["inner"] == "center" and
      H1["inner"] == "center" and W1["ded"] == W1["plate"] == W1["end"] == "center",
      "headings typed in the text, the dedication, the bookplate and THE END keep their own (centred) look")

# ============================================================================ 7. projects
proj = os.path.join(TMP, "titles.bookforge")
bf_edit.save_project(proj, B, S1, "1.8.0")
b7, s7 = bf_edit.load_project(proj)
check(all(s7[k] == S1[k] for k in TITLE_LAYOUT_KEYS + ("chapter_sink",)), "saved in a project and read back")
F7 = make(b7, dict(s7))
check(chars(F7["pdf"]) == chars(F1["pdf"]) and
      zipfile.ZipFile(F7["epub"]).read("OEBPS/style.css") == zipfile.ZipFile(F1["epub"]).read("OEBPS/style.css") and
      [(p.text, p.alignment, p.paragraph_format.space_before) for p in docx.Document(F7["docx"]).paragraphs] ==
      [(p.text, p.alignment, p.paragraph_format.space_before) for p in D1.paragraphs] and
      chars(F7["cover"]) == chars(F1["cover"]),
      "the reopened project makes the same PDF, EPUB stylesheet, Word paragraphs and cover")
old = {k: v for k, v in S1.items() if k not in TITLE_LAYOUT_KEYS and k != "chapter_sink"}
bf_edit.save_project(os.path.join(TMP, "old.bookforge"), B, old, "1.7.0")
b8, s8 = bf_edit.load_project(os.path.join(TMP, "old.bookforge"))
F8 = make(b8, dict(s8))
check(chars(F8["pdf"]) == chars(F0["pdf"]) and chars(F8["cover"]) == chars(F0["cover"]) and
      zipfile.ZipFile(F8["epub"]).read("OEBPS/style.css") == zipfile.ZipFile(F0["epub"]).read("OEBPS/style.css") and
      [(p.text, p.alignment, p.paragraph_format.space_before) for p in docx.Document(F8["docx"]).paragraphs] ==
      [(p.text, p.alignment, p.paragraph_format.space_before) for p in D0.paragraphs],
      "an older project without the title settings: the standard values in every format (PDF and cover as before)")

# ============================================================================ window
try:
    import tkinter as tk
    _probe = tk.Tk()
    _probe.destroy()
    have_display = True
except Exception as e:                     # noqa: BLE001
    have_display = False
    print(f"SKIP  window checks (no display: {e})")

if have_display:
    from tkinter import messagebox, filedialog
    SAID = []
    for k in ("askyesno", "askyesnocancel", "showinfo", "showerror", "showwarning", "askokcancel"):
        setattr(messagebox, k, (lambda kind: (lambda title="", message="", **kw: SAID.append((kind, title, message))
                                              or True))(k))
    PICK = []
    filedialog.askopenfilename = lambda **kw: PICK.pop(0) if PICK else ""
    filedialog.asksaveasfilename = lambda **kw: PICK.pop(0) if PICK else ""
    import BookForge as BF
    import bf_theme as th
    app = BF.App()
    app._no_save = True
    errors = []
    app.report_callback_exception = lambda *a: errors.append("".join(__import__("traceback").format_exception(*a)))

    def pump(sec=0.2):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            time.sleep(0.01)

    def idle(sec=240):
        end = time.time() + sec
        while time.time() < end:
            app.update()
            if not app.busy:
                return True
            time.sleep(0.02)
        return False

    app.var("source_mode").set("paste")
    app._mode_changed()
    app.texts["paste_text"].insert("1.0", "".join(f"Chapter {i}\n\n" + "Words of the chapter. " * 80 + "\n\n"
                                                  for i in range(1, 4)))
    app.var("title").set("Window Title")
    app.scan(force=True)
    idle(60)
    app.shelf.select(th.tab_index("out"))
    pump(0.2)
    combos = []
    todo = [app.pages[th.tab_index("out")]]
    while todo:
        w = todo.pop()
        todo.extend(w.winfo_children())
        if w.winfo_class() == "TCombobox" and str(w.cget("textvariable")) == str(app.var("cover_align")):
            combos.append(w)
    check(combos and list(combos[0].cget("values")) == ["Same as title page", "Centered", "Left", "Right"] and
          app.var("cover_align").get() == "Same as title page" and app.var("cover_pos").get() == 32.0,
          "Output > Cover has the cover title alignment (Same as title page / Centered / Left / Right) and position")
    for k, v in CHOSEN.items():
        app.var(k).set(v)
    outdir = tempfile.mkdtemp(dir=TMP)
    app.var("out_dir").set(outdir)
    app.var("open_when_done").set(False)
    for k, v in dict(ALL, fmt_impose=False, fmt_txt=False).items():
        app.var(k).set(v)
    idle()
    app.build()
    idle()
    pump(0.2)
    got = {os.path.splitext(f)[1]: os.path.join(r, f) for r, _d, fs in os.walk(outdir) for f in fs}
    ep = got.get(".epub")
    cv = next((os.path.join(r, f) for r, _d, fs in os.walk(outdir) for f in fs if f.endswith("Cover.pdf")), None)
    wd = got.get(".docx")
    check(ep and epub_look(ep, "Window Title", "text-align")[0] == "left" and
          epub_look(ep, "Chapter 2", "text-align")[0] == "right" and wd and
          walign(None, para(docx.Document(wd), "Window Title")) == "left" and cv and
          abs(word(cv, "Title")[1]["x1"] - (pdfplumber.open(cv).pages[0].width - pad)) < 2,
          "the window: Make the book writes the chosen title settings into the EPUB, Word file and cover")
    proj = os.path.join(TMP, "window.bookforge")
    PICK.append(proj)
    idle()
    app.save_project_dialog()
    for k in CHOSEN:
        app.var(k).set(DEFAULTS[k])
    PICK.append(proj)
    app.open_project_dialog()
    pump(0.3)
    check(all(app.var(k).get() == v for k, v in CHOSEN.items()), "the window: saved in a project and back on reopening")
    book9, sett = bf_edit.load_project(proj)
    for k in TITLE_LAYOUT_KEYS:
        sett.pop(k, None)
    old = os.path.join(TMP, "window_old.bookforge")
    bf_edit.save_project(old, book9, sett, "1.7.0")
    PICK.append(old)
    app.open_project_dialog()
    pump(0.3)
    check(all(app.var(k).get() == DEFAULTS[k] for k in TITLE_LAYOUT_KEYS),
          "the window: an older project opens with the standard title and cover settings")
    check(not errors, f"no window errors {errors[:1]}")
    app.destroy()

print(f"\n{len(FAILS)} failure(s)")
sys.exit(1 if FAILS else 0)
